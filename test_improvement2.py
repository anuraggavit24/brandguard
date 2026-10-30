"""Checks for Improvement 2: posters and images, documents, web pages, brand assets and accessibility, OCR upgrades.
Run: python test_improvement2.py"""
import io, json, os, re, subprocess, sys, tempfile
os.environ["DATA_DIR"] = tempfile.mkdtemp()
os.environ.pop("GROQ_API_KEY", None)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import main
import ocr
import formats as FMT
import assets as AST
import generic as GEN
import realestate as RE
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

D = os.environ["DATA_DIR"]
FONT = next((f for f in ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
                          r"C:\Windows\Fonts\arialbd.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf")
             if os.path.exists(f)), "DejaVuSans-Bold.ttf")
fails = []


def check(name, cond, extra=""):
    print(("PASS " if cond else "FAIL ") + name + (f"  [{extra}]" if extra and not cond else ""))
    if not cond:
        fails.append(name)


def font(n):
    return ImageFont.truetype(FONT, n)


def make_logo(path):
    im = Image.new("RGB", (240, 120), "white")
    d = ImageDraw.Draw(im)
    d.ellipse((8, 8, 112, 112), outline=(31, 58, 95), width=10)
    d.rectangle((40, 40, 80, 80), fill=(242, 169, 0))
    d.text((124, 30), "ZM", fill=(31, 58, 95), font=font(56))
    im.save(path)
    return im


def make_poster(path, logo=None, price="Rs 7.1 lakh", low_contrast=True, bg="white", band=None):
    im = Image.new("RGB", (1200, 1600), bg)
    d = ImageDraw.Draw(im)
    if band:
        d.rectangle((0, 1250, 1200, 1600), fill=band)
    d.text((80, 120), "Zippy Hatchback", fill=(20, 20, 20), font=font(96))
    d.text((80, 420), "VXi", fill=(20, 20, 20), font=font(80))
    d.text((80, 530), price, fill=(20, 20, 20), font=font(80))
    d.text((80, 800), "Festive discount 5%", fill=(20, 20, 20), font=font(64))
    if low_contrast:
        d.text((80, 1000), "Terms and conditions apply", fill=(200, 200, 200), font=font(56))
    d.text((80, 1180), "Prices are ex-showroom Pune", fill=(30, 30, 30), font=font(14))
    if logo is not None:
        im.paste(logo.resize((360, 180)), (780, 1380))
    im.save(path)


# ---------------- OCR upgrades
logo_path = os.path.join(D, "logo.png")
logo = make_logo(logo_path)
poster = os.path.join(D, "poster.png")
make_poster(poster, logo)
lines = ocr.read_image(poster)
texts = [l["text"] for l in (lines or [])]
check("poster text is read", lines and any("Zippy Hatchback" in t for t in texts), str(texts))
check("a price on its own line is paired with the label above it",
      any(re.search(r"VXi.*7\.1 lakh", t) for t in texts), str(texts))
check("lines keep their position and height", all(l.get("box") and l.get("text_h") for l in lines))
check("the same text read by two passes is kept once", len(texts) == len({ocr._key(t) for t in texts}), str(texts))
check("missing language packs are skipped, not fatal", "eng" in ocr.langs())
check("pairing does not join two label lines", ocr.pair_lines([
    {"text": "Welcome", "raw": "Welcome", "box": (0, 0, 100, 20)}, {"text": "Visit today", "raw": "Visit today", "box": (0, 30, 100, 20)}])[1]["text"] == "Visit today")

# ---------------- contrast maths
check("WCAG contrast black on white is 21", round(AST.contrast_ratio((0, 0, 0), (255, 255, 255))) == 21)
check("light grey on white fails 4.5", AST.contrast_ratio((200, 200, 200), (255, 255, 255)) < 4.5)
check("palette parsing", AST.parse_palette("#1f3a5f, F2A900 and #fff") == ["#1F3A5F", "#F2A900", "#FFFFFF"])
mk = main.make_finding
cfg_reg = {"assets": {"palette": ["#1F3A5F", "#F2A900"], "check_palette": True, "require_logo": True, "logo": logo_path,
                      "min_contrast": 4.5, "min_text_pct": 1.2}}
f, ch = AST.check_image(poster, lines, cfg_reg, mk)
types = [x["type"] for x in f]
check("low-contrast text needs a person", "LOW_CONTRAST" in types, str(types))
check("tiny text gets advice", "SMALL_TEXT" in types, str(types))
check("the logo is found on the poster", "LOGO_MISSING" not in types and any(c[0] == "Brand logo" and "found." in c[2] for c in ch), str(ch))
no_logo = os.path.join(D, "nologo.png")
make_poster(no_logo, None, low_contrast=False)
f2, _ = AST.check_image(no_logo, ocr.read_image(no_logo), cfg_reg, mk)
check("a required logo that is missing needs a person", "LOGO_MISSING" in [x["type"] for x in f2])
check("good contrast is not flagged", "LOW_CONTRAST" not in [x["type"] for x in f2], str([x["evidence"] for x in f2]))
red = os.path.join(D, "red.png")
make_poster(red, logo, low_contrast=False, band=(220, 30, 40))
f3, _ = AST.check_image(red, ocr.read_image(red), cfg_reg, mk)
check("a large non-brand colour is pointed out", "OFF_PALETTE" in [x["type"] for x in f3])
check("white and black are neutral, not off-brand", "OFF_PALETTE" not in [x["type"] for x in f2])

# ---------------- web page fetching safety
class Resp:
    def __init__(self, code, body=b"", ctype="text/html", loc=None):
        self.status_code, self._b, self.encoding = code, body, "utf-8"
        self.headers = {"content-type": ctype}
        if loc:
            self.headers["location"] = loc

    def iter_content(self, n):
        yield self._b


check("localhost is refused", not FMT.public_host("localhost") and not FMT.public_host("127.0.0.1"))
check("private network is refused", not FMT.public_host("10.0.0.5") and not FMT.public_host("169.254.169.254"))
real_public = FMT.public_host
FMT.public_host = lambda h: h != "internal.test"
try:
    FMT.fetch_page("https://blog.test/a", getter=lambda *a, **k: Resp(302, loc="https://internal.test/admin"))
    check("a redirect to a private address is refused", False)
except ValueError as ex_:
    check("a redirect to a private address is refused", "not a public" in str(ex_))
html_page = (b"<html><head><title>New Zippy launch</title><meta name='description' content='Big news'>"
             b"<script>var x='guaranteed returns';</script></head><body><nav>Menu</nav><h1>Zippy Hatchback</h1>"
             b"<p>Get the Zippy Hatchback VXi for Rs 7.1 lakh with a flat 10% discount.</p></body></html>")
title, desc, text = FMT.fetch_page("https://blog.test/a", getter=lambda *a, **k: Resp(200, html_page))
check("page title and text are read; scripts and menus are skipped", title == "New Zippy launch" and "10% discount" in text
      and "guaranteed returns" not in text and "Menu" not in text)
try:
    FMT.fetch_page("https://blog.test/a.zip", getter=lambda *a, **k: Resp(200, b"PK", "application/zip"))
    check("a non-page link is refused", False)
except ValueError:
    check("a non-page link is refused", True)
try:
    FMT.fetch_page("ftp://x.test/a")
    check("only http and https", False)
except ValueError:
    check("only http and https", True)

# ---------------- end to end through the app
with TestClient(main.app) as a:
    a.post("/signup", data={"company": "Zippy Motors", "name": "Meera", "email": "m@z.test", "password": "a-long-password-1",
                            "industry": "ecommerce"})
    BID = main.q("SELECT id FROM brands WHERE name='Zippy Motors'")[0]["id"]
    a.post("/admin/sample")
    car = [p for p in RE.brand_projects(BID) if p["official_name"] == "Zippy Hatchback"][0]
    r = a.post("/admin/assets", data={"palette": "#1F3A5F, #F2A900", "min_contrast": "4.5", "min_text_pct": "1.2",
                                      "require_logo": "1", "check_palette": "1"},
               files={"logo": ("logo.png", open(logo_path, "rb"), "image/png")})
    reg = json.loads(main.q("SELECT registry FROM brands WHERE id=?", (BID,))[0]["registry"])
    check("brand assets are saved with the logo", reg["assets"]["require_logo"] and os.path.exists(reg["assets"]["logo"]))
    check("rules page shows the brand assets card", "Brand assets" in a.get("/admin/rules").text)
    bad = a.post("/admin/assets", data={"min_contrast": "abc"})
    check("bad numbers are refused", "must be numbers" in bad.text)
    r = a.post("/admin/team/invite", data={"email": "d@x.test", "kind": "agency", "project_id": [str(car["id"])]})
    tok = re.search(r"/invite/([\w-]+)", r.text).group(1)
    p = TestClient(main.app)
    p.post(f"/invite/{tok}", data={"name": "Ad Agency", "password": "agency-strong-pass"})

    def send(path, name, mime):
        r = p.post("/upload", data={"brand_id": str(BID), "language": "en", "project_id": ""},
                   files={"file": (name, open(path, "rb"), mime)}, follow_redirects=False)
        loc = r.headers.get("location", "")
        if not loc.startswith("/s/"):
            return None, None, r.text
        sid = int(loc.rsplit("/", 1)[-1])
        row = main.q("SELECT * FROM submissions WHERE id=?", (sid,))[0]
        return row["status"], main.rep_of(row), sid

    st, rep, sid = send(poster, "poster.png", "image/png")
    types = [f["type"] for f in rep["findings"]]
    cl = {(c["kind"], c["verdict"]) for c in rep["claims"]}
    check("a poster is reviewed: price and discount verified from the image", ("price", "VERIFIED") in cl and ("custom", "VERIFIED") in cl, f"{cl}")
    check("poster gets accessibility findings, not a speech warning", "LOW_CONTRAST" in types and "NO_SPEECH_DETECTED" not in types, str(types))
    check("poster with low contrast needs a person, not auto-approval", st == "SUPERVISION", st)
    pg = p.get(f"/s/{sid}").text
    check("report says 'In the image' and shows the asset section", "In the image" in pg and "Brand assets and accessibility" in pg)
    wrong = os.path.join(D, "wrong.png")
    make_poster(wrong, logo, price="Rs 6.5 lakh", low_contrast=False)
    st, rep, _ = send(wrong, "wrong.png", "image/png")
    check("a wrong price on a poster is a violation", st == "REJECTED" and ("price", "VIOLATION") in {(c["kind"], c["verdict"]) for c in rep["claims"]},
          f"{st} {[(c['kind'], c['verdict'], c['heard']) for c in rep['claims']]}")

    import docx
    doc = docx.Document()
    doc.add_heading("Zippy Hatchback festive offer", 1)
    doc.add_paragraph("Book the Zippy Hatchback VXi at Rs 7.1 lakh with a 10% discount. 100% safe investment.")
    dpath = os.path.join(D, "offer.docx")
    doc.save(dpath)
    st, rep, sid = send(dpath, "offer.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    types = [f["type"] for f in rep["findings"]]
    check("a Word document is reviewed: wrong discount and banned phrase caught", st == "REJECTED" and "CUSTOM_CHECK" in types
          and "BANNED_PHRASE" in types, f"{st} {types}")
    check("documents say 'In the document'", "In the document" in p.get(f"/s/{sid}").text)
    tpath = os.path.join(D, "blank.txt")
    open(tpath, "w").write("   ")
    st, rep, msg = send(tpath, "blank.txt", "text/plain")
    check("an empty document is explained, not crashed", st is None and "No text" in msg)
    zpath = os.path.join(D, "x.exe")
    open(zpath, "wb").write(b"MZ")
    st, rep, msg = send(zpath, "x.exe", "application/octet-stream")
    check("unsupported files are refused with the accepted list", st is None and "pdf, docx" in msg)

    FMT_fetch = FMT.fetch_page
    FMT.fetch_page = lambda url, getter=None: FMT_fetch(url, getter=lambda *a_, **k_: Resp(200, html_page))
    r = p.post("/upload_page", data={"brand_id": str(BID), "url": "https://blog.test/zippy"}, follow_redirects=False)
    sid = int(r.headers["location"].rsplit("/", 1)[-1])
    row = main.q("SELECT * FROM submissions WHERE id=?", (sid,))[0]
    rep = main.rep_of(row)
    check("a blog page is reviewed and the inflated discount is caught", row["status"] == "REJECTED"
          and "CUSTOM_CHECK" in [f["type"] for f in rep["findings"]] and row["source_url"] == "https://blog.test/zippy", row["status"])
    FMT.public_host = real_public
    r = p.post("/upload_page", data={"brand_id": str(BID), "url": "http://127.0.0.1:8000/admin"})
    check("the app refuses to fetch its own server", "not a public web page" in r.text)
    check("submit page offers posters, documents and web pages", ".pdf,.docx" in p.get("/submit").text and "/upload_page" in p.get("/submit").text)

print("\n%d failed" % len(fails) if fails else "\nALL PASSED")
sys.exit(1 if fails else 0)
