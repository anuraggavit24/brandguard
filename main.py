#!/usr/bin/env python3
"""BrandGuard web app (Round 2) - link ingestion, caption check, stricter auto-approval."""

import hashlib
import io
import html
import json
import math
import os
import re
import secrets
import sqlite3
import subprocess
import time
import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import quote, urlparse

import requests
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response
from starlette.middleware.sessions import SessionMiddleware

import accounts as ACC
import jobs as JOBS
import learning as LRN
import ai_review as AIR
import generic as GEN
import importer as IMP
import formats as FMT
import assets as AST
import mailer
import packs
import rewrite as RW
import ocr as OCR
import realestate as RE
import ui
from engine import (LAKH, UNITS, classify_price, extract_prices, fmt_money, fmt_time,
                    norm, review)

# short spelling seen in captions and titles, for example "Rs 90L"
UNITS.setdefault("l", LAKH)

DATA_DIR = Path(os.environ.get("DATA_DIR", "data"))
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = DATA_DIR / "brandguard.db"
GROQ_KEY = os.environ.get("GROQ_API_KEY", "")
MODE = os.environ.get("TRANSCRIBE_MODE", "groq")  # groq | local
LINKS_ENABLED = os.environ.get("LINKS_ENABLED", "1") == "1"
MAX_MB = 60
MAX_SECONDS = 180
VOICE_MIN_AVG = float(os.environ.get("VOICE_MIN_AVG", "0.6"))
VOICE_MAX_LOW = float(os.environ.get("VOICE_MAX_LOW", "0.4"))
THRESHOLD = 0.65
AV_EXT = {".mp4", ".mov", ".mkv", ".webm", ".m4a", ".mp3", ".wav", ".aac", ".ogg"}
ALLOWED = AV_EXT | FMT.IMAGE_EXT | FMT.DOC_EXT
LINK_HOSTS = ("youtube.com", "youtu.be", "instagram.com", "facebook.com", "fb.watch")
DEMO = {"admin": 1, "partner": 2, "freelancer": 3}
ROLE_LABEL = {"admin": "Builder / Brand Manager", "submitter": "Creator / Partner"}

# comparative and urgency claims that need supporting evidence
UNSUPPORTED = [
    "distress", "below market", "market se kam", "lowest price", "best price", "cheapest",
    "last few", "only few", "few units left", "limited units", "limited period", "hurry",
    "से कम प्राइस", "से कम कीमत", "मार्केट से कम", "सबसे कम", "सबसे सस्ता", "सिर्फ आज", "आखिरी मौका",
]

app = FastAPI(title="BrandGuard")
def _secret_key():
    # never fall back to a guessable key: use SECRET_KEY, else a random key kept in the data folder
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    f = DATA_DIR / ".secret_key"
    if not f.exists():
        f.write_text(secrets.token_hex(32), encoding="utf-8")
    return f.read_text(encoding="utf-8").strip()


app.add_middleware(SessionMiddleware, secret_key=_secret_key(), max_age=3 * 86400, same_site="lax")

e = html.escape


# ------------------------------------------------------------------ database
def q(sql, args=()):
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    try:
        return con.execute(sql, args).fetchall()
    finally:
        con.close()


def ex(sql, args=()):
    con = sqlite3.connect(DB_PATH)
    try:
        cur = con.execute(sql, args)
        con.commit()
        return cur.lastrowid
    finally:
        con.close()


REG1 = {
    "brand": "Demo Developer", "project": "Demo Heights", "rera": "P00000000000",
    "configs": [{"name": "2 BHK", "price": 13000000}, {"name": "3 BHK", "price": 19000000}],
    "possession": "December 2028",
    "banned_phrases": ["guaranteed returns", "best in city", "price will double",
                       "100% safe", "no hidden charges"],
    "confidential": ["dealer commission", "internal discount", "unreleased project", "margin slab"],
    "allowed_contacts": ["9999999999"],
    "glossary": ["Demo Heights", "RERA", "BHK", "possession", "crore", "lakh",
                 "दीड कोटी", "सव्वा कोटी", "साडे तीन कोटी", "डेढ़ करोड़", "कोटी", "ताबा"],
}
REG2 = {
    "brand": "Sample Builders", "project": "Sample Residency", "rera": "P11111111111",
    "configs": [{"name": "1 BHK", "price": 6500000}, {"name": "2 BHK", "price": 9500000}],
    "possession": "June 2029",
    "banned_phrases": ["guaranteed returns", "price will double", "only today"],
    "confidential": ["dealer commission", "internal discount"],
    "allowed_contacts": ["8888888888"],
    "glossary": ["Sample Residency", "RERA", "BHK", "crore", "lakh", "कोटी", "लाख"],
}


def init_db():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH)
    con.executescript("""
    CREATE TABLE IF NOT EXISTS brands (id INTEGER PRIMARY KEY, name TEXT, registry TEXT);
    CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY, name TEXT, email TEXT, role TEXT, brand_id INTEGER);
    CREATE TABLE IF NOT EXISTS memberships (user_id INTEGER, brand_id INTEGER);
    CREATE TABLE IF NOT EXISTS submissions (
        id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, brand_id INTEGER, filename TEXT,
        sha256 TEXT, language TEXT, status TEXT, verdict TEXT, report TEXT, created_at TEXT,
        approval_id TEXT, decision_note TEXT, dispute_note TEXT);
    CREATE TABLE IF NOT EXISTS audit (
        id INTEGER PRIMARY KEY AUTOINCREMENT, sub_id INTEGER, at TEXT, actor TEXT, action TEXT, note TEXT);
    """)
    try:
        con.execute("ALTER TABLE submissions ADD COLUMN source_url TEXT")
    except sqlite3.OperationalError:
        pass
    try:
        con.execute("ALTER TABLE audit ADD COLUMN brand_id INTEGER")
    except sqlite3.OperationalError:
        pass
    if con.execute("SELECT COUNT(*) FROM brands").fetchone()[0] == 0:
        con.execute("INSERT INTO brands VALUES (1, ?, ?)", ("Demo Developer", json.dumps(REG1, ensure_ascii=False)))
        con.execute("INSERT INTO brands VALUES (2, ?, ?)", ("Sample Builders", json.dumps(REG2, ensure_ascii=False)))
        con.execute("INSERT INTO users VALUES (1, 'Builder Admin', 'admin@demo.test', 'admin', 1)")
        con.execute("INSERT INTO users VALUES (2, 'Rahul Realty (Channel Partner)', 'partner@demo.test', 'submitter', NULL)")
        con.execute("INSERT INTO users VALUES (3, 'Priya Freelance (Freelancer)', 'freelancer@demo.test', 'submitter', NULL)")
        con.execute("INSERT INTO memberships VALUES (2, 1)")
        con.execute("INSERT INTO memberships VALUES (3, 1)")
        con.execute("INSERT INTO memberships VALUES (3, 2)")
    con.execute("UPDATE users SET name='Priya Freelance (Freelancer)', email='freelancer@demo.test' "
                "WHERE id=3 AND name LIKE '%Fvideoanc%'")
    # the two demo companies start with the general and real estate starter rules, like a real signup would
    import packs as _packs
    for bid, bname in con.execute("SELECT id, name FROM brands WHERE id IN (1, 2)").fetchall():
        reg_ = json.loads(con.execute("SELECT registry FROM brands WHERE id=?", (bid,)).fetchone()[0] or "{}")
        if bname in ("Demo Developer", "Sample Builders") and reg_.get("packs_seeded") != 2:
            _packs.apply(reg_, ["real_estate"])
            reg_["packs_seeded"] = 2
            con.execute("UPDATE brands SET registry=? WHERE id=?", (json.dumps(reg_, ensure_ascii=False), bid))
    con.commit()
    con.close()


init_db()
RE.setup(DB_PATH)
RE.seed_if_empty()
LRN.setup(DB_PATH)
AIR.setup(DB_PATH)
GEN.setup(DB_PATH)
IMP.setup(DB_PATH)
ACC.setup(DB_PATH)
JOBS.setup(DB_PATH)


def log(sid, actor, action, note="", brand_id=None):
    ex("INSERT INTO audit (sub_id, at, actor, action, note, brand_id) VALUES (?,?,?,?,?,?)",
       (sid, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), actor, action, note, brand_id))


def me(request):
    uid = request.session.get("uid")
    if not uid:
        return None
    rows = q("SELECT * FROM users WHERE id=?", (uid,))
    return rows[0] if rows else None


def get_sub(sid):
    rows = q("""SELECT s.*, u.name AS uname, b.name AS bname FROM submissions s
                JOIN users u ON u.id=s.user_id JOIN brands b ON b.id=s.brand_id WHERE s.id=?""", (sid,))
    return rows[0] if rows else None


def can_view(u, s):
    return (u["role"] == "admin" and s["brand_id"] == u["brand_id"]) or s["user_id"] == u["id"]


def new_approval_id():
    return f"BG-{datetime.now().year}-{secrets.token_hex(3).upper()}"


# ---------------------------------------------------------------------- pages
LANG_LABEL = {"mr": "Marathi", "hi": "Hindi", "en": "English", "auto": "Auto-detect"}
QUEUE_STATUSES = ("SUPERVISION", "DISPUTED")
APPROVED_STATUSES = ("APPROVED", "AUTO_APPROVED")
REJECTED_STATUSES = ("REJECTED", "BLOCKED", "REJECTED_FINAL")
PAGE_SIZE = 15

SUB_SQL = """SELECT s.*, u.name AS uname, b.name AS bname FROM submissions s
             JOIN users u ON u.id=s.user_id JOIN brands b ON b.id=s.brand_id """


def user_brands(u):
    return q("SELECT b.* FROM brands b JOIN memberships m ON m.brand_id=b.id WHERE m.user_id=? ORDER BY b.id",
             (u["id"],))


def current_company(request, u):
    """The company a submitter is working in. Always one of their own memberships."""
    ids = [b["id"] for b in user_brands(u)]
    if not ids:
        return None
    cid = request.session.get("company") if request is not None else None
    return cid if cid in ids else ids[0]


def page(title, body, user=None, active="", crumbs=None, request=None):
    if not user:
        return HTMLResponse(ui.landing(body))
    flash = request.session.pop("flash", "") if request is not None else ""
    company, queue_n = None, 0
    if user["role"] == "submitter":
        company = (current_company(request, user), [(b["id"], b["name"]) for b in user_brands(user)])
    else:
        queue_n = q("SELECT COUNT(*) AS c FROM submissions WHERE brand_id=? AND status IN ('SUPERVISION','DISPUTED')",
                    (user["brand_id"],))[0]["c"]
    demo = [] if user["pw_hash"] else [(k, lbl, DEMO[k] == user["id"]) for k, lbl in
                                       (("admin", "Builder"), ("partner", "Channel Partner"), ("freelancer", "Freelancer"))]
    general = False
    if user["role"] == "admin" and user["brand_id"]:
        try:
            general = load_reg(user).get("domain", "property") != "property"
        except Exception:
            general = False
    return HTMLResponse(ui.shell(title, body, user, ROLE_LABEL[user["role"]], active, crumbs,
                                 queue_n, company, flash, demo, general))


def err_page(u, msg, request=None):
    back = "/submit" if u["role"] == "submitter" else "/"
    body = (ui.page_header("Could not complete this request")
            + ui.alert(e(msg), "bad", "Something went wrong.")
            + f'<p style="margin-top:20px"><a class="btn" href="{back}">Go back</a></p>')
    return page("Could not complete", body, u, "", [("Error", None)], request)


def rep_of(row):
    try:
        return json.loads(row["report"] or "{}")
    except (TypeError, ValueError):
        return {}


def risk_badge(rep):
    order = ["LOW", "MEDIUM", "HIGH", "CRITICAL"]
    sev = [f.get("severity", "LOW") for f in rep.get("findings", [])]
    if not sev:
        return ui.badge("No findings", "ok", "check-circle")
    top = max(sev, key=lambda s: order.index(s) if s in order else 0)
    tone = {"CRITICAL": "bad", "HIGH": "bad", "MEDIUM": "warn"}.get(top, "neutral")
    return ui.badge(f"{top.title()} risk", tone, "alert" if tone != "neutral" else "info")


def sub_table(rows, show_user=True, label="Submissions", review_label="Review"):
    if not rows:
        return ui.empty("Nothing here yet", "Submissions will appear here as soon as they are made.")
    heads = ["Submission"] + (["Creator"] if show_user else []) + ["Company", "Submitted", "Findings", "Status", "Action"]
    out = []
    for r in rows:
        rep = rep_of(r)
        link = f' {ui.icon("link", 14)}' if r["source_url"] else ""
        fname = e(r["filename"] or "Untitled")
        cells = [f'<a href="/s/{r["id"]}"><b>#{r["id"]}</b></a>{link}<span class="trunc" title="{fname}">{fname}</span>']
        if show_user:
            cells.append(e(r["uname"]))
        n = f'{rep.get("flags", 0)} violation(s), {rep.get("supervision", 0)} to check'
        cells += [e(r["bname"]), f'<span class="muted">{e(r["created_at"])}</span>',
                  f'{risk_badge(rep)}<br><span class="muted">{n}</span>', ui.chip(r["status"]),
                  f'<a class="btn sm {"" if r["status"] in QUEUE_STATUSES and show_user else "secondary"}" '
                  f'href="/s/{r["id"]}" aria-label="{review_label if (r["status"] in QUEUE_STATUSES and show_user) else "View"} submission {r["id"]}">'
                  f'{review_label if (r["status"] in QUEUE_STATUSES and show_user) else "View"}</a>']
        out.append(cells)
    return ui.table(heads, out, label)


# --------------------------------------------------------------------- landing
@app.get("/health")
def health():
    return {"ok": True}


@app.get("/ai-status")
def ai_status(request: Request):
    """Builder-only check that the text AI works: models tried, and the reason if it fails."""
    import llm
    if not admin_only(request):
        return RedirectResponse("/", 303)
    out = {"key_set": llm.configured(), "models_in_order": llm.models(), "test": None, "reason": ""}
    if llm.configured():
        try:
            got = llm.chat_json('Reply with JSON only: {"ok": true}', "ping", timeout=20)
            out["test"] = "ok" if got.get("ok") is True else f"unexpected reply: {str(got)[:80]}"
        except Exception as ex_:
            out["test"], out["reason"] = "failed", str(ex_)
    return out


@app.get("/ocr-status")
def ocr_status(request: Request):
    if not admin_only(request):
        return JSONResponse({"error": "builder only"}, 403)
    import shutil
    return {"ocr_available": OCR.available(), "tesseract": OCR.TESSERACT,
            "ffmpeg": shutil.which("ffmpeg"), "ffprobe": shutil.which("ffprobe"),
            "fps": OCR.OCR_FPS, "budget": OCR.OCR_BUDGET}


@app.get("/ocr-test")
def ocr_test(request: Request, name: str = "approval.mp4"):
    if not admin_only(request):
        return JSONResponse({"error": "builder only"}, 403)
    import time, traceback
    p = Path("demo_audio") / Path(name).name
    if not p.exists():
        return {"error": "file not found", "looked_for": str(p)}
    out = {"file": str(p), "has_video": OCR.has_video(p)}
    t = time.time()
    try:
        with __import__("tempfile").TemporaryDirectory() as d:
            fr = OCR._frames(p, d, 30)
            out["frames_extracted"] = len(fr)
            out["first_frame_text"] = OCR._ocr_one(fr[0][1]) if fr else None
    except Exception:
        out["exception"] = traceback.format_exc()[-1500:]
    out["secs"] = round(time.time() - t, 1)
    return out


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    u = me(request)
    if not u:
        def role(key, ic, title, text, btn):
            return (f'<div class="card role-card"><span class="ic-lg">{ui.icon(ic, 22)}</span><h2>{title}</h2>'
                    f'<p class="muted">{text}</p><a class="btn" href="/enter/{key}">{btn}{ui.icon("arrow-right", 16)}</a></div>')
        body = f"""
        <div style="margin-bottom:28px"><h1>Review videos before they go live</h1>
        <p class="page-desc">BrandGuard listens to a marketing video in Marathi, Hindi or English, checks every price
        and claim against the brand's approved facts, and sends anything uncertain to a person. Nothing is approved
        without a recorded decision.</p></div>
        <div class="cols-2" style="margin-bottom:28px">
          <div class="card role-card"><span class="ic-lg">{ui.icon("building", 22)}</span><h2>Set up your company</h2>
            <p class="muted">Create a company account, add your brand rules and projects, then invite channel partners, freelancers and agencies.</p>
            <a class="btn" href="/signup">Create company account{ui.icon("arrow-right", 16)}</a></div>
          <div class="card role-card"><span class="ic-lg">{ui.icon("users", 22)}</span><h2>Sign in</h2>
            <p class="muted">Already have an account, or an invite link from a company? Sign in or open the link you were sent.</p>
            <a class="btn secondary" href="/login">Sign in{ui.icon("arrow-right", 16)}</a></div>
        </div>
        <p class="sec-title">Or try the demo with a ready-made role</p>
        <div class="cols-3">
          {role("admin", "shield", "Builder / Brand Manager",
                "Set brand rules and approved facts, review flagged and disputed items, then approve or reject.",
                "Enter as Builder")}
          {role("partner", "upload", "Channel Partner",
                "Submit a video or paste a link for one company and see the result.", "Enter as Partner")}
          {role("freelancer", "building", "Freelancer",
                "Works with two companies under one login and switches between them.", "Enter as Freelancer")}
        </div>
        <p style="margin-top:20px"><a class="btn secondary" href="/test-audio">Judges: demo files to download</a></p>
        <p class="muted" style="margin-top:12px">The demo roles are shared sample accounts with fictional data. Real accounts are separate.</p>"""
        return page("Welcome", body)
    return admin_dash(request, u) if u["role"] == "admin" else submitter_dash(request, u)


@app.get("/enter/{key}")
def enter(request: Request, key: str):
    uid = DEMO.get("freelancer" if key == "fvideoancer" else key)
    if uid:
        request.session["uid"] = uid
        request.session.pop("company", None)
    return RedirectResponse("/", 303)


@app.get("/company")
def switch_company(request: Request, id: int = 0):
    u = me(request)
    if u and u["role"] == "submitter" and id in [b["id"] for b in user_brands(u)]:
        request.session["company"] = id
    ref = urlparse(request.headers.get("referer", "")).path
    return RedirectResponse(ref if ref.startswith("/") and not ref.startswith("//") and not ref.startswith("/s/") else "/", 303)


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/", 303)


# ------------------------------------------------------------ submitter views
def submitter_dash(request, u):
    cid = current_company(request, u)
    brand = q("SELECT name FROM brands WHERE id=?", (cid,))
    cname = brand[0]["name"] if brand else "your company"
    mine = q(SUB_SQL + "WHERE s.user_id=? AND s.brand_id=? ORDER BY s.id DESC", (u["id"], cid))
    n = Counter(r["status"] for r in mine)
    approved = sum(n[s] for s in APPROVED_STATUSES)
    rejected = sum(n[s] for s in REJECTED_STATUSES)
    in_review = n["SUPERVISION"] + n["DISPUTED"] + n["NEEDS_CONFIRMATION"]
    first = u["name"].split(" (")[0].split()[0]
    multi = len(user_brands(u)) > 1

    welcome = f"""<section class="card welcome" aria-labelledby="wl"><div class="card-head"><h2 id="wl">Welcome, {e(first)}</h2>
      <button type="button" class="btn secondary sm" data-dismiss="creator">Dismiss</button></div>
      <p>Submit a video and BrandGuard checks it against {e(cname)}'s approved prices and rules before it is published.
      Uncertain cases are sent to a person, so you always get a clear answer.</p>
      <ol class="steps"><li>Submit content</li><li>Analyze content</li><li>Review findings</li><li>Resolve issues</li><li>Receive a decision</li></ol></section>"""
    attn = ""
    if n["REJECTED"] + n["BLOCKED"]:
        attn = ui.alert(f'{n["REJECTED"] + n["BLOCKED"]} submission(s) were rejected. Open them to see what to fix. '
                        'You can resubmit a corrected file, or dispute the decision if you think it is wrong.',
                        "warn", "Needs your attention.")
    scope = (f'<p class="muted">Showing activity for <b>{e(cname)}</b>. Use the company switcher at the top to change.</p>'
             if multi else "")
    body = (ui.page_header("My Dashboard", f"Your video submissions for {e(cname)}.",
                           f'<a class="btn" href="/submit">{ui.icon("upload", 16)}Submit a Video</a>')
            + '<div class="stack">' + welcome + attn + scope
            + '<div class="stats">'
            + ui.stat("Submissions", len(mine), "", "/reviews")
            + ui.stat("In review", in_review, "Waiting for a person", "/reviews?status=supervision", "warn")
            + ui.stat("Approved", approved, "", "/reviews?status=approved", "ok")
            + ui.stat("Rejected", rejected, "Fix and resubmit, or dispute", "/reviews?status=rejected", "bad")
            + '</div>'
            + ui.card("Recent submissions",
                      sub_table(mine[:5], show_user=False, label="Recent submissions") if mine else
                      ui.empty("No submissions yet", f"Submit your first video for {e(cname)} to get started.",
                               '<a class="btn" href="/submit">Submit a Video</a>'),
                      actions='<a href="/reviews">View all</a>' if len(mine) > 5 else "")
            + '</div>')
    return page("My Dashboard", body, u, "overview", [], request)


@app.get("/submit", response_class=HTMLResponse)
def submit_form(request: Request):
    u = me(request)
    if not u or u["role"] != "submitter":
        return RedirectResponse("/", 303)
    brands = user_brands(u)
    cid = current_company(request, u)
    opts = "".join(f'<option value="{b["id"]}"{" selected" if b["id"] == cid else ""}>{e(b["name"])}</option>'
                   for b in brands)
    langs = "".join(f'<option value="{k}">{v}</option>' for k, v in
                    (("mr", "Marathi"), ("hi", "Hindi"), ("en", "English"), ("auto", "Auto-detect")))
    projs = RE.authorized_projects(u["id"], cid)
    popts = '<option value="">Detect automatically (recommended)</option>' + "".join(
        f'<option value="{p["id"]}">{e(p["official_name"])}{" - " + e(p["city"]) if p["city"] else ""}</option>' for p in projs)
    company = (f'<label for="brand_id">Company</label><select id="brand_id" name="brand_id">{opts}</select>'
               f'<label for="project_id">Project <span class="muted">(optional)</span></label>'
               f'<select id="project_id" name="project_id">{popts}</select>'
               f'<p class="hint">BrandGuard finds the project from what is said and written. You only need to choose one if '
               f'you want to be explicit. You can only pick projects you have been given access to'
               f'{"" if projs else ". You have none for this company yet, so ask the company to grant access"}.</p>')
    lang = (f'<label for="language">Spoken language</label><select id="language" name="language">{langs}</select>'
            '<p class="hint">Choose the language spoken in the video. If unsure, pick Auto-detect.</p>')
    exts = ",".join(sorted(ALLOWED))
    file_form = f"""
    <form method="post" action="/upload" enctype="multipart/form-data" data-lock>
      <div class="dropzone" id="dropzone">{ui.icon("upload", 32)}
        <p><b>Drag a video, audio, poster, image or document here</b></p>
        <p class="muted">Accepted: {", ".join(x.lstrip(".") for x in sorted(ALLOWED))}. Up to {MAX_MB} MB.
        The first {MAX_SECONDS // 60} minutes are reviewed.</p>
        <label class="btn secondary" for="file">Browse files</label>
        <input id="file" type="file" name="file" required data-max="{MAX_MB * 1024 * 1024}" data-ext="{exts}"
               accept="video/*,audio/*,image/png,image/jpeg,image/webp,.pdf,.docx,.txt" aria-describedby="file-err"></div>
      <div class="picked" id="picked">{ui.icon("file", 18)}<b id="picked-name"></b><span class="sp"></span>
        <button type="button" class="btn secondary sm" id="remove-file">Remove</button></div>
      <p class="req" id="file-err" role="alert" hidden></p>
      {company}{lang}
      <p data-busy hidden class="muted" role="status">Reviewing. On this free server it can take a few minutes. Please keep this page open.</p>
      <button id="submit-file" class="btn" type="submit" data-busy-label="Reviewing...">Submit for Compliance Review</button>
    </form>"""
    if LINKS_ENABLED:
        link_form = f"""
    {ui.alert("Link submission is supported on a laptop or desktop. On a phone, download the video and upload the file instead.", "info")}
    <form method="post" action="/upload_link" data-lock style="margin-top:16px">
      <label for="url">Link to a public video</label>
      <input id="url" type="url" name="url" required placeholder="https://www.youtube.com/shorts/..." aria-describedby="url-h">
      <p class="hint" id="url-h">YouTube Shorts, Instagram videos and Facebook videos. Both the audio and the post's title and
      description are reviewed. Private posts cannot be fetched.</p>
      {company}{lang}
      <p data-busy hidden class="muted" role="status">Fetching and reviewing. This can take up to a minute. Please keep this page open.</p>
      <button class="btn" type="submit" data-busy-label="Fetching...">Fetch and Submit for Compliance Review</button>
    </form>"""
    else:
        link_form = ui.alert(
            "Link checking is switched off on this hosted demo, because YouTube, Instagram and Facebook block automated "
            "downloads from shared cloud servers. It works in the local build on a laptop or desktop. To check a posted video "
            "here, download it and use <b>Upload a file</b>.", "warn", "Links are unavailable here.")
    page_form = f"""
    <form method="post" action="/upload_page" data-lock style="margin-top:22px">
      <label for="purl">Link to a blog post or web page</label>
      <input id="purl" type="url" name="url" required placeholder="https://www.example.com/blog/new-launch">
      <p class="hint">The page's text is read and checked like a document: prices, banned claims, product facts and wording.
      Only public pages can be read.</p>
      {company}
      <button class="btn" type="submit" data-busy-label="Reading...">Check this page</button>
    </form>"""
    link_form = link_form + page_form
    body = (ui.page_header("Submit Content for Review",
                           "BrandGuard listens to the video, checks every price and claim against the brand's approved facts "
                           "and shows you the result. Anything uncertain goes to a person for a decision.")
            + ui.alert("This demo runs on a free Render server with very little CPU and memory. To review a video, BrandGuard turns the speech into text and then reads every frame of the video, so it can take a few minutes. The first request after a quiet period is slower still, because the server has to wake up. Please keep this page open until the result appears. On a normal server the same review takes seconds. Short videos (under 15 seconds) are quickest.", "warn", "Please be patient: free server.")
            + '<div class="cols"><div>'
            + ui.card("", f"""
              <div class="method" role="tablist" aria-label="Submission method">
                <button type="button" class="btn" role="tab" id="tab-file" aria-selected="true" aria-controls="panel-file">{ui.icon("upload", 16)}Upload a file</button>
                <button type="button" class="btn" role="tab" id="tab-link" aria-selected="false" aria-controls="panel-link">{ui.icon("link", 16)}Paste a link</button></div>
              <div id="panel-file" role="tabpanel" aria-labelledby="tab-file">{file_form}</div>
              <div id="panel-link" role="tabpanel" aria-labelledby="tab-link" hidden>{link_form}</div>""")
            + '</div><div>'
            + ui.card("What happens next", """<ol class="muted" style="padding-left:20px;margin:0">
              <li>The speech is transcribed and the on-screen text is read, then both are checked. On the free demo server this can take a few minutes.</li>
              <li>You see one of: auto-approved, rejected, or sent to a person.</li>
              <li>If rejected, fix the video and resubmit, or dispute the finding.</li>
              <li>Nothing is approved just because it uploaded.</li></ol>""")
            + '</div></div>')
    return page("Submit a Video", body, u, "submit", [("Submit a Video", None)], request)


@app.get("/help", response_class=HTMLResponse)
def help_page(request: Request):
    u = me(request)
    if not u:
        return RedirectResponse("/", 303)
    legend = [[ui.chip(k), t] for k, t in (
        ("AUTO_APPROVED", "Every price matched the approved facts and nothing needed a person. An approval ID was issued."),
        ("SUPERVISION", "BrandGuard was not sure (for example, unclear audio or a possible lakh/crore mix-up). A builder will decide."),
        ("REJECTED", "A claim breaks the brand rules, such as a wrong price or a banned phrase. Fix and resubmit, or dispute."),
        ("BLOCKED", "Confidential company information was mentioned. Remove it before resubmitting."),
        ("DISPUTED", "You disagreed with a rejection. A builder will uphold or overturn it."),
        ("APPROVED", "A builder approved this after review."),
        ("REJECTED_FINAL", "A builder upheld the rejection after review."))]
    body = (ui.page_header("Help and Guidelines", "How reviews work and what each status means.")
            + '<div class="stack">'
            + ui.card("What BrandGuard checks", """<ul class="muted" style="margin:0;padding-left:20px">
              <li>Prices spoken or written, compared with the approved price list (including forms like "dedh crore" or "sava crore").</li>
              <li>Banned claims and confidential company terms.</li>
              <li>Phone numbers and emails that are not on the approved contact list.</li>
              <li>Comparative or urgency claims that need evidence, such as "lowest price" or "limited units".</li>
              <li>For linked posts, the title and description as well as the audio.</li>
              <li>Posters, banners and images (text read from the image), documents (PDF, Word, text) and blog pages.</li>
              <li>For images: text contrast, text size, brand colours and the brand logo, when the company has set them.</li>
              <li>The company's own product facts, such as discounts, warranties, fees, rates and dates.</li></ul>""")
            + ui.card("What each status means", ui.table(["Status", "Meaning"], legend, "Status meanings"))
            + ui.card("Not available", """<p class="muted" style="margin:0">Checking a RERA number against the official
              register, identifying fonts, and finding a logo that has been redrawn or heavily recoloured.</p>""")
            + '</div>')
    return page("Help and Guidelines", body, u, "help", [("Help and Guidelines", None)], request)


def save_upload(up: UploadFile):
    ext = Path(up.filename or "").suffix.lower()
    if ext not in ALLOWED:
        raise ValueError("Unsupported file type. Upload a video or audio file (mp4, mov, m4a, wav...), an image "
                         "(png, jpg, webp) or a document (pdf, docx, txt).")
    dst = UPLOAD_DIR / (uuid.uuid4().hex + ext)
    h = hashlib.sha256()
    size = 0
    too_big = False
    with open(dst, "wb") as f:
        while True:
            chunk = up.file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_MB * 1024 * 1024:
                too_big = True
                break
            h.update(chunk)
            f.write(chunk)
    if too_big:
        dst.unlink(missing_ok=True)
        raise ValueError(f"File is larger than {MAX_MB} MB.")
    return dst, h.hexdigest()


def host_ok(url: str) -> bool:
    try:
        p = urlparse(url.strip())
    except Exception:
        return False
    if p.scheme not in ("http", "https") or not p.hostname:
        return False
    h = p.hostname.lower()
    return any(h == d or h.endswith("." + d) for d in LINK_HOSTS)


def fetch_link(url: str):
    """Download one public video from an allowed host. Returns (path, sha256, title, description)."""
    import yt_dlp
    from yt_dlp.utils import DownloadError

    fail = ("Could not fetch this link. The post may be private, removed, or blocked for automated "
            "access. Please download the video and upload the file instead.")
    stem = uuid.uuid4().hex
    opts = {
        "outtmpl": str(UPLOAD_DIR / (stem + ".%(ext)s")),
        "format": "bv*[height<=480]+ba[format_note*=original]/bv*[height<=480]+ba/b[height<=480]/b",
        "merge_output_format": "mp4",
        "noplaylist": True,
        "quiet": True,
        "no_warnings": True,
        "socket_timeout": 20,
        "retries": 1,
        "max_filesize": MAX_MB * 1024 * 1024,
        "match_filter": yt_dlp.utils.match_filter_func("duration <=? 300"),
    }
    try:
        with yt_dlp.YoutubeDL(opts) as ydl:
            info = ydl.extract_info(url.strip(), download=True)
    except DownloadError:
        raise RuntimeError(fail)
    except Exception:
        raise RuntimeError(fail)
    if not info:
        raise RuntimeError("This video is longer than 5 minutes or larger than the size limit, "
                           "so it was not fetched. Please upload a shorter file.")
    files = [f for f in UPLOAD_DIR.glob(stem + ".*") if f.suffix.lower() in AV_EXT]
    if not files:
        raise RuntimeError(fail)
    path = files[0]
    if path.stat().st_size > MAX_MB * 1024 * 1024:
        path.unlink(missing_ok=True)
        raise RuntimeError(f"The video is larger than {MAX_MB} MB.")
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    title = (info.get("title") or "Linked video").strip()[:200]
    desc = (info.get("description") or "").strip()[:1500]
    return path, h.hexdigest(), title, desc


def extract_audio(src: Path, dst: Path):
    cmd = ["ffmpeg", "-y", "-i", str(src), "-vn", "-ac", "1", "-ar", "16000",
           "-t", str(MAX_SECONDS), str(dst)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    if r.returncode != 0 or not dst.exists() or dst.stat().st_size < 1000:
        raise RuntimeError("Could not read audio from this file. Check that it has sound and try again.")


def transcribe_groq(wav: Path, lang: str, glossary):
    if not GROQ_KEY:
        raise RuntimeError("The speech service key is not configured on the server.")
    data = {"model": "whisper-large-v3", "response_format": "verbose_json", "temperature": "0"}
    if glossary:
        data["prompt"] = ", ".join(glossary)[:400]
    if lang in ("mr", "hi", "en"):
        data["language"] = lang
    with open(wav, "rb") as f:
        r = requests.post("https://api.groq.com/openai/v1/audio/transcriptions",
                          headers={"Authorization": f"Bearer {GROQ_KEY}"},
                          files={"file": ("audio.wav", f)}, data=data, timeout=120)
    if r.status_code != 200:
        raise RuntimeError(f"The speech service returned an error ({r.status_code}). Please try again.")
    j = r.json()
    segs = [{"start": round(s["start"], 2), "end": round(s["end"], 2), "text": s["text"].strip(),
             "conf": round(math.exp(s.get("avg_logprob", -1.0)), 3)} for s in j.get("segments", [])]
    if not segs and (j.get("text") or "").strip():
        segs = [{"start": 0, "end": 0, "text": j["text"].strip(), "conf": 0.5}]
    return segs


def transcribe(wav: Path, lang: str, glossary):
    if MODE == "local":
        from engine import transcribe as local_transcribe
        segs, _ = local_transcribe(wav, os.environ.get("LOCAL_MODEL", "medium"),
                                   lang if lang in ("mr", "hi", "en") else None, glossary, "auto", "auto")
        return segs
    return transcribe_groq(wav, lang, glossary)


def record_lessons(sid, rep, brand_id):
    ids = [c["lesson_id"] for c in rep.get("cleared", [])]
    LRN.record_uses(sid, ids)
    for c in rep.get("cleared", []):
        log(sid, "BrandGuard engine", "LESSON_APPLIED",
            f'"{c["evidence"]}" not flagged: lesson #{c["lesson_id"]} from submission #{c["source_sub"]}.', brand_id)


def apply_lessons(rep, brand_id):
    """Clear wording flags that a builder already decided are harmless in this exact context (precedent)."""
    inp = rep.get("inputs", {})
    rep["cleared"], keep = [], []
    for f in rep["findings"]:
        les = None
        if f["type"] in LRN.LEARNABLE:
            src = f.get("source", "audio")
            if src == "caption":
                cands = [inp.get("caption", "")]
            elif src == "screen":
                cands = [l["text"] for l in inp.get("screen", []) if abs(l["start"] - f["time"]) < 0.01] or [l["text"] for l in inp.get("screen", [])]
            else:
                cands = [x["text"] for x in inp.get("segments", []) if abs(x["start"] - f["time"]) < 0.01] or [x["text"] for x in inp.get("segments", [])]
            trig = LRN.tokens(f.get("evidence", ""))
            for text in cands:
                if LRN._occurrences(LRN.tokens(text), trig):
                    les = LRN.match(brand_id, f["type"], f.get("evidence", ""), text)
                    break
        if les:
            rep["cleared"].append({"lesson_id": les["id"], "type": f["type"], "evidence": f["evidence"], "time": f["time"],
                                   "summary": les["summary"], "source_sub": les["source_sub"]})
        else:
            keep.append(f)
    rep["findings"] = keep


def finalize(rep):
    """Recount findings and set the verdict from the final list."""
    live = [f for f in rep["findings"] if not f.get("dismissed")]
    fl = [f for f in live if f["status"] == "FLAG"]
    su = [f for f in live if f["status"] == "SUPERVISION"]
    rep["flags"], rep["supervision"] = len(fl), len(su)
    if any(f["severity"] == "CRITICAL" for f in live):
        rep["verdict"] = "BLOCKED - confidential data"
    elif fl:
        rep["verdict"] = "REJECT - fix and resubmit"
    elif su:
        rep["verdict"] = "SUPERVISION - manager check needed"
    else:
        rep["verdict"] = "PASS - auto-clear candidate"


def make_finding(status, severity, ftype, t, evidence, why, fix, conf=None, source="audio"):
    return {"status": status, "severity": severity, "type": ftype, "time": t, "evidence": evidence,
            "explanation": why, "suggestion": fix, "confidence": conf, "source": source}


def redact(text, reg):
    out = text or ""
    for w in reg.get("confidential", []):
        if w.strip():
            out = re.sub(re.escape(norm(w)), "[REDACTED]", out, flags=re.I)
    return out


def glossary_for(reg, brand_id, user_id):
    terms = list(reg.get("glossary", []))
    for p in RE.authorized_projects(user_id, brand_id):
        terms += RE.names_of(p)
    seen, out = set(), []
    for t in terms:
        if t and t.lower() not in seen:
            seen.add(t.lower())
            out.append(t)
    return out


def transcribe_file(src: Path, lang: str, glossary):
    wav = src.parent / (src.stem + "_16k.wav")
    try:
        extract_audio(src, wav)
        return transcribe(wav, lang, glossary)
    finally:
        wav.unlink(missing_ok=True)

        

def _close(a, b):
    """True if two words differ by at most one letter (two for long words): OCR and speech-to-text slips."""
    if a == b:
        return True
    if len(b) < 4 or abs(len(a) - len(b)) > 2:
        return False
    limit = 2 if len(b) >= 8 else 1
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i] + [0] * len(b)
        for j, cb in enumerate(b, 1):
            cur[j] = min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb))
        prev = cur
        if min(prev) > limit:
            return False
    return prev[-1] <= limit


def confidential_figures(text, confidential):
    """A confidential subject stated with a number ("2% commission", "कमिशन 2 टक्के") even when the exact confidential
    phrase ("dealer commission") is not used. The subject is the longest word of each confidential phrase, matched by
    sound so Devanagari spellings count. Returns [(confidential phrase, heard words)]."""
    import phon
    toks = phon.tokens(text)
    out = []
    for ph in confidential:
        words = [w for w in re.findall(r"[\w\u0900-\u097F]+", (ph or "").lower()) if len(phon.skel(w)) >= 3]
        if not words:
            continue
        head = max(words, key=lambda w: len(phon.skel(w)))
        hs = phon.skel(head)
        for i, t in enumerate(toks):
            if t != head and phon.skel(t) != hs:
                continue
            near = toks[max(0, i - 4):i + 5]
            if any(re.fullmatch(r"\d+(?:[.,]\d+)?%?", w) for w in near) or "%" in text:
                out.append((ph, " ".join(near)))
                break
    return out


def near_banned(text, banned):
    """Banned phrases (of at least two words) that appear in the text with small spelling slips but not exactly."""
    toks = re.findall(r"[\w%']+", norm(text or "").lower())
    out = []
    low = " ".join(toks)
    for ph in banned:
        pt = re.findall(r"[\w%']+", norm(ph).lower())
        if len(pt) < 2 or not ph.isascii() or (" " + " ".join(pt) + " ") in (" " + low + " "):
            continue
        for i in range(len(toks) - len(pt) + 1):
            window = toks[i:i + len(pt)]
            if all(_close(w, p) for w, p in zip(window, pt)) and sum(w != p for w, p in zip(window, pt)) <= 2:
                out.append(ph)
                break
    return out


def analyze(u, brand_id, segs, caption, reg, hint_pid=None, confirmed=None, is_link=False, screen=None, use_lessons=True,
            kind="video"):
    """Rules plus project-aware claim checks. Never approves anything it could not fully check."""
    rep = review(segs, reg, THRESHOLD, check_prices=False)
    caption_parts = RE._split_caption(caption) if caption.strip() else []
    if caption.strip():
        cap = review([{"start": 0, "end": 0, "text": caption, "conf": 1.0}], reg, THRESHOLD, check_prices=False)
        for f in cap["findings"]:
            f["source"] = "caption"
            rep["findings"].append(f)

    screen_lines = list((screen or {}).get("lines", []))
    if screen_lines:  # banned phrases and confidential terms drawn on screen
        scr = review([{"start": l["start"], "end": l["end"], "text": l["text"], "conf": l["conf"]} for l in screen_lines],
                     reg, THRESHOLD, check_prices=False)
        for f in scr["findings"]:
            f["source"] = "screen"
            rep["findings"].append(f)
    sources = [(s["text"], s["start"], "audio") for s in segs]
    sources += [(l["text"], l["start"], "screen") for l in screen_lines]
    if caption.strip():
        sources.append((caption, 0, "caption"))
    # a banned phrase misread by OCR or speech-to-text ("watch your meney grow") is not an exact match, but a person must see it
    exact = {(f.get("source"), norm(str(f["evidence"])).lower()) for f in rep["findings"] if f["type"] == "BANNED_PHRASE"}
    near_seen = set()
    secret_hit = {(f.get("source")) for f in rep["findings"] if f["type"] == "CONFIDENTIAL_DATA"}
    for text, t, srcname in sources:
        for phrase, heard in confidential_figures(text, reg.get("confidential", [])):
            if srcname in secret_hit or (srcname, "c:" + phrase) in near_seen:
                continue
            near_seen.add((srcname, "c:" + phrase))
            rep["findings"].append(make_finding(
                "SUPERVISION", "HIGH", "CONFIDENTIAL_LIKELY", t, heard[:120],
                f'A figure is given for something the company keeps confidential ("{phrase}").',
                "Remove the figure unless the company has approved sharing it.", None, srcname))
        for phrase in near_banned(text, reg.get("banned_phrases", [])):
            if (srcname, phrase.lower()) in exact or (srcname, phrase.lower()) in near_seen:
                continue
            near_seen.add((srcname, phrase.lower()))
            rep["findings"].append(make_finding(
                "SUPERVISION", "MEDIUM", "BANNED_PHRASE_LIKELY", t, text[:120],
                f'This looks like the banned phrase "{phrase}" with a letter or two misread by '
                f'{"on-screen text reading" if srcname == "screen" else "speech recognition"}.',
                "A person should check the wording. If it is the banned claim, remove it.", None, srcname))
    seen = set()
    for text, t, srcname in sources:
        low = norm(text).lower()
        for w in UNSUPPORTED:
            if w in low and (w, srcname) not in seen:
                seen.add((w, srcname))
                rep["findings"].append(make_finding(
                    "SUPERVISION", "MEDIUM", "UNSUPPORTED_CLAIM", t, w,
                    f'"{w}" is a comparative or urgency claim that needs supporting evidence.',
                    "Add proof from the approved facts or offers, or remove the claim.", None, srcname))

    authorised = RE.authorized_projects(u["id"], brand_id)
    all_p = RE.brand_projects(brand_id)
    texts = [("audio", i, s["text"]) for i, s in enumerate(segs)] + [("caption", i, c) for i, c in enumerate(caption_parts)]
    texts += [("screen", i, l["text"]) for i, l in enumerate(screen_lines)]
    if confirmed:
        idres = {"state": "CONFIRMED", "project_ids": list(confirmed), "evidence": {}, "out_of_scope": [],
                 "segments": {}, "note": "Project confirmed by the submitter."}
    else:
        idres = RE.identify(texts, authorised, all_p, hint_pid, reg.get("domain", "property"))
    state = idres["state"]
    names = {p["id"]: p["official_name"] for p in all_p}
    rep["identity"] = {"state": state, "project_ids": idres["project_ids"], "note": idres["note"],
                       "evidence": {str(k): v for k, v in idres["evidence"].items()},
                       "names": {str(k): names.get(k, "?") for k in set(idres["project_ids"]) | set(idres["out_of_scope"])}}
    rep["project_ids"] = list(idres["project_ids"])
    rep["inputs"] = {"segments": [{"start": s["start"], "end": s.get("end", s["start"]), "text": redact(s["text"], reg),
                                   "conf": s.get("conf", 1.0)} for s in segs],
                     "caption": redact(caption, reg),
                     "screen": [{"start": l["start"], "end": l["end"], "text": redact(l["text"], reg),
                                 "raw": redact(l.get("raw", l["text"]), reg), "conf": l["conf"]} for l in screen_lines]}
    if screen is not None:
        rep["screen_meta"] = {"frames": screen.get("frames", 0), "total": screen.get("total", 0), "partial": bool(screen.get("partial"))}
    rep["policy"] = RE.POLICY_VERSION
    rep["kind"] = kind
    rep["claims"], rep["facts_used"] = [], []
    rep["needs_confirmation"] = state in ("LIKELY", "AMBIGUOUS", "NOT_FOUND", "CONFLICT") and bool(segs or caption.strip())
    rep["candidates"] = [p["id"] for p in authorised] if state in ("NOT_FOUND",) else list(idres["project_ids"])

    checks = [("Project identification", "done" if not rep["needs_confirmation"] else "pending",
               idres["note"] or "Project identified from the content.")]
    if state == "NOT_AUTHORIZED":
        rep["findings"].append(make_finding(
            "SUPERVISION", "HIGH", "UNAUTHORIZED_PROJECT", 0,
            ", ".join(names.get(p, "?") for p in idres["out_of_scope"]),
            "The content seems to promote a project this account has not been given access to.",
            "Remove that project, or ask the company to grant access."))
    elif state == "SELECTED":
        rep["findings"].append(make_finding(
            "SUPERVISION", "LOW", "IDENTITY_WEAK", 0, names.get(idres["project_ids"][0], ""),
            "The content never names the project, so the project could only be taken from your selection.",
            "A person should confirm this video is about the selected project."))
    if idres["out_of_scope"] and state not in ("NOT_AUTHORIZED",):
        rep["findings"].append(make_finding(
            "SUPERVISION", "HIGH", "UNAUTHORIZED_PROJECT", 0, ", ".join(names.get(p, "?") for p in idres["out_of_scope"]),
            "The content also mentions a project this account has not been given access to.",
            "Remove that project, or ask the company to grant access."))

    if not rep["needs_confirmation"] and idres["project_ids"]:
        by_proj, unassigned = RE.assign_segments(segs, caption_parts, idres, all_p, screen_lines)
        claims, used = RE.validate(by_proj, all_p, THRESHOLD, reg.get("domain", "property"))
        rep["claims"], rep["facts_used"] = claims, used
        for c in claims:
            if c["verdict"] != RE.VERIFIED:
                rep["findings"].append(RE.claim_to_finding(c))
        for src_, t_, conf_, text_ in unassigned:
            if extract_prices(text_) or RE.AREA_RE.search(norm(text_)):
                rep["findings"].append(make_finding(
                    "SUPERVISION", "MEDIUM", "PROJECT_UNCERTAIN", t_, text_[:80],
                    "A price or area was mentioned, but it could not be linked to one project.",
                    "A person should check which project this refers to.", conf_, src_))

    rep["passed"] = sum(1 for c in rep["claims"] if c["verdict"] == RE.VERIFIED)
    speech = bool(segs)
    if kind == "video":
        checks.append(("Spoken claims", "done" if speech else "incomplete",
                       "Speech transcribed and checked." if speech else "No speech could be detected, so spoken claims were not checked."))
    else:
        n_lines = len(screen_lines)
        what = {"image": "Text in the image", "document": "Text in the document"}.get(kind, "Text")
        checks.append((what, "done" if n_lines else "incomplete",
                       f"{n_lines} line(s) of text read and checked." if n_lines else
                       ("No text could be read." if screen is not None else "Text could not be read on this server (OCR unavailable).")))
    if is_link:
        has_cap = bool(caption.strip())
        checks.append(("Post title and description", "done" if has_cap else "incomplete",
                       "Checked." if has_cap else "No title or description was available."))
    checks.append(("Project facts (price, area, bathrooms, possession, RERA, offers, distance)"
                   if reg.get("domain", "property") == "property" else "Product facts (price and your own fact types)",
                   "done" if not rep["needs_confirmation"] and idres["project_ids"] else "pending",
                   "Checked against the approved facts of the identified project."
                   if idres["project_ids"] and not rep["needs_confirmation"] else "Waiting for the project to be confirmed."))
    if kind != "video":
        pass
    elif screen is None:
        checks.append(("On-screen text", "not_available",
                       "Not analysed. This file has no video, or on-screen reading was off or unavailable."))
    elif screen.get("partial"):
        checks.append(("On-screen text", "incomplete",
                       f"Only {screen['frames']} of {screen['total']} frames were read in the time allowed."))
        rep["findings"].append(make_finding(
            "SUPERVISION", "MEDIUM", "SCREEN_PARTIAL", 0, f"{screen['frames']}/{screen['total']} frames",
            "Not every frame could be read in the time allowed, so some on-screen text may not have been checked.",
            "A person should watch the rest of the video."))
    else:
        checks.append(("On-screen text", "done",
                       f"{screen.get('frames', 0)} frames read (2 per second plus every scene change); "
                       f"{len(screen_lines)} distinct text lines found."))
    spoken_p = {p["value"] for s in segs for p in extract_prices(s["text"])}
    shown_p = {p["value"] for l in screen_lines for p in extract_prices(l["text"])}
    if spoken_p and shown_p and not (spoken_p & shown_p):
        rep["findings"].append(make_finding(
            "SUPERVISION", "HIGH", "VOICE_SCREEN_MISMATCH", 0,
            "voice: " + ", ".join(fmt_money(v) for v in sorted(spoken_p)) + " / screen: " + ", ".join(fmt_money(v) for v in sorted(shown_p)),
            "The price spoken in the video is different from the price shown on screen.",
            "Make the voice and the on-screen price the same, and match the approved price."))
    if kind == "video":
        checks.append(("Voice versus screen", "done" if (spoken_p and shown_p) else "incomplete",
                       "Spoken and on-screen prices compared." if (spoken_p and shown_p)
                       else "Needs both a spoken price and an on-screen price to compare."))
    checks.append(("RERA authoritative verification", "not_available",
                   "Not available. A RERA number is only compared with the BrandGuard registry unless an administrator has recorded an official check."))
    rep["checks"] = checks

    if use_lessons:
        apply_lessons(rep, brand_id)
    if kind != "video" and not screen_lines:
        rep["findings"].append(make_finding(
            "SUPERVISION", "MEDIUM", "NO_TEXT_FOUND", 0, "-",
            "No text could be read from this file, so nothing was checked.", "A person should look at the file."))
    elif kind == "video" and not segs:
        rep["findings"].append(make_finding(
            "SUPERVISION", "MEDIUM", "NO_SPEECH_DETECTED", 0, "-",
            "No speech could be detected in this file.", "A human should check the file."))
    elif (not rep["needs_confirmation"] and rep["passed"] == 0
          and not any(f["status"] == "FLAG" for f in rep["findings"])
          and not any(f["type"] == "NO_VERIFIABLE_CLAIM" for f in rep["findings"])):
        rep["findings"].append(make_finding(
            "SUPERVISION", "LOW", "NO_VERIFIABLE_CLAIM", 0, "-",
            "Nothing in this content could be matched against approved facts, so nothing was verified.",
            "A person should confirm that no price or fact needs checking."))
    # reliability across voices: very unclear speech is never auto-approved
    if segs:
        confs = [float(x.get("conf", 1.0)) for x in segs]
        avg = sum(confs) / len(confs)
        low = sum(1 for c_ in confs if c_ < 0.5)
        rep["voice"] = {"avg": round(avg, 2), "low_lines": low, "lines": len(confs)}
        clear = avg >= VOICE_MIN_AVG and low / len(confs) <= VOICE_MAX_LOW
        rep["checks"].append(("Speech clarity", "done" if clear else "incomplete",
                              f"Average recognition {round(avg * 100)}%, {low} of {len(confs)} lines unclear."))
        if not clear:
            rep["findings"].append(make_finding(
                "SUPERVISION", "MEDIUM", "SPEECH_UNCLEAR", 0, f"{round(avg * 100)}% average recognition",
                "The speech was hard to recognise (noise, accent or low volume), so the transcript may contain mistakes.",
                "A person should listen to the video before it is approved.", avg))
    finish_report(rep)
    return rep


def finish_report(rep):
    """Show repeated advice once, number the findings, set the verdict. Safe to call again after adding findings."""
    seen_notes, kept = set(), []
    for f in rep["findings"]:
        if f["status"] == "NOTE":
            k = (f["type"], f["evidence"], f.get("source"))
            if k in seen_notes:
                continue
            seen_notes.add(k)
        kept.append(f)
    rep["findings"] = kept
    for i_, f in enumerate(rep["findings"]):
        f["i"] = i_
    finalize(rep)
    if rep["needs_confirmation"] and not any(f["severity"] == "CRITICAL" for f in rep["findings"]):
        rep["verdict"] = "CONFIRM - project needs confirming"


def add_ai_review(rep, brand_id, reg):
    """Run the AI reviewer on the finished report. Never raises and never changes a verdict on its own:
    paraphrase matches become 'needs a human check', everything else is advice."""
    try:
        units = AIR.units_from(rep.get("inputs", {}))
        result, state = AIR.get_or_run(brand_id, units, reg)
        rep["ai_review"] = {"state": state, "text": AIR.STATE_TEXT.get(state, ""),
                            "why": AIR.last_reason if state == "error" else "",
                            "summary": (result or {}).get("summary", ""),
                            "summary_en": (result or {}).get("summary_en", ""),
                            "tone_consistent": (result or {}).get("tone_consistent", True)}
        for f in AIR.findings_for(result, brand_id, make_finding):
            rep["findings"].append(f)
        n = sum(1 for f in rep["findings"] if f.get("ai"))
        rep["checks"].append(("AI review (paraphrase, tone, terminology, clarity)",
                              "done" if state in ("fresh", "cached") else "not_available",
                              AIR.STATE_TEXT.get(state, "") + (f" {n} point(s) raised." if n else "")))
        finish_report(rep)
    except Exception:
        rep["ai_review"] = {"state": "error", "text": AIR.STATE_TEXT["error"], "summary": "", "tone_consistent": True}
    return rep


def status_of(rep):
    """Turn a finished report into a status. Approval needs certain identity, a verified claim and no open issue."""
    if rep["verdict"].startswith("BLOCKED"):
        return "BLOCKED"
    if rep.get("needs_confirmation"):
        return "NEEDS_CONFIRMATION"
    if rep["flags"]:
        return "REJECTED"
    if rep["supervision"]:
        return "SUPERVISION"
    ident = rep.get("identity", {}).get("state")
    if ident in ("MATCHED", "MULTIPLE", "CONFIRMED") and rep.get("passed", 0) >= 1:
        return "AUTO_APPROVED"
    return "SUPERVISION"


def allowed_brand(u, brand_id):
    return bool(q("SELECT 1 FROM memberships WHERE user_id=? AND brand_id=?", (u["id"], brand_id)))


def store_projects(sid, rep, actor):
    ex("DELETE FROM submission_projects WHERE sub_id=?", (sid,))
    for pid in rep.get("project_ids", []):
        ex("INSERT INTO submission_projects (sub_id, project_id, identity_state, evidence, confirmed_by) VALUES (?,?,?,?,?)",
           (sid, pid, rep.get("identity", {}).get("state"),
            rep.get("identity", {}).get("evidence", {}).get(str(pid), ""),
            actor if rep.get("identity", {}).get("state") == "CONFIRMED" else None))


def review_submission(u, brand_id, path, sha, filename, language, source_url=None, caption="", hint_pid=None):
    """Review one file, then delete the uploaded copy. Only the report, transcript and file hash are kept, so the
    file can still be verified later by its hash, and uploads do not fill the disk."""
    try:
        return _review_submission(u, brand_id, path, sha, filename, language, source_url, caption, hint_pid)
    finally:
        remove_upload(path)


def remove_upload(path):
    if os.environ.get("KEEP_UPLOADS") == "1":
        return
    try:
        p = Path(path).resolve()
        if UPLOAD_DIR.resolve() in p.parents and p.is_file():
            p.unlink()
    except Exception:
        pass


def sweep_uploads(max_age_hours=24):
    """Delete leftover uploads (from crashes) older than a day, except files still waiting in the bulk queue."""
    if os.environ.get("KEEP_UPLOADS") == "1" or not UPLOAD_DIR.exists():
        return 0
    try:
        waiting = {str(Path(r["path"]).resolve()) for r in q("SELECT path FROM jobs WHERE status IN ('QUEUED','PROCESSING')")}
    except Exception:
        waiting = set()
    cutoff, n = time.time() - max_age_hours * 3600, 0
    for f in UPLOAD_DIR.iterdir():
        try:
            if f.is_file() and f.stat().st_mtime < cutoff and str(f.resolve()) not in waiting:
                f.unlink()
                n += 1
        except Exception:
            pass
    return n


def _review_submission(u, brand_id, path, sha, filename, language, source_url=None, caption="", hint_pid=None):
    """Review one file and store the result. Returns (submission_id, None) or (None, message)."""
    brand = q("SELECT * FROM brands WHERE id=?", (brand_id,))[0]
    reg = json.loads(brand["registry"])
    if not RE.authorized_projects(u["id"], brand_id):
        return None, (f"You have not been given access to any {brand['name']} project yet. "
                      "Ask the company to grant you access, then submit again.")
    if hint_pid and not RE.can_use_project(u["id"], hint_pid):
        return None, "You do not have access to that project."
    try:
        kind = FMT.kind_of(path)
        if kind == "image":
            segs, lines = [], FMT.read_image(path)
            screen = {"lines": lines, "frames": 1, "total": 1, "partial": False} if lines is not None else None
        elif kind == "document":
            segs = []
            screen = {"lines": FMT.read_document(path), "frames": 0, "total": 0, "partial": False}
        else:
            segs = transcribe_file(path, language, glossary_for(reg, brand_id, u["id"]))
            screen = None if source_url else OCR.read_screen(path, MAX_SECONDS)
        rep = analyze(u, brand_id, segs, caption, reg, hint_pid, is_link=bool(source_url) and kind == "video", screen=screen,
                      kind=kind)
        if kind == "image" and screen is not None:
            af, ac = AST.check_image(path, screen["lines"], reg, make_finding)
            rep["findings"] += af
            rep["checks"] += ac
            finish_report(rep)
        add_ai_review(rep, brand_id, reg)
    except Exception as ex_:
        msg = str(ex_) if isinstance(ex_, (ValueError, RuntimeError)) else "Something went wrong while reviewing this file."
        return None, msg

    dup = q("SELECT id FROM submissions WHERE sha256=? AND brand_id=? ORDER BY id LIMIT 1", (sha, brand_id))
    if dup:
        rep["duplicate_of"] = dup[0]["id"]
    status = status_of(rep)
    approval = new_approval_id() if status == "AUTO_APPROVED" else None
    sid = ex("""INSERT INTO submissions (user_id, brand_id, filename, sha256, language, status, verdict, report,
                created_at, approval_id, source_url) VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
             (u["id"], brand_id, filename, sha, language, status, rep["verdict"],
              json.dumps(rep, ensure_ascii=False), datetime.now().strftime("%Y-%m-%d %H:%M"),
              approval, source_url))
    store_projects(sid, rep, u["name"])
    record_lessons(sid, rep, brand_id)
    how = "SUBMITTED_LINK" if source_url else "SUBMITTED"
    log(sid, u["name"], how, f"{filename} for {brand['name']}", brand_id)
    log(sid, "BrandGuard engine", "AUTO_REVIEW",
        f"{rep['verdict']} ({rep['flags']} flags, {rep['supervision']} to check). Identity: {rep['identity']['state']}", brand_id)
    if dup:
        log(sid, "BrandGuard engine", "DUPLICATE_FOUND", f"Identical file already submitted as #{dup[0]['id']}", brand_id)
    if approval:
        log(sid, "BrandGuard engine", "AUTO_APPROVED", f"{approval} (policy {RE.POLICY_VERSION})", brand_id)
    return sid, None


def run_review(u, brand_id, path, sha, filename, language, source_url=None, caption="", hint_pid=None):
    sid, msg = review_submission(u, brand_id, path, sha, filename, language, source_url, caption, hint_pid)
    return err_page(u, msg) if msg else RedirectResponse(f"/s/{sid}", 303)




@app.post("/confirm/{sid}")
async def confirm_project(request: Request, sid: int):
    u = me(request)
    s = get_sub(sid)
    if not u or u["role"] != "submitter" or not s or s["user_id"] != u["id"]:
        return RedirectResponse("/", 303)
    if s["status"] != "NEEDS_CONFIRMATION":
        return RedirectResponse(f"/s/{sid}", 303)
    form = await request.form()
    ids = []
    for v in form.getlist("project_id"):
        if str(v).isdigit():
            ids.append(int(v))
    ids = list(dict.fromkeys(ids))
    if not ids:
        request.session["flash"] = "Choose at least one project to continue."
        return RedirectResponse(f"/s/{sid}", 303)
    for pid in ids:  # backend check, whatever the form sent
        if not RE.can_use_project(u["id"], pid) or RE.get_project(pid)["brand_id"] != s["brand_id"]:
            return err_page(u, "You do not have access to one of those projects.")
    old = rep_of(s)
    reg = json.loads(q("SELECT registry FROM brands WHERE id=?", (s["brand_id"],))[0]["registry"])
    segs = old.get("inputs", {}).get("segments", [])
    caption = old.get("inputs", {}).get("caption", "")
    scr_old = old.get("inputs", {}).get("screen")
    screen = None
    if scr_old is not None and old.get("screen_meta"):
        screen = {"lines": [{"start": l["start"], "end": l["end"], "text": l["text"], "raw": l.get("raw", l["text"]),
                             "conf": l["conf"]} for l in scr_old], **old["screen_meta"]}
    kind = old.get("kind", "video")
    if kind != "video" and screen is None and scr_old is not None:
        screen = {"lines": [{"start": l["start"], "end": l["end"], "text": l["text"], "raw": l.get("raw", l["text"]),
                             "conf": l["conf"]} for l in scr_old], "frames": 1, "total": 1, "partial": False}
    rep = analyze(u, s["brand_id"], segs, caption, reg, confirmed=ids, is_link=bool(s["source_url"]) and kind == "video",
                  screen=screen, kind=kind)
    keep = [f for f in old.get("findings", []) if f["type"] in ("LOW_CONTRAST", "SMALL_TEXT", "OFF_PALETTE", "LOGO_MISSING")]
    if keep:
        rep["findings"] += keep
        rep["checks"] += [c for c in old.get("checks", []) if c[0] in ("Text contrast and size", "Brand colours", "Brand logo")]
        finish_report(rep)
    add_ai_review(rep, s["brand_id"], reg)
    if old.get("duplicate_of"):
        rep["duplicate_of"] = old["duplicate_of"]
    status = status_of(rep)
    approval = new_approval_id() if status == "AUTO_APPROVED" else None
    ex("UPDATE submissions SET status=?, verdict=?, report=?, approval_id=? WHERE id=?",
       (status, rep["verdict"], json.dumps(rep, ensure_ascii=False), approval, sid))
    store_projects(sid, rep, u["name"])
    record_lessons(sid, rep, s["brand_id"])
    names = ", ".join(RE.get_project(i)["official_name"] for i in ids)
    log(sid, u["name"], "PROJECT_CONFIRMED", f"Confirmed project: {names}", s["brand_id"])
    log(sid, "BrandGuard engine", "AUTO_REVIEW", f"{rep['verdict']} ({rep['flags']} flags, {rep['supervision']} to check)", s["brand_id"])
    if approval:
        log(sid, "BrandGuard engine", "AUTO_APPROVED", f"{approval} (policy {RE.POLICY_VERSION})", s["brand_id"])
    return RedirectResponse(f"/s/{sid}", 303)


@app.post("/upload")
def upload(request: Request, brand_id: int = Form(...), language: str = Form("auto"),
           project_id: str = Form(""), file: UploadFile = File(...)):
    u = me(request)
    if not u or u["role"] != "submitter":
        return RedirectResponse("/", 303)
    if not allowed_brand(u, brand_id):
        return err_page(u, "You are not linked to that company.")
    try:
        path, sha = save_upload(file)
    except Exception as ex_:
        return err_page(u, str(ex_) if isinstance(ex_, (ValueError, RuntimeError)) else "Upload failed.")
    return run_review(u, brand_id, path, sha, file.filename, language,
                      hint_pid=int(project_id) if project_id.isdigit() else None)


@app.post("/upload_link")
def upload_link(request: Request, brand_id: int = Form(...), language: str = Form("auto"),
                url: str = Form(...), project_id: str = Form("")):
    u = me(request)
    if not u or u["role"] != "submitter":
        return RedirectResponse("/", 303)
    if not allowed_brand(u, brand_id):
        return err_page(u, "You are not linked to that company.")
    url = url.strip()
    if not host_ok(url):
        return err_page(u, "Please paste a link from YouTube, Instagram or Facebook.")
    try:
        path, sha, title, desc = fetch_link(url)
    except Exception as ex_:
        return err_page(u, str(ex_) if isinstance(ex_, (ValueError, RuntimeError)) else "Could not fetch this link.")
    caption = (title + "\n" + desc).strip()
    return run_review(u, brand_id, path, sha, title[:80], language, source_url=url, caption=caption,
                      hint_pid=int(project_id) if project_id.isdigit() else None)


@app.post("/upload_page")
def upload_page(request: Request, brand_id: int = Form(...), url: str = Form(...), project_id: str = Form("")):
    """A blog post or any public web page: its text is fetched and checked like a document."""
    u = me(request)
    if not u or u["role"] != "submitter":
        return RedirectResponse("/", 303)
    if not allowed_brand(u, brand_id):
        return err_page(u, "You are not linked to that company.")
    try:
        title, desc, text = FMT.fetch_page(url)
    except ValueError as ex_:
        return err_page(u, str(ex_))
    body = "\n\n".join(x for x in (title, desc, text) if x).strip()
    if len(body) < 20:
        return err_page(u, "Almost no text was found on that page.")
    path = UPLOAD_DIR / (uuid.uuid4().hex + ".txt")
    path.write_text(body, encoding="utf-8")
    sha = hashlib.sha256(body.encode("utf-8")).hexdigest()
    return run_review(u, brand_id, path, sha, (title or url)[:80], "auto", source_url=url.strip(),
                      hint_pid=int(project_id) if project_id.isdigit() else None)


# ------------------------------------------------------------- result / review
TYPE_LABEL = {
    "PRICE_MISMATCH": "Price does not match the approved price",
    "PRICE_UNIT_CHECK": "Possible unit or transcription error",
    "PRICE_LOW_CONFIDENCE": "Price heard with low confidence",
    "PRICE_NOT_IN_REGISTRY": "Price is not close to any approved price",
    "UNCLEAR_AMOUNT": "Amount could not be read",
    "BANNED_PHRASE": "Banned phrase used",
    "CONFIDENTIAL_DATA": "Confidential information mentioned",
    "UNKNOWN_CONTACT": "Contact detail not on the approved list",
    "UNSUPPORTED_CLAIM": "Claim needs supporting evidence",
    "NO_SPEECH_DETECTED": "No speech detected",
    "NO_VERIFIABLE_CLAIM": "Nothing could be verified",
    "AREA_CHECK": "Area does not match the approved area",
    "BATHROOMS_CHECK": "Bathroom count does not match",
    "POSSESSION_CHECK": "Possession claim needs attention",
    "RERA_CHECK": "RERA number needs attention",
    "OFFER_CHECK": "Offer claim needs attention",
    "DISTANCE_CHECK": "Distance claim does not match",
    "BOOKING_CHECK": "Booking amount does not match",
    "UNAUTHORIZED_PROJECT": "Project this account cannot promote",
    "PROJECT_UNCERTAIN": "Could not link a claim to one project",
    "IDENTITY_WEAK": "Project not named in the content",
    "CUSTOM_CHECK": "Product fact does not match the approved value",
    "BANNED_PHRASE_LIKELY": "Looks like a banned phrase (misread)",
    "CONFIDENTIAL_LIKELY": "Confidential figure mentioned",
    "LOW_CONTRAST": "Text is hard to read (low contrast)", "SMALL_TEXT": "Text smaller than the brand minimum",
    "OFF_PALETTE": "Colour outside the brand palette", "LOGO_MISSING": "Brand logo not found", "NO_TEXT_FOUND": "No text found",
    "TERMINOLOGY": "Preferred wording", "TONE": "Tone of voice", "SPEECH_UNCLEAR": "Speech was hard to recognise",
    "BANNED_PARAPHRASE": "Same meaning as a banned phrase (AI)", "AI_TONE": "Tone (AI review)",
    "AI_TERMINOLOGY": "Terminology (AI review)", "AI_CLARITY": "Hard to follow (AI review)",
    "AI_ADVICE": "Claim may need evidence (AI review)", "AI_LOWCONF": "Words that may be misheard (AI review)",
}
for _k in ("AREA", "BATHROOMS", "POSSESSION", "RERA", "OFFER", "DISTANCE", "BOOKING", "PRICE", "CUSTOM"):
    TYPE_LABEL[_k + "_UNVERIFIABLE"] = "No approved fact to compare with (not a confirmed error)"
UNCERTAIN = {"PRICE_UNIT_CHECK", "PRICE_LOW_CONFIDENCE", "UNCLEAR_AMOUNT", "PRICE_NOT_IN_REGISTRY"}
SECTIONS = [
    ("Price verification", lambda t: t.startswith("PRICE_") and not t.endswith("_UNVERIFIABLE") or t == "UNCLEAR_AMOUNT"),
    ("Product facts (your own fact types: discount, warranty, fees, dates...)", lambda t: t.startswith("CUSTOM_")),
    ("Property facts (area, bathrooms, possession, RERA, offers, distance)",
     lambda t: not t.startswith("CUSTOM_") and (t.endswith("_CHECK") and not t.startswith("PRICE_") or t.endswith("_UNVERIFIABLE"))),
    ("Project and permission", lambda t: t in ("UNAUTHORIZED_PROJECT", "PROJECT_UNCERTAIN", "IDENTITY_WEAK")),
    ("Restricted or banned language", lambda t: t in ("BANNED_PHRASE", "BANNED_PHRASE_LIKELY")),
    ("Confidential information", lambda t: t in ("CONFIDENTIAL_DATA", "CONFIDENTIAL_LIKELY")),
    ("Contact information", lambda t: t == "UNKNOWN_CONTACT"),
    ("Claims requiring evidence", lambda t: t == "UNSUPPORTED_CLAIM"),
    ("Different wording of a banned phrase (AI)", lambda t: t == "BANNED_PARAPHRASE"),
    ("Wording and tone suggestions", lambda t: t in ("TERMINOLOGY", "TONE", "AI_TONE", "AI_TERMINOLOGY", "AI_CLARITY",
                                                     "AI_ADVICE", "AI_LOWCONF")),
    ("Speech clarity", lambda t: t == "SPEECH_UNCLEAR"),
    ("Brand assets and accessibility", lambda t: t in ("LOW_CONTRAST", "SMALL_TEXT", "OFF_PALETTE", "LOGO_MISSING")),
]
SUGGESTABLE = ("BANNED_PHRASE", "UNSUPPORTED_CLAIM", "TERMINOLOGY", "TONE")
ACTION_LABEL = {
    "SUBMITTED": "Submission received", "SUBMITTED_LINK": "Link submission received",
    "AUTO_REVIEW": "Analysis completed", "AUTO_APPROVED": "Auto-approved", "APPROVED": "Approved by reviewer",
    "REJECTED": "Rejected by reviewer", "DISPUTE_RAISED": "Dispute raised",
    "DISPUTE_OVERTURNED": "Dispute resolved: decision overturned", "DISPUTE_UPHELD": "Dispute resolved: decision upheld",
    "RULE_ADDED": "Brand rule added", "RULE_CHANGED": "Brand rule changed", "RULE_REMOVED": "Brand rule removed",
    "FACTS_UPDATED": "Brand details updated", "PROJECT_CONFIRMED": "Project confirmed by submitter",
    "DUPLICATE_FOUND": "Identical file found", "PROJECT_ADDED": "Project added", "PROJECT_EDITED": "Project details changed",
    "FACT_SET": "Approved fact saved", "FACT_ENDED": "Approved fact withdrawn", "ACCESS_GRANTED": "Project access granted",
    "ACCESS_REVOKED": "Project access revoked", "UNIT_ADDED": "Property type added", "RERA_RECORDED": "RERA check recorded",
    "PRICE_SAVED": "Approved price saved", "PRICE_REMOVED": "Approved price removed",
    "RULES_IMPORTED": "Brand rules imported from document",
    "LESSON_LEARNED": "Lesson written from dispute", "LESSON_APPLIED": "Earlier decision applied",
    "LESSON_CHANGED": "Lesson changed", "COMPANY_CREATED": "Company created", "INVITE_CREATED": "Invite created",
    "INVITE_ACCEPTED": "Invite accepted", "INVITE_CANCELLED": "Invite cancelled", "MEMBER_REMOVED": "Member removed",
    "REWRITE_SUGGESTED": "Wording suggestion created", "PACK_APPLIED": "Industry rules added", "FACT_TYPE_SET": "Fact type saved", "FACT_TYPE_REMOVED": "Fact type removed",
    "IMPORT_APPLIED": "Products and facts imported", "SAMPLE_LOADED": "Sample data loaded", "AI_DISMISSED": "AI concern dismissed", "AI_BANNED": "AI concern added to banned phrases", "BATCH_STARTED": "Bulk batch queued", "BATCH_DONE": "Bulk batch finished",
}
DECISION_TEXT = {
    "AUTO_APPROVED": ("ok", "check-circle", "Every price matched the approved facts and nothing needed a person to check."),
    "APPROVED": ("ok", "check-circle", "A reviewer looked at this submission and approved it."),
    "REJECTED": ("bad", "x-circle", "One or more claims break the brand rules. Fix the issues below and resubmit, or dispute the decision if you think it is wrong."),
    "BLOCKED": ("bad", "x-circle", "Confidential company information was mentioned, so this cannot be approved. Remove it and resubmit."),
    "REJECTED_FINAL": ("bad", "x-circle", "A reviewer looked at this submission and upheld the rejection."),
    "NEEDS_CONFIRMATION": ("warn", "search", "BrandGuard is not certain which project this video is about, so it did not check prices against a guess. Choose the project below to continue."),
    "SUPERVISION": ("warn", "alert", "BrandGuard could not be sure about some parts of this content, so a person must check it. This is not a confirmed violation."),
    "DISPUTED": ("warn", "scale", "The creator disagreed with the rejection. A reviewer must uphold or overturn it."),
}


def finding_html(f, sid=None, rep=None, can_suggest=False, can_decide=False):
    tone = {"FLAG": "bad", "SUPERVISION": "warn", "NOTE": "neutral"}.get(f["status"], "ok")
    if f.get("dismissed"):
        tone = "neutral"
    kind = f["type"]
    label = TYPE_LABEL.get(kind, kind.replace("_", " ").title())
    cap = f.get("source") == "caption"
    where = ("Post title or description" if cap else f'On screen at {fmt_time(f["time"])}' if f.get("source") == "screen"
             else f'Spoken at {fmt_time(f["time"])}')
    k_ = (rep or {}).get("kind", "video")
    if k_ != "video" and f.get("source") == "screen":
        where = {"image": "In the image", "document": "In the document"}.get(k_, "In the content")
    chip_txt = {"SUPERVISION": "Needs a human check", "NOTE": "Suggestion"}.get(f["status"], "Violation")
    if f.get("dismissed"):
        chip_txt = "Dismissed by a builder"
    status_chip = ui.badge(chip_txt, tone, {"warn": "alert", "neutral": "info"}.get(tone, "x-circle"))
    conf = ""
    if f.get("confidence") is not None and not cap:
        conf = f' <span class="muted">Recognition confidence {round(float(f["confidence"]) * 100)}%</span>'
    note = ""
    if kind in UNCERTAIN:
        note = ('<dt>Note</dt><dd>This is an uncertainty, not a confirmed violation. The amount may have been misheard or '
                'a unit (lakh or crore) may be wrong, so a person should listen to this moment.</dd>')
    sug = ""
    if sid is not None and rep is not None and "i" in f and kind in SUGGESTABLE:
        got = (rep.get("suggestions") or {}).get(str(f["i"]))
        if got and got.get("rewrite"):
            by = "Rule-based replacement" if got.get("by") == "rule" else "AI-written. Check it before using it."
            sug = (f'<dt>Suggested wording</dt><dd><div class="muted">Instead of: {e(got["sentence"])}</div>'
                   f'<div><b>{e(got["rewrite"])}</b></div><div class="muted">{by}</div></dd>')
        elif got:
            sug = f'<dt>Suggested wording</dt><dd class="muted">{e(got.get("error") or "No suggestion available.")}</dd>'
        elif can_suggest:
            sug = (f'<dt>Suggested wording</dt><dd><form method="post" action="/suggest/{sid}/{f["i"]}" style="display:inline">'
                   f'<button class="btn secondary sm" type="submit">Suggest compliant wording</button></form></dd>')
    if f.get("ai"):
        if f.get("alternative"):
            sug += (f'<dt>Suggested wording</dt><dd><b>{e(f["alternative"])}</b>'
                    f'<div class="muted">Written by the AI reviewer. Check it before using it.</div></dd>')
        if can_decide and sid is not None and "i" in f and f["type"] == "BANNED_PARAPHRASE" and not f.get("dismissed") and not f.get("confirmed"):
            sug += (f'<dt>Decision</dt><dd><form method="post" action="/ai/{sid}/{f["i"]}/ban" style="display:inline" data-lock>'
                    f'<button class="btn sm" type="submit">Yes, same claim: add to banned phrases</button></form> '
                    f'<form method="post" action="/ai/{sid}/{f["i"]}/dismiss" style="display:inline" data-lock>'
                    f'<button class="btn secondary sm" type="submit">No, harmless: dismiss</button></form></dd>')
        elif f.get("confirmed"):
            sug += '<dt>Decision</dt><dd>Confirmed as the same claim and added to the banned phrases.</dd>'
    return (f'<div class="finding {tone}"><div class="finding-top">{status_chip}<b>{e(label)}</b>{conf}</div>'
            f'<dl><dt>Where</dt><dd>{e(where)}</dd>'
            f'<dt>Detected</dt><dd>{e(str(f["evidence"]))}</dd>'
            f'<dt>Why it matters</dt><dd>{e(f["explanation"])}</dd>{note}'
            f'<dt>What to do</dt><dd>{e(f["suggestion"])}</dd>{sug}</dl></div>')


VERDICT_BADGE = {
    "VERIFIED": ("Verified", "ok", "check-circle"), "VIOLATION": ("Violation", "bad", "x-circle"),
    "UNVERIFIABLE": ("Cannot verify", "warn", "alert"), "REVIEW": ("Person must review", "warn", "alert"),
    "INCOMPLETE": ("Incomplete", "warn", "alert"), "NA": ("Not applicable", "neutral", "info"),
}
KIND_LABEL = {"price": "Price", "booking": "Booking amount", "area": "Area", "bathrooms": "Bathrooms",
              "possession": "Possession", "rera": "RERA", "offer": "Offer", "distance": "Distance", "custom": "Product fact"}
IDENT_TEXT = {
    "MATCHED": ("ok", "Project identified from the content."), "CONFIRMED": ("ok", "Project confirmed by the submitter."),
    "MULTIPLE": ("info", "Several projects are mentioned. Each claim was checked against its own project."),
    "LIKELY": ("warn", "Likely match, but not certain."), "AMBIGUOUS": ("warn", "The name fits more than one project."),
    "NOT_FOUND": ("warn", "No project could be identified from the content."),
    "CONFLICT": ("warn", "The selected project differs from the project in the content."),
    "SELECTED": ("warn", "The content does not name the project. Only your selection was used."),
    "NOT_AUTHORIZED": ("bad", "The content promotes a project this account cannot use."),
}


def claims_table(claims):
    rows = []
    for c in claims:
        b = VERDICT_BADGE.get(c["verdict"], VERDICT_BADGE["REVIEW"])
        when = ("Caption" if c.get("source") == "caption" else
                f'{fmt_time(c.get("time") or 0)} (screen)' if c.get("source") == "screen" else fmt_time(c.get("time") or 0))
        off = e(str(c.get("official") or "-"))
        if c.get("fact_version"):
            off += f'<br><span class="muted">{e(c["fact_version"])}</span>'
        if c.get("source_title"):
            off += f'<br><span class="muted">Source: {e(c["source_title"])}</span>'
        why = e(c.get("reason") or "")
        if c.get("level") == "registry_match":
            why += " " + ui.badge("Registry match only", "warn", "info")
        elif c.get("level") == "authoritative":
            why += " " + ui.badge("Confirmed with official register", "ok", "check")
        rows.append([e(when), e(KIND_LABEL.get(c["kind"], c["kind"])), e(c.get("project") or "-"),
                     e(str(c.get("heard") or "")), off, ui.badge(b[0], b[1], b[2]), why])
    return ui.table(["Where", "Claim", "Project", "Said", "Approved value", "Result", "Why"], rows, "Claim checks")


def identity_card(rep):
    idn = rep.get("identity")
    if not idn:
        return ""
    tone, text = IDENT_TEXT.get(idn["state"], ("info", ""))
    names = idn.get("names", {})
    plist = ", ".join(names.get(str(i), "?") for i in idn.get("project_ids", []))
    inner = f"<p>{e(text)} {e(idn.get('note') or '')}</p>"
    if plist:
        inner += f"<p><b>Project:</b> {e(plist)}</p>"
    ev = [f"{e(names.get(k, 'Project'))}: {e(v)}" for k, v in idn.get("evidence", {}).items() if v]
    if ev:
        inner += '<p class="muted">Evidence: ' + "; ".join(ev) + "</p>"
    return ui.card("Project identification", inner)


def checks_card(rep):
    chk = rep.get("checks")
    if not chk:
        return ""
    lab = {"done": ("Done", "ok", "check-circle"), "pending": ("Waiting", "warn", "clock"),
           "incomplete": ("Incomplete", "warn", "alert"), "not_available": ("Not performed", "neutral", "info")}
    items = "".join(f'<li style="padding:10px 0;border-bottom:1px solid var(--border, #e5e7eb)">'
                    f'<div style="display:flex;gap:8px;align-items:center;justify-content:space-between;flex-wrap:wrap">'
                    f'<b>{e(n)}</b>{ui.badge(*lab.get(st, lab["pending"]))}</div>'
                    f'<div class="muted" style="margin-top:4px">{e(note)}</div></li>' for n, st, note in chk)
    return ui.card("What was and was not checked", f'<ul style="list-style:none;margin:0;padding:0">{items}</ul>',
                   f"Policy version {e(rep.get('policy', ''))}. A check that was not performed is never counted as passed.")


def timeline_html(rows):
    if not rows:
        return ui.empty("No events recorded")
    items = "".join(f'<li><span class="when">{e(a["at"])}</span><span><b>{e(ACTION_LABEL.get(a["action"], a["action"]))}</b>'
                    f' <span class="muted">by {e(a["actor"])}</span>'
                    f'{"<br>" + e(a["note"]) if a["note"] else ""}</span></li>' for a in rows)
    return f'<ul class="timeline">{items}</ul>'

    

@app.get("/s/{sid}", response_class=HTMLResponse)
def view_submission(request: Request, sid: int):
    u = me(request)
    if not u:
        return RedirectResponse("/", 303)
    s = get_sub(sid)
    if not s or not can_view(u, s):
        return page("Not found", ui.page_header("Submission not found")
                    + ui.card("", ui.empty("We could not find that submission",
                                           "It may not exist, or you may not have access to it.",
                                           '<a class="btn" href="/">Back to dashboard</a>')),
                    u, "", [("Not found", None)], request)
    rep = rep_of(s)
    admin = u["role"] == "admin"
    tone, ic, why = DECISION_TEXT.get(s["status"], ("neutral", "info", ""))
    findings = rep.get("findings", [])

    # ---- decision summary
    counts = (f'{ui.badge(str(rep.get("flags", 0)) + " violation(s)", "bad", "x-circle")} '
              f'{ui.badge(str(rep.get("supervision", 0)) + " to check", "warn", "alert")} '
              f'{ui.badge(str(rep.get("passed", 0)) + " claim(s) verified", "ok", "check-circle")}')
    approval = ""
    if s["approval_id"]:
        approval = (f'<dl class="kv" style="margin-top:14px"><dt>Approval ID</dt><dd><b>{e(s["approval_id"])}</b></dd>'
                    f'<dt>File hash (SHA-256)</dt><dd class="mono">{e(s["sha256"])}</dd></dl>')
    summary = ui.card("", f'<div class="verdict {tone}">{ui.icon(ic, 26)}<div><h2>{e(ui.STATUS.get(s["status"], (s["status"],))[0])}</h2>'
                      f'<p>{e(why)}</p><p style="margin-top:10px">{counts}</p></div></div>{approval}')

    # ---- actions
    actions, dialogs = "", ""
    said_val = (s["dispute_said"] or "") if "dispute_said" in s.keys() else ""
    if not admin and s["user_id"] == u["id"] and s["status"] in ("REJECTED", "BLOCKED"):
        actions = ui.card("What would you like to do?",
                          f'<p class="muted">Fix the issues and <a href="/submit">submit a corrected video</a>, or dispute the '
                          f'decision if you believe a finding is wrong.</p>'
                          f'<button type="button" class="btn secondary" data-dialog="dlg-dispute">{ui.icon("scale", 16)}Raise a dispute</button>')
        dialogs = (
            f'<dialog id="dlg-dispute" class="dlg" aria-labelledby="dlg-dispute-t">'
            f'<form method="post" action="/dispute/{sid}" data-lock>'
            f'<h2 id="dlg-dispute-t">Raise a dispute</h2>'
            f'<p class="muted">Explain why the finding is wrong. Mention the time in the video if you can. A builder will review it.</p>'
            f'<label for="dsp-said">What did you actually say? <span class="muted">(optional, helps the builder check quickly)</span></label>'
            f'<input id="dsp-said" name="said" maxlength="80" placeholder="Example: 1 crore 30 lakh">'
            f'<label for="dsp-note">Why is the finding wrong? <span class="req">(required)</span></label>'
            f'<textarea id="dsp-note" name="note" rows="3" maxlength="1000" required></textarea>'
            f'<div class="dlg-actions"><button type="button" class="btn secondary" data-close>Cancel</button>'
            f'<button type="submit" class="btn primary">Submit dispute</button></div></form></dialog>')
    elif not admin and s["status"] == "DISPUTED":
        you_said = f'<p><b>You said:</b> {e(said_val)}</p>' if said_val else ""
        actions = ui.card("Your dispute", f'<p>{e(s["dispute_note"] or "")}</p>{you_said}'
                          '<p class="muted">Waiting for a builder to review it. You will see the outcome here.</p>')
    if admin and s["brand_id"] == u["brand_id"] and s["status"] in QUEUE_STATUSES:
        disputed = s["status"] == "DISPUTED"
        ctx = ""
        if disputed:
            ctx = f'<p><b>Creator\'s explanation:</b> {e(s["dispute_note"] or "")}</p>'
            if said_val:
                ctx += f'<p><b>What the creator says they said:</b> {e(said_val)}</p>'
            ctx += ('<p class="muted">If you overturn a wording flag (a banned or unsupported phrase), an AI writes a lesson '
                    'from both notes. Next time the same phrase appears in the same context it passes by precedent. '
                    'Anything new is still flagged or sent to a person. Prices and facts are never learned. '
                    'Manage lessons under Lessons.</p>')
        yes, no = ("Overturn decision", "Uphold rejection") if disputed else ("Approve", "Reject")
        actions = ui.card("Your decision", ctx + '<p class="muted">Review the findings, then record a decision. '
                          'It is saved to the audit trail.</p>'
                          f'<div class="form-actions"><button type="button" class="btn success" data-dialog="dlg-approve">{ui.icon("check", 16)}{yes}</button>'
                          f'<button type="button" class="btn danger" data-dialog="dlg-reject">{ui.icon("x", 16)}{no}</button></div>')
        dialogs = (ui.confirm_dialog("dlg-approve", f"{yes}?",
                                     "This approves the content and issues an approval ID." if not disputed else
                                     "This approves the content and issues an approval ID. The creator's dispute is resolved in their favour.",
                                     f"/decide/{sid}", yes, "success", ("note", "Note", False), {"decision": "approve"})
                   + ui.confirm_dialog("dlg-reject", f"{no}?",
                                       "This is a final decision for this submission." if not disputed else
                                       "The original rejection stands and the dispute is closed.",
                                       f"/decide/{sid}", no, "danger", ("note", "Reason for the creator", True),
                                       {"decision": "reject"}))
    if s["decision_note"]:
        actions += ui.card("Reviewer note", f'<p style="margin:0">{e(s["decision_note"])}</p>')

    # ---- left column: details + transcript
    lang = s["language"] if s["language"] in ("mr", "hi", "en") else ""
    src = ""
    if s["source_url"]:
        src = (f'<dt>Source link</dt><dd><a href="{e(s["source_url"])}" target="_blank" rel="noopener noreferrer">'
               f'{e(s["source_url"])}</a></dd>')
    ckind = rep.get("kind", "video")
    klabel = {"image": "Image", "document": "Document" if not s["source_url"] else "Web page"}.get(ckind, "Video")
    langrow = (f'<dt>Spoken language</dt><dd>{e(LANG_LABEL.get(s["language"], s["language"] or ""))}</dd>' if ckind == "video" else "")
    details = ui.card("Submission details", f"""<dl class="kv"><dt>{klabel}</dt><dd>{e(s["filename"] or "Untitled")}</dd>
      <dt>Company</dt><dd>{e(s["bname"])}</dd><dt>Submitted by</dt><dd>{e(s["uname"])}</dd>
      <dt>Submitted</dt><dd>{e(s["created_at"])}</dd>{langrow}{src}</dl>"""
                      + (ui.alert("This was fetched from a link. The platform's copy can differ from the original file, so the "
                                  "file-hash check does not apply to link submissions.", "info") if s["source_url"] else ""))
    tr = rep.get("transcript", [])
    if tr:
        rows = "".join(
            f'<div class="tr-row{" low" if t["conf"] < THRESHOLD else ""}"><time>{fmt_time(t["start"])}</time>'
            f'<span{f" lang={lang}" if lang else ""}>{e(t["text"])}</span>'
            f'<span class="muted conf">{round(t["conf"] * 100)}%{" (low)" if t["conf"] < THRESHOLD else ""}</span></div>'
            for t in tr)
        transcript = ui.card("Transcript and recognition confidence",
                             f'<p class="muted">Confidential terms are redacted. Highlighted lines had low recognition '
                             f'confidence and may contain mistakes.</p>{rows}')
    elif ckind != "video":
        transcript = ""
    else:
        transcript = ui.card("Transcript and recognition confidence", ui.empty("No transcript available"))
    scr_lines = rep.get("inputs", {}).get("screen") or []
    meta = rep.get("screen_meta")
    if meta is not None:
        if scr_lines:
            srows = "".join(
                f'<div class="tr-row"><time>{fmt_time(l["start"])}</time><span>{e(l.get("raw") or l["text"])}</span>'
                f'<span class="muted conf">{round(l["conf"] * 100)}%</span></div>' for l in scr_lines)
        else:
            srows = ui.empty("No text was found on screen")
        if ckind == "video":
            transcript += ui.card("Text read from the screen",
                                  f'<p class="muted">{meta["frames"]} frames read (2 per second plus every scene change). '
                                  f'Each line is shown once, at the first time it appeared.</p>{srows}')
        else:
            transcript += ui.card("Text read from the image" if ckind == "image" else "Text of the document",
                                  '<p class="muted">Each line that was checked. On posters, a price on its own line is joined '
                                  'with the label above it, so it is checked against the right variant.</p>' + srows)

    cleared = rep.get("cleared") or []
    if cleared:
        crows = "".join(
            f'<div class="tr-row"><time>{fmt_time(c["time"])}</time><span><b>"{e(c["evidence"])}"</b> not flagged. {e(c["summary"])} '
            f'<a href="/s/{c["source_sub"]}">Decision in #{c["source_sub"]}</a></span></div>' for c in cleared)
        transcript += ui.card("Cleared by earlier builder decisions",
                              '<p class="muted">A builder already decided this wording is harmless in this context. '
                              'Anything new would still be flagged.</p>' + crows)

    # ---- right column: grouped findings
    parts, clean = [], []
    used = set()
    can_sug = (admin and s["brand_id"] == u["brand_id"]) or s["user_id"] == u["id"]
    can_dec = admin and s["brand_id"] == u["brand_id"]
    claims = rep.get("claims", [])
    for title, test in SECTIONS:
        fs = [f for f in findings if test(f["type"])]
        used.update(id(f) for f in fs)
        inner = ""
        if title == "Price verification" and claims and "verdict" in claims[0]:
            pass
        elif title == "Price verification" and claims:
            crow = []
            for c in claims:
                when = "Caption" if c.get("source") == "caption" else fmt_time(c["time"])
                crow.append([e(when), e(c["heard"]), e(c["config"]), e(c["approved"]),
                             {"PASS": ui.badge("Matches", "ok", "check-circle"), "FLAG": ui.badge("Mismatch", "bad", "x-circle")}.get(
                                 c["status"], ui.badge("Needs check", "warn", "alert"))])
            inner += ui.table(["Time", "Detected", "Compared with", "Approved", "Result"], crow, "Price claims")
            inner += '<p></p>'
        inner += "".join(finding_html(f, sid, rep, can_sug, can_dec) for f in fs)
        if inner:
            parts.append(f'<h3 class="sec-title">{e(title)}</h3>{inner}')
        else:
            clean.append(title.lower())
    other = [f for f in findings if id(f) not in used]
    if other:
        parts.append('<h3 class="sec-title">Other checks</h3>' + "".join(finding_html(f, sid, rep, can_sug, can_dec) for f in other))
    if not findings and not claims:
        parts.append(ui.empty("No issues found", "Nothing in this video broke a rule."))
    if clean and (findings or claims):
        parts.append(f'<p class="muted" style="margin-top:12px">{ui.icon("check-circle", 14)} No issues found in: {", ".join(clean)}.</p>')
    if s["source_url"]:
        cf = [f for f in findings if f.get("source") == "caption"]
        parts.append('<h3 class="sec-title">Link metadata</h3><p class="muted">'
                     + (f'{len(cf)} finding(s) came from the post title or description (marked above).' if cf
                        else "The post title and description were checked and had no findings.") + "</p>")
    air = rep.get("ai_review") or {}
    ai_card = ""
    if air:
        body_ = f'<p class="muted">{e(air.get("text", ""))}</p>'
        if air.get("why") and admin:
            body_ += f'<p class="muted" style="font-size:13px">Reason: {e(air["why"])}</p>'
        if air.get("summary"):
            en = air.get("summary_en") or ""
            body_ = (f'<p><b>What was said:</b> {e(air["summary"])}</p>'
                     + (f'<p class="muted"><b>In English:</b> {e(en)}</p>' if en and en != air["summary"] else "")
                     + body_)
        ai_card = ui.card("AI review", body_, "Advice only. The AI never approves or rejects. Rules and people decide.")
    findings_card = ui.card("Compliance findings", "".join(parts))
    new_claims = [c for c in claims if "verdict" in c]
    claims_card = (ui.card("Claim-by-claim checks", claims_table(new_claims),
                           "Every claim found in the content, what it was compared with, and the result.") if new_claims else "")
    extra = identity_card(rep) + checks_card(rep)
    notes = ""
    if rep.get("duplicate_of") and rep["duplicate_of"] != sid:
        notes += ui.alert(f'An identical file was already submitted as #{rep["duplicate_of"]}.', "info", "Duplicate file.")
    if s["approval_id"]:
        rs, why_ = RE.reuse_status(sid, rep, s["user_id"], s["brand_id"])
        if rs == "OK":
            notes += ui.alert("Facts used for this approval are still current and the creator still has access. "
                              "This is a check at this moment, not a permanent guarantee.", "ok", "Reuse check: no change found.")
        else:
            notes += ui.alert("; ".join(why_) + ". The original approval is unchanged, but this content should be re-checked before it is used again.",
                              "warn", "Re-check needed before reuse.")
    confirm = ""
    if not admin and s["user_id"] == u["id"] and s["status"] == "NEEDS_CONFIRMATION":
        cands = [RE.get_project(i) for i in rep.get("candidates", [])]
        auth = RE.authorized_projects(u["id"], s["brand_id"])
        opts_ = {p["id"]: p for p in auth}
        for c_ in cands:
            if c_ and c_["id"] not in opts_ and False:
                pass
        pool = [p for p in auth] if not cands or any(c_ is None for c_ in cands) else [opts_[c_["id"]] for c_ in cands if c_ and c_["id"] in opts_]
        if not pool:
            pool = auth
        boxes = "".join(f'<label class="check"><input type="checkbox" name="project_id" value="{p["id"]}"> '
                        f'{e(p["official_name"])}{" (" + e(p["city"]) + ")" if p["city"] else ""}</label>' for p in pool)
        confirm = ui.card("Which project is this video about?",
                          f'<form method="post" action="/confirm/{sid}" data-lock><p class="muted">Tick every project the video promotes. '
                          f'Prices and facts are then checked against those projects only.</p>{boxes}'
                          f'<p></p><button class="btn" type="submit">Confirm and continue</button></form>')

    au = q("SELECT * FROM audit WHERE sub_id=? ORDER BY id", (sid,))
    audit = ui.card("History", timeline_html(au))

    crumbs = [("Review Queue" if admin else "My Submissions", "/queue" if admin else "/reviews"), (f"Submission #{sid}", None)]
    body = (ui.page_header(f"Submission #{sid}", e(s["filename"] or ""), '<a class="btn secondary" href="/reviews">Back to list</a>')
            + '<div class="stack">' + summary + notes + confirm + actions + '<div class="cols"><div class="stack">' + details + extra + transcript
            + '</div><div class="stack">' + claims_card + ai_card + findings_card + audit + '</div></div></div>' + dialogs)
    return page(f"Submission {sid}", body, u, "queue" if admin and s["status"] in QUEUE_STATUSES else "reviews", crumbs, request)


@app.post("/dispute/{sid}")
def dispute(request: Request, sid: int, note: str = Form(""), said: str = Form("")):
    u, s = me(request), get_sub(sid)
    if u and s and s["user_id"] == u["id"] and s["status"] in ("REJECTED", "BLOCKED") and note.strip():
        ex("UPDATE submissions SET status='DISPUTED', dispute_note=?, dispute_said=? WHERE id=?",
           (note.strip()[:1000], said.strip()[:80], sid))
        detail = note.strip()[:300]
        if said.strip():
            detail += f" (creator says they said: {said.strip()[:80]})"
        log(sid, u["name"], "DISPUTE_RAISED", detail)
        request.session["flash"] = "Dispute sent to the builder."
    return RedirectResponse(f"/s/{sid}", 303)


@app.post("/decide/{sid}")
def decide(request: Request, sid: int, decision: str = Form(...), note: str = Form("")):
    u, s = me(request), get_sub(sid)
    if (u and s and u["role"] == "admin" and s["brand_id"] == u["brand_id"]
            and s["status"] in ("SUPERVISION", "DISPUTED")):
        was_dispute = s["status"] == "DISPUTED"
        if decision == "approve":
            status, approval = "APPROVED", new_approval_id()
            action = "DISPUTE_OVERTURNED" if was_dispute else "APPROVED"
        else:
            status, approval = "REJECTED_FINAL", None
            action = "DISPUTE_UPHELD" if was_dispute else "REJECTED"
        ex("UPDATE submissions SET status=?, approval_id=?, decision_note=? WHERE id=?",
           (status, approval, note.strip()[:1000], sid))
        log(sid, u["name"], action, (approval or "") + (" " + note.strip()[:300] if note.strip() else ""))
        if was_dispute and decision == "approve":
            try:
                made = LRN.create_from_dispute(get_sub(sid), rep_of(s), s["dispute_note"] or "", note.strip()[:300], u["name"])
                for m in made:
                    if m["reused"]:
                        continue
                    log(sid, "BrandGuard engine", "LESSON_LEARNED",
                        f'Lesson #{m["id"]} ({"active" if m["active"] else "waiting for approval"}): {m["summary"]}', u["brand_id"])
                if made and any(not m["reused"] for m in made):
                    request.session["flash"] = ("Decision recorded. BrandGuard wrote a lesson from this dispute. "
                                                "Prices and facts are never learned. Review it under Lessons.")
                    return RedirectResponse(f"/s/{sid}", 303)
            except Exception:
                pass  # learning must never block a decision
        request.session["flash"] = "Decision recorded: " + ui.STATUS[status][0] + "."
    return RedirectResponse(f"/s/{sid}", 303)

    

# ------------------------------------------------------------------ list views
TABS = {
    "all": ("All", None), "supervision": ("Supervision Required", ("SUPERVISION",)),
    "disputed": ("Disputed", ("DISPUTED",)), "approved": ("Approved", APPROVED_STATUSES),
    "rejected": ("Rejected", REJECTED_STATUSES),
}


def scope_sql(request, u):
    if u["role"] == "admin":
        return "s.brand_id=?", [u["brand_id"]]
    return "s.user_id=? AND s.brand_id=?", [u["id"], current_company(request, u)]


def list_view(request, mode):
    u = me(request)
    if not u or (mode == "queue" and u["role"] != "admin"):
        return RedirectResponse("/", 303)
    admin = u["role"] == "admin"
    tab_keys = ("all", "supervision", "disputed") if mode == "queue" else tuple(TABS)
    st = request.query_params.get("status", "all")
    st = st if st in tab_keys else "all"
    term = request.query_params.get("q", "").strip()[:80]
    try:
        pg = max(1, int(request.query_params.get("page", "1")))
    except ValueError:
        pg = 1
    where, args = scope_sql(request, u)
    if mode == "queue":
        where += " AND s.status IN ('SUPERVISION','DISPUTED')"
    counts = Counter({r["status"]: r["c"] for r in q(
        f"SELECT s.status, COUNT(*) AS c FROM submissions s WHERE {where} GROUP BY s.status", tuple(args))})

    def tab_count(k):
        sts = TABS[k][1]
        return sum(counts.values()) if sts is None else sum(counts[x] for x in sts)

    w2, a2 = where, list(args)
    if TABS[st][1]:
        w2 += " AND s.status IN (%s)" % ",".join("?" * len(TABS[st][1]))
        a2 += list(TABS[st][1])
    if term:
        w2 += " AND (s.filename LIKE ? OR u.name LIKE ? OR s.approval_id LIKE ? OR CAST(s.id AS TEXT)=?)"
        a2 += [f"%{term}%", f"%{term}%", f"%{term}%", term.lstrip("#")]
    total = q(f"SELECT COUNT(*) AS c FROM submissions s JOIN users u ON u.id=s.user_id WHERE {w2}", tuple(a2))[0]["c"]
    pages = max(1, math.ceil(total / PAGE_SIZE))
    pg = min(pg, pages)
    rows = q(SUB_SQL + f"WHERE {w2} ORDER BY s.id DESC LIMIT {PAGE_SIZE} OFFSET {(pg - 1) * PAGE_SIZE}", tuple(a2))

    base = "/queue" if mode == "queue" else "/reviews"
    tabs = "".join(
        f'<a href="{base}?status={k}{"&q=" + quote(term) if term else ""}"{" aria-current=true" if k == st else ""}>'
        f'{e(TABS[k][0])} ({tab_count(k)})</a>' for k in tab_keys)
    toolbar = (f'<div class="toolbar"><nav class="tabs" aria-label="Filter by status">{tabs}</nav></div>'
               f'<div class="toolbar"><form method="get" action="{base}" role="search"><input type="hidden" name="status" value="{st}">'
               f'<label for="q" class="sr-only" style="position:absolute;left:-9999px">Search submissions</label>'
               f'<input id="q" type="search" name="q" value="{e(term)}" placeholder="Search by file, creator or approval ID">'
               f'<button class="btn secondary" type="submit">{ui.icon("search", 16)}Search</button></form></div>')
    if rows:
        listing = sub_table(rows, show_user=admin, label="Submissions") + ui.pager(
            f"{base}?status={st}" + (f"&q={quote(term)}" if term else ""), pg, pages)
    elif term or st != "all":
        listing = ui.empty("No matching submissions", "Try a different filter or search term.",
                           f'<a class="btn secondary" href="{base}">Clear filters</a>')
    else:
        listing = ui.empty("Nothing waiting for you" if mode == "queue" else "No submissions yet",
                           "New items that need a decision will appear here." if mode == "queue" else
                           "Submit a video to see its review here.",
                           '' if admin else '<a class="btn" href="/submit">Submit a Video</a>')
    if mode == "queue":
        head = ui.page_header("Review Queue", f"{sum(counts.values())} submission(s) need a decision. Open one to see the evidence and decide.")
        crumbs, active = [("Review Queue", None)], "queue"
    else:
        head = ui.page_header("All Reviews" if admin else "My Submissions",
                              "Every submission and its current status." if admin else "Every video you have submitted for the selected company.",
                              "" if admin else '<a class="btn" href="/submit">Submit a Video</a>')
        crumbs, active = [("All Reviews" if admin else "My Submissions", None)], "reviews"
    return page(("Review Queue" if mode == "queue" else "Reviews"), head + ui.card("", toolbar + listing), u, active, crumbs, request)


@app.get("/queue", response_class=HTMLResponse)
def queue_page(request: Request):
    return list_view(request, "queue")


@app.get("/reviews", response_class=HTMLResponse)
def reviews_page(request: Request):
    return list_view(request, "reviews")


# --------------------------------------------------------------------- disputes
@app.get("/disputes", response_class=HTMLResponse)
def disputes_page(request: Request):
    u = me(request)
    if not u:
        return RedirectResponse("/", 303)
    admin = u["role"] == "admin"
    where, args = scope_sql(request, u)
    rows = q(SUB_SQL + f"WHERE {where} AND s.id IN (SELECT sub_id FROM audit WHERE action='DISPUTE_RAISED') ORDER BY s.id DESC",
             tuple(args))
    open_rows = [r for r in rows if r["status"] == "DISPUTED"]
    done = [r for r in rows if r["status"] != "DISPUTED"]

    def tbl(rs, resolved):
        if not rs:
            return ui.empty("No resolved disputes" if resolved else "No open disputes",
                            "Disputes appear here after a creator challenges a rejection." if admin else
                            "If a rejection looks wrong, open the submission and choose Raise a dispute.")
        heads = ["Submission"] + (["Creator"] if admin else []) + ["Explanation", "Outcome" if resolved else "Status", "Action"]
        out = []
        for r in rs:
            outcome = ({"APPROVED": ui.badge("Decision overturned", "ok", "check-circle"),
                        "REJECTED_FINAL": ui.badge("Decision upheld", "bad", "x-circle")}.get(r["status"], ui.chip(r["status"]))
                       if resolved else ui.chip(r["status"]))
            lbl = "Resolve" if admin and not resolved else "View"
            out.append([f'<a href="/s/{r["id"]}"><b>#{r["id"]}</b></a><span class="trunc">{e(r["filename"] or "")}</span>']
                       + ([e(r["uname"])] if admin else [])
                       + [f'<span class="trunc" style="max-width:340px" title="{e(r["dispute_note"] or "")}">{e(r["dispute_note"] or "")}</span>',
                          outcome, f'<a class="btn sm" href="/s/{r["id"]}" aria-label="{lbl} dispute for submission {r["id"]}">{lbl}</a>'])
        return ui.table(heads, out, "Disputes")
    body = (ui.page_header("Disputes", "A creator can dispute a rejection. A builder then upholds the rejection or overturns it. "
                           "Every resolution is recorded in the audit trail.")
            + '<div class="stack">' + ui.card(f"Open ({len(open_rows)})", tbl(open_rows, False))
            + ui.card(f"Resolved ({len(done)})", tbl(done, True)) + '</div>')
    return page("Disputes", body, u, "disputes", [("Disputes", None)], request)


# ------------------------------------------------------------------- approvals
@app.get("/approvals", response_class=HTMLResponse)
def approvals_page(request: Request):
    u = me(request)
    if not u:
        return RedirectResponse("/", 303)
    admin = u["role"] == "admin"
    where, args = scope_sql(request, u)
    term = request.query_params.get("q", "").strip()[:60]
    w2, a2 = where + " AND s.approval_id IS NOT NULL", list(args)
    if term:
        w2 += " AND (s.approval_id LIKE ? OR s.filename LIKE ?)"
        a2 += [f"%{term}%", f"%{term}%"]
    rows = q(f"""SELECT s.*, u.name AS uname, b.name AS bname,
                 (SELECT MAX(at) FROM audit a WHERE a.sub_id=s.id AND a.action IN ('AUTO_APPROVED','APPROVED','DISPUTE_OVERTURNED')) AS approved_at
                 FROM submissions s JOIN users u ON u.id=s.user_id JOIN brands b ON b.id=s.brand_id
                 WHERE {w2} ORDER BY s.id DESC LIMIT 200""", tuple(a2))
    search = (f'<div class="toolbar"><form method="get" action="/approvals" role="search">'
              f'<label for="q" style="position:absolute;left:-9999px">Search approval records</label>'
              f'<input id="q" type="search" name="q" value="{e(term)}" placeholder="Search by approval ID or file">'
              f'<button class="btn secondary" type="submit">{ui.icon("search", 16)}Search</button></form></div>')
    if rows:
        out = []
        for r in rows:
            h = r["sha256"] or ""
            verif = (ui.badge("Not applicable (link)", "neutral", "link") if r["source_url"]
                     else ui.badge("Verifiable by file", "ok", "file-check"))
            out.append([f'<b>{e(r["approval_id"])}</b>', f'<a href="/s/{r["id"]}">#{r["id"]}</a><span class="trunc">{e(r["filename"] or "")}</span>',
                        e(r["bname"]), f'<span class="muted">{e(r["approved_at"] or r["created_at"])}</span>',
                        f'<span class="mono" title="{e(h)}">{e(h[:12])}…</span>' if h else "-", verif])
        listing = ui.table(["Approval ID", "Submission", "Company", "Approved", "SHA-256", "Verification"], out, "Approval records")
    else:
        listing = ui.empty("No approval records yet", "An approval ID is issued when a submission is auto-approved or a reviewer approves it.")
    body = (ui.page_header("Approval Records", "Each approved file has an approval ID and a SHA-256 fingerprint. "
                           "You can check a file against these records on the Verify a File page.",
                           '<a class="btn secondary" href="/verify">Verify a File</a>')
            + ui.card("", search + listing))
    return page("Approval Records", body, u, "approvals", [("Approval Records", None)], request)


# ----------------------------------------------------------------------- audit
@app.get("/audit", response_class=HTMLResponse)
def audit_page(request: Request):
    u = me(request)
    if not u:
        return RedirectResponse("/", 303)
    admin = u["role"] == "admin"
    term = request.query_params.get("q", "").strip()[:60]
    try:
        pg = max(1, int(request.query_params.get("page", "1")))
    except ValueError:
        pg = 1
    if admin:
        where, args = "(s.brand_id=? OR a.brand_id=?)", [u["brand_id"], u["brand_id"]]
    else:
        where, args = "s.user_id=? AND s.brand_id=?", [u["id"], current_company(request, u)]
    if term:
        where += " AND (a.actor LIKE ? OR a.action LIKE ? OR a.note LIKE ? OR CAST(a.sub_id AS TEXT)=?)"
        args += [f"%{term}%", f"%{term.upper().replace(' ', '_')}%", f"%{term}%", term.lstrip("#")]
    frm = "FROM audit a LEFT JOIN submissions s ON s.id=a.sub_id"
    total = q(f"SELECT COUNT(*) AS c {frm} WHERE {where}", tuple(args))[0]["c"]
    pages = max(1, math.ceil(total / 25))
    pg = min(pg, pages)
    rows = q(f"SELECT a.* {frm} WHERE {where} ORDER BY a.id DESC LIMIT 25 OFFSET {(pg - 1) * 25}", tuple(args))
    search = (f'<div class="toolbar"><form method="get" action="/audit" role="search">'
              f'<label for="q" style="position:absolute;left:-9999px">Search the audit trail</label>'
              f'<input id="q" type="search" name="q" value="{e(term)}" placeholder="Search by person, action or submission number">'
              f'<button class="btn secondary" type="submit">{ui.icon("search", 16)}Search</button></form></div>')
    if rows:
        out = [[f'<span class="muted">{e(a["at"])}</span>', e(a["actor"]), f'<b>{e(ACTION_LABEL.get(a["action"], a["action"]))}</b>',
                f'<a href="/s/{a["sub_id"]}">#{a["sub_id"]}</a>' if a["sub_id"] else "-", e(a["note"] or "")] for a in rows]
        listing = ui.table(["When", "Who", "Event", "Submission", "Details"], out, "Audit trail") + \
            ui.pager("/audit" + (f"?q={quote(term)}" if term else ""), pg, pages)
    else:
        listing = ui.empty("No events found", "Events appear as submissions are made and decided." if not term else "Try a different search.")
    body = (ui.page_header("Audit Trail", "A history of what happened and who did it, newest first. This is a record of past events, "
                           "not the current status of a submission. File verification checks are not recorded here.")
            + ui.card("", search + listing))
    return page("Audit Trail", body, u, "audit", [("Audit Trail", None)], request)


# ---------------------------------------------------------------- admin views
def setup_checklist(u, brand):
    """Getting-started steps for a new company. Returns '' once every step is done or the builder hid it."""
    reg = json.loads(brand["registry"] or "{}")
    ob = reg.get("onboarding", {})
    if ob.get("hidden"):
        return ""
    bid = u["brand_id"]
    general = reg.get("domain", "property") != "property"
    things = "products" if general else "projects"
    has_facts = bool(RE._q("SELECT 1 FROM facts f JOIN projects p ON p.id=f.project_id WHERE p.brand_id=? LIMIT 1", (bid,)))
    has_team = bool(q("SELECT 1 FROM invites WHERE brand_id=? LIMIT 1", (bid,)) or q("SELECT 1 FROM memberships WHERE brand_id=? LIMIT 1", (bid,)))
    has_sub = bool(q("SELECT 1 FROM submissions WHERE brand_id=? LIMIT 1", (bid,)))
    steps = [
        (bool(ob.get("rules_seen")), "Check your brand rules", "Your industry's starter rules are already in. Remove anything that does not apply, add your own banned phrases and confidential terms, and set brand colours and logo.", "/admin/rules", "Open Brand Rules"),
        (has_facts, f"Add your {things} and their approved facts", "Prices, discounts, dates and other facts that every claim is checked against. Fastest: import the spreadsheet, or load sample data to try it out.", "/admin/projects", f"Open {things.title()} and Facts"),
        (has_team, "Invite your team", "Channel partners, freelancers and agencies get a one-time link and can only submit for what you allow.", "/admin/team", "Invite people"),
        (has_sub, "See your first review", "When someone submits content, the result appears in the Review Queue. To see a result now, use the demo files with a test partner account.", "/test-audio", "Open demo files"),
    ]
    if all(s[0] for s in steps):
        return ""
    done = sum(1 for s in steps if s[0])
    def _step_link(ok, href, btn):
        if ok:
            return ""
        return f'<a class="btn secondary sm" style="margin-top:8px" href="{href}">{e(btn)}</a>'

    items = "".join(f'<li class="{"done" if ok else ""}"><span class="dot">{"✓" if ok else n}</span><div><b>{e(t)}</b>'
                    f'<div class="muted">{e(d)}</div>{_step_link(ok, href, btn)}</div></li>'
                    for n, (ok, t, d, href, btn) in enumerate(steps, 1))
    return (f'<section class="card" aria-labelledby="gs"><div class="card-head"><h2 id="gs">Get started: {done} of {len(steps)} done</h2>'
            f'<span><button type="button" class="btn secondary sm" data-tour-start>Take the tour</button> '
            f'<form method="post" action="/admin/onboarding/hide" style="display:inline"><button class="btn secondary sm" type="submit">Hide</button></form></span></div>'
            f'<ol class="checklist">{items}</ol></section>')


@app.post("/admin/onboarding/hide")
def onboarding_hide(request: Request):
    u = admin_only(request)
    if u:
        reg = load_reg(u)
        reg.setdefault("onboarding", {})["hidden"] = True
        save_reg(u, reg)
    return RedirectResponse("/", 303)


def admin_dash(request, u):
    bid = u["brand_id"]
    brand = q("SELECT * FROM brands WHERE id=?", (bid,))[0]
    allrows = q(SUB_SQL + "WHERE s.brand_id=? ORDER BY s.id DESC", (bid,))
    queue = [r for r in allrows if r["status"] in QUEUE_STATUSES]
    n = Counter(r["status"] for r in allrows)
    approved = sum(n[s] for s in APPROVED_STATUSES)
    rejected = sum(n[s] for s in REJECTED_STATUSES)

    types = Counter()
    for r in allrows:
        for f in rep_of(r).get("findings", []):
            types[f["type"]] += 1
    ov = q("""SELECT a.action, COUNT(*) AS c FROM audit a JOIN submissions s ON s.id=a.sub_id
              WHERE s.brand_id=? AND a.action IN ('DISPUTE_OVERTURNED','DISPUTE_UPHELD') GROUP BY a.action""", (bid,))
    ovd = {r["action"]: r["c"] for r in ov}
    recent = q("""SELECT a.* FROM audit a LEFT JOIN submissions s ON s.id=a.sub_id
                  WHERE s.brand_id=? OR a.brand_id=? ORDER BY a.id DESC LIMIT 6""", (bid, bid))

    welcome = f"""<section class="card welcome" aria-labelledby="wl"><div class="card-head"><h2 id="wl">Welcome to your compliance workspace</h2>
      <button type="button" class="btn secondary sm" data-dismiss="builder">Dismiss</button></div>
      <p>Creators submit videos for {e(brand["name"])}. BrandGuard checks them against your approved facts and rules, and sends
      anything uncertain to you. Start with the Review Queue.</p>
      <ol class="steps"><li>Review submissions</li><li>Examine evidence</li><li>Approve, reject, or request supervision</li><li>Record the decision</li></ol></section>"""
    setup = setup_checklist(u, brand)
    if setup:
        welcome = setup
    if queue:
        prio = sub_table(queue[:5], show_user=True, label="Priority reviews")
    else:
        prio = ui.empty("You are all caught up", "Nothing is waiting for a decision right now.")
    if types:
        common = "<ul style='margin:0;padding-left:20px'>" + "".join(
            f"<li>{e(TYPE_LABEL.get(k, k))}: <b>{v}</b></li>" for k, v in types.most_common(5)) + "</ul>"
        insights = (common + f'<p style="margin-top:12px">Disputes overturned: <b>{ovd.get("DISPUTE_OVERTURNED", 0)}</b> · '
                    f'upheld: <b>{ovd.get("DISPUTE_UPHELD", 0)}</b></p>'
                    '<p class="muted">Findings that creators successfully dispute may point to guidelines that need clearer wording.</p>')
    else:
        insights = ui.empty("No findings yet", "Insights appear after submissions are reviewed.")
    body = (ui.page_header("Compliance Overview",
                           f"Review activity for {e(brand['name'])}. Pending decisions are listed first.",
                           f'<a class="btn" href="/queue">{ui.icon("inbox", 16)}Review Queue ({len(queue)})</a>')
            + '<div class="stack">' + welcome
            + '<div class="stats">'
            + ui.stat("Awaiting review", len(queue), f"{n['SUPERVISION']} supervision, {n['DISPUTED']} disputed", "/queue", "info")
            + ui.stat("Supervision required", n["SUPERVISION"], "BrandGuard was unsure", "/queue?status=supervision", "warn")
            + ui.stat("Approved", approved, "Auto and reviewer approvals", "/reviews?status=approved", "ok")
            + ui.stat("Rejected", rejected, "", "/reviews?status=rejected", "bad")
            + '</div>'
            + ui.card("Priority reviews", prio, actions='<a href="/queue">Open queue</a>' if queue else "")
            + '<div class="grid-2">' + ui.card("Recent activity", timeline_html(recent), actions='<a href="/audit">Full audit trail</a>')
            + ui.card("Compliance insights", insights) + '</div></div>')
    return page("Overview", body, u, "overview", [], request)

    

def lines(s):
    return [x.strip() for x in (s or "").splitlines() if x.strip()]


def parse_price(s):
    s = s.strip()
    try:
        return int(float(s.replace(",", "")))
    except ValueError:
        pass
    found = extract_prices(s)
    return found[0]["value"] if found else None


def load_reg(u):
    return json.loads(q("SELECT registry FROM brands WHERE id=?", (u["brand_id"],))[0]["registry"])


def save_reg(u, reg):
    ex("UPDATE brands SET name=?, registry=? WHERE id=?",
       ((reg.get("brand") or "").strip() or "Brand", json.dumps(reg, ensure_ascii=False), u["brand_id"]))


def admin_only(request):
    u = me(request)
    return u if u and u["role"] == "admin" else None


RULE_KINDS = [
    ("banned_phrases", "Banned phrases", "Claims creators must never make. Any match is a violation.", "e.g. guaranteed returns"),
    ("confidential", "Confidential terms", "Internal terms that must never be spoken or written. A match blocks approval and is redacted in transcripts.", "e.g. dealer commission"),
    ("allowed_contacts", "Approved contact numbers", "Other phone numbers or emails are sent to a person to check.", "e.g. 9999999999"),
    ("terminology", "Preferred wording", "Words to avoid and what to say instead, written as: avoid => prefer. A match is shown as a suggestion only. It never rejects a video.", "e.g. cheap => affordable"),
    ("avoid_tone", "Tone to avoid", "Pushy or hype words that do not fit the brand voice. A match is shown as a suggestion only.", "e.g. hurry up"),
    ("glossary", "Speech glossary", "Project names and key terms that help the speech recogniser spell things correctly. This does not create violations.", "e.g. Demo Heights"),
]


@app.get("/admin/rules", response_class=HTMLResponse)
def rules_page(request: Request, err: str = ""):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    reg = load_reg(u)
    if not reg.get("onboarding", {}).get("rules_seen"):
        reg.setdefault("onboarding", {})["rules_seen"] = True
        save_reg(u, reg)
    term = request.query_params.get("q", "").strip().lower()[:60]
    cards, dialogs = [], []
    for key, title, desc, ph in RULE_KINDS:
        items = [x for x in reg.get(key, []) if not term or term in x.lower()]
        rows = []
        for i, x in enumerate(items):
            did = f"{key}-{i}"
            rows.append([f'<span lang="mr">{e(x)}</span>', ui.badge("Active", "ok", "check-circle"),
                         f'<button type="button" class="btn secondary sm" data-dialog="ed-{did}" aria-label="Edit {e(x)}">Edit</button> '
                         f'<button type="button" class="btn secondary sm" data-dialog="rm-{did}" aria-label="Remove {e(x)}">Remove</button>'])
            dialogs.append(ui.confirm_dialog(f"rm-{did}", "Remove this rule?",
                                             f"<b>{e(x)}</b> will no longer be checked. This affects new submissions only.",
                                             "/admin/rules/remove", "Remove rule", "danger", None, {"field": key, "value": x}))
            dialogs.append(
                f'<dialog id="ed-{did}" class="dlg" aria-labelledby="ed-{did}-t"><form method="post" action="/admin/rules/edit" data-lock>'
                f'<input type="hidden" name="field" value="{key}"><input type="hidden" name="old" value="{e(x)}">'
                f'<h2 id="ed-{did}-t">Edit rule</h2><label for="ed-{did}-v">Rule</label>'
                f'<input id="ed-{did}-v" name="value" value="{e(x)}" required maxlength="200">'
                f'<p class="muted">Changing a rule affects new submissions only. Earlier decisions are not re-evaluated.</p>'
                f'<div class="dlg-actions"><button type="button" class="btn secondary" data-close>Cancel</button>'
                f'<button type="submit" class="btn">Save change</button></div></form></dialog>')
        tbl = (ui.table(["Rule", "Status", "Actions"], rows, title) if rows else
               ui.empty("No matching rules" if term else "No rules yet", "Add one below." if not term else ""))
        add = (f'<form method="post" action="/admin/rules/add" class="toolbar" style="margin:14px 0 0" data-lock>'
               f'<input type="hidden" name="field" value="{key}"><label for="add-{key}" style="position:absolute;left:-9999px">Add to {e(title)}</label>'
               f'<input id="add-{key}" name="value" required maxlength="200" placeholder="{e(ph)}" style="flex:1;min-width:200px;margin:0">'
               f'<button class="btn" type="submit">{ui.icon("plus", 16)}Add</button></form>')
        cards.append(ui.card(f"{title} ({len(reg.get(key, []))})", tbl + add, desc))
    search = (f'<div class="toolbar"><form method="get" action="/admin/rules" role="search"><label for="q" style="position:absolute;left:-9999px">Search rules</label>'
              f'<input id="q" type="search" name="q" value="{e(term)}" placeholder="Search all rules">'
              f'<button class="btn secondary" type="submit">{ui.icon("search", 16)}Search</button></form></div>')
    body = (ui.page_header("Brand Rules", f"The rules every submission for {e(reg.get('brand', ''))} is checked against. "
                           "Changes apply to new submissions only.",
                           f'<a class="btn" href="/admin/guide">{ui.icon("upload", 16)}Import from PDF</a>')
            + '<div class="stack">' + (ui.alert(e(err), "bad") if err else "") + ui.card("", search)
            + pack_card() + assets_card(reg) + "".join(cards) + "</div>" + "".join(dialogs))
    return page("Brand Rules", body, u, "rules", [("Brand Rules", None)], request)


def assets_card(reg):
    c = AST.config(reg)
    logo = ("Logo uploaded." if c.get("logo") and os.path.exists(c["logo"]) else "No logo uploaded.")
    return ui.card("Brand assets and accessibility (images and posters)",
                   f'<form method="post" action="/admin/assets" enctype="multipart/form-data" data-lock><div class="field-grid">'
                   f'<div><label for="pal">Brand colours (hex)</label><input id="pal" name="palette" value="{e(", ".join(c["palette"]))}" '
                   f'placeholder="#1F3A5F, #F2A900"></div>'
                   f'<div><label for="mc">Minimum text contrast</label><input id="mc" name="min_contrast" value="{c["min_contrast"]:g}"></div>'
                   f'<div><label for="mt">Smallest text (% of image height)</label><input id="mt" name="min_text_pct" value="{c["min_text_pct"]:g}"></div>'
                   f'<div><label for="lg">Logo (PNG or JPG)</label><input id="lg" type="file" name="logo" accept="image/png,image/jpeg"></div></div>'
                   f'<p><label><input type="checkbox" name="check_palette" value="1" {"checked" if c["check_palette"] else ""}> '
                   f'Point out large areas in non-brand colours</label><br>'
                   f'<label><input type="checkbox" name="require_logo" value="1" {"checked" if c["require_logo"] else ""}> '
                   f'Every poster must show the logo</label></p>'
                   f'<p class="hint">{logo} Contrast 4.5 is the WCAG level for normal text (large text needs 3). Fonts are not checked.</p>'
                   f'<button class="btn" type="submit">Save brand assets</button></form>')


@app.post("/admin/assets")
async def save_assets(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    f = await request.form()
    reg = load_reg(u)
    c = AST.config(reg)
    try:
        c["palette"] = AST.parse_palette(f.get("palette", ""))
        c["min_contrast"] = min(21.0, max(1.0, float(f.get("min_contrast") or 4.5)))
        c["min_text_pct"] = min(20.0, max(0.1, float(f.get("min_text_pct") or 1.2)))
    except ValueError:
        request.session["flash"] = "Contrast and text size must be numbers."
        return RedirectResponse("/admin/rules", 303)
    c["check_palette"] = bool(f.get("check_palette"))
    c["require_logo"] = bool(f.get("require_logo"))
    up = f.get("logo")
    if up is not None and getattr(up, "filename", ""):
        data = await up.read()
        import cv2
        import numpy as np
        img = cv2.imdecode(np.frombuffer(data[:3 * 1024 * 1024], np.uint8), cv2.IMREAD_COLOR) if data else None
        if img is None or len(data) > 3 * 1024 * 1024:
            request.session["flash"] = "The logo must be a PNG or JPG image under 3 MB."
            return RedirectResponse("/admin/rules", 303)
        (DATA_DIR / "assets").mkdir(parents=True, exist_ok=True)
        dst = DATA_DIR / "assets" / f"brand_{u['brand_id']}_logo.png"
        cv2.imwrite(str(dst), img)               # re-encoded, so only pixels are kept
        c["logo"] = str(dst)
    if c["require_logo"] and not c.get("logo"):
        c["require_logo"] = False
        request.session["flash"] = "Upload a logo before making it required."
    reg["assets"] = c
    save_reg(u, reg)
    log(None, u["name"], "RULE_CHANGED", "Brand assets and accessibility settings updated", u["brand_id"])
    request.session.setdefault("flash", "Brand assets saved.")
    return RedirectResponse("/admin/rules", 303)


def pack_card():
    rows = ""
    for k, lbl in packs.OPTIONS:
        if not k:
            continue
        p, n = packs.PACKS[k], packs.counts(k)
        src = "".join(f"<li><b>{e(t)}</b><br><span class='muted'>{e(d)}</span></li>" for t, d in p.get("sources", []))
        rows += (f'<details style="margin:0 0 10px"><summary><b>{e(lbl)}</b> '
                 f'<span class="muted">({n["banned"]} banned phrases, {n["terminology"]} terms, {n["tone"]} tone notes)</span></summary>'
                 f'<div style="padding:8px 0 0 14px"><p class="muted" style="margin:0 0 6px">Based on:</p><ul>{src}</ul>'
                 f'<form method="post" action="/admin/rules/pack" data-lock>'
                 f'<input type="hidden" name="pack" value="{k}"><button class="btn secondary sm" type="submit">Add {e(lbl)} rules</button></form>'
                 f'</div></details>')
    return ui.card("Starter rules by industry",
                   f'<p class="muted">Add a ready list of commonly challenged claims for your industry. Phrases already in your '
                   f'rules are skipped, and every phrase can be edited or removed below. Open an industry to see which public '
                   f'rules the list is based on. These lists are a starting point, not legal advice: some phrases may be '
                   f'allowed for your product, so review them and remove what does not apply.</p>{rows}')


@app.post("/admin/rules/pack")
def rules_pack(request: Request, pack: str = Form("")):
    u = admin_only(request)
    if not u or pack not in packs.PACKS:
        return RedirectResponse("/admin/rules", 303)
    reg = load_reg(u)
    added = packs.apply(reg, [pack])
    save_reg(u, reg)
    n = sum(added.values())
    log(None, u["name"], "PACK_APPLIED", f"{packs.PACKS[pack]['label']}: {n} rule(s) added", u["brand_id"])
    request.session["flash"] = (f"{packs.PACKS[pack]['label']}: {n} rule(s) added." if n
                                else "All of those rules were already there.")
    return RedirectResponse("/admin/rules", 303)


def rule_field(field):
    return field if field in {k for k, *_ in RULE_KINDS} else None


@app.post("/admin/rules/add")
def rule_add(request: Request, field: str = Form(""), value: str = Form("")):
    u, field, value = admin_only(request), rule_field(field), value.strip()[:200]
    if not u or not field:
        return RedirectResponse("/", 303)
    reg = load_reg(u)
    if not value:
        return RedirectResponse("/admin/rules?err=Enter+a+value+to+add.", 303)
    if value.lower() in [x.lower() for x in reg.get(field, [])]:
        return RedirectResponse("/admin/rules?err=That+rule+already+exists.", 303)
    reg.setdefault(field, []).append(value)
    save_reg(u, reg)
    log(None, u["name"], "RULE_ADDED", f"{field}: {value}", u["brand_id"])
    request.session["flash"] = "Rule added. It applies to new submissions."
    return RedirectResponse("/admin/rules", 303)


@app.post("/admin/rules/edit")
def rule_edit(request: Request, field: str = Form(""), old: str = Form(""), value: str = Form("")):
    u, field, value = admin_only(request), rule_field(field), value.strip()[:200]
    if not u or not field:
        return RedirectResponse("/", 303)
    reg = load_reg(u)
    items = reg.get(field, [])
    if old not in items or not value:
        return RedirectResponse("/admin/rules?err=Could+not+find+that+rule.", 303)
    if value.lower() != old.lower() and value.lower() in [x.lower() for x in items]:
        return RedirectResponse("/admin/rules?err=That+rule+already+exists.", 303)
    items[items.index(old)] = value
    save_reg(u, reg)
    log(None, u["name"], "RULE_CHANGED", f"{field}: {old} -> {value}", u["brand_id"])
    request.session["flash"] = "Rule updated. It applies to new submissions."
    return RedirectResponse("/admin/rules", 303)


@app.post("/admin/rules/remove")
def rule_remove(request: Request, field: str = Form(""), value: str = Form("")):
    u, field = admin_only(request), rule_field(field)
    if not u or not field:
        return RedirectResponse("/", 303)
    reg = load_reg(u)
    if value in reg.get(field, []):
        reg[field].remove(value)
        save_reg(u, reg)
        log(None, u["name"], "RULE_REMOVED", f"{field}: {value}", u["brand_id"])
        request.session["flash"] = "Rule removed."
    return RedirectResponse("/admin/rules", 303)


def facts_view(request, u, reg, err="", form=None):
    f = form or {}
    rows, dialogs = [], []
    for i, c in enumerate(reg.get("configs", [])):
        rows.append([f'<b>{e(c["name"])}</b>', f'{ui.money(c["price"])} <span class="muted">({e(fmt_money(c["price"]))})</span>',
                     ui.badge("Checked on every review", "ok", "check-circle"),
                     f'<button type="button" class="btn secondary sm" data-dialog="ec-{i}" aria-label="Edit {e(c["name"])}">Edit</button> '
                     f'<button type="button" class="btn secondary sm" data-dialog="rc-{i}" aria-label="Remove {e(c["name"])}">Remove</button>'])
        dialogs.append(ui.confirm_dialog(f"rc-{i}", "Remove this approved price?",
                                         f"<b>{e(c['name'])}</b> will no longer be used to check prices. New submissions only.",
                                         "/admin/facts/price/remove", "Remove price", "danger", None, {"name": c["name"]}))
        dialogs.append(
            f'<dialog id="ec-{i}" class="dlg" aria-labelledby="ec-{i}-t"><form method="post" action="/admin/facts/price" data-lock>'
            f'<input type="hidden" name="old" value="{e(c["name"])}"><h2 id="ec-{i}-t">Edit approved price</h2>'
            f'<label for="ec-{i}-n">Configuration</label><input id="ec-{i}-n" name="name" value="{e(c["name"])}" required maxlength="80">'
            f'<label for="ec-{i}-p">Approved price</label><input id="ec-{i}-p" name="price" value="{c["price"]}" required>'
            f'<p class="hint">Rupees, e.g. 13000000, or 1.3 Cr, or 90 lakh. The saved value is shown back in full.</p>'
            f'<div class="dlg-actions"><button type="button" class="btn secondary" data-close>Cancel</button>'
            f'<button type="submit" class="btn">Save price</button></div></form></dialog>')
    price_tbl = (ui.table(["Configuration", "Approved price", "Engine", "Actions"], rows, "Approved prices") if rows
                 else ui.empty("No approved prices", "Add at least one price. Without it, prices cannot be verified."))
    add = (f'<form method="post" action="/admin/facts/price" data-lock style="margin-top:16px"><div class="field-grid">'
           f'<div><label for="pn">Configuration <span class="req">(required)</span></label><input id="pn" name="name" required maxlength="80" '
           f'value="{e(f.get("name", ""))}" placeholder="e.g. 2 BHK"></div>'
           f'<div><label for="pp">Approved price <span class="req">(required)</span></label><input id="pp" name="price" required '
           f'value="{e(f.get("price", ""))}" placeholder="e.g. 13000000 or 1.3 Cr" aria-describedby="pph"></div></div>'
           f'<p class="hint" id="pph">Enter rupees, or write 90 lakh / 1.3 Cr. The saved value is shown back in full so you can check lakh versus crore.</p>'
           f'<button class="btn" type="submit">{ui.icon("plus", 16)}Add approved price</button></form>')
    d = reg
    details = (f'<form method="post" action="/admin/setup" data-lock><div class="field-grid">'
               f'<div><label for="brand">Brand name</label><input id="brand" name="brand" value="{e(d.get("brand", ""))}"></div>'
               f'<div><label for="project">Project name</label><input id="project" name="project" value="{e(d.get("project", ""))}"></div>'
               f'<div><label for="rera">RERA number</label><input id="rera" name="rera" value="{e(d.get("rera", ""))}"></div>'
               f'<div><label for="poss">Possession date</label><input id="poss" name="possession" value="{e(d.get("possession", ""))}"></div></div>'
               f'<button class="btn" type="submit">Save details</button></form>')
    status = ui.table(["Fact", "Stored", "Used by the compliance engine"], [
        ["Approved prices", ui.badge("Stored", "ok", "check"), ui.badge("Checked on every review", "ok", "check-circle")],
        ["RERA number", ui.badge("Stored", "ok", "check"), ui.badge("Not checked yet", "neutral", "clock")],
        ["Possession date", ui.badge("Stored", "ok", "check"), ui.badge("Not checked yet", "neutral", "clock")]], "What is checked")
    body = (ui.page_header("Approved Facts", f"The reference information that content for {e(d.get('brand', ''))} is checked against. "
                           "Changes apply to new submissions only.")
            + '<div class="stack">' + (ui.alert(e(err), "bad") if err else "")
            + ui.card("Approved prices", price_tbl + add, "Spoken and written prices are compared with these.")
            + ui.card("Brand and project details", details)
            + ui.card("What is stored and what is checked", status,
                      "RERA numbers and possession dates are kept as reference only. Verifying them is planned, not available today.")
            + "</div>" + "".join(dialogs))
    return page("Approved Facts", body, u, "facts", [("Approved Facts", None)], request)


@app.get("/admin/setup")
def setup_redirect():
    return RedirectResponse("/admin/facts", 303)


@app.post("/admin/setup")
def setup_save(request: Request, brand: str = Form(None), project: str = Form(None), rera: str = Form(None),
               possession: str = Form(None), banned: str = Form(None), confidential: str = Form(None),
               contacts: str = Form(None), glossary: str = Form(None)):
    """Update only the fields that were sent. Price configurations have their own routes."""
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    reg = load_reg(u)
    for key, val in (("brand", brand), ("project", project), ("rera", rera), ("possession", possession)):
        if val is not None:
            reg[key] = val.strip()[:120]
    for key, val in (("banned_phrases", banned), ("confidential", confidential),
                     ("allowed_contacts", contacts), ("glossary", glossary)):
        if val is not None:
            reg[key] = lines(val)
    save_reg(u, reg)
    log(None, u["name"], "FACTS_UPDATED", "Brand and project details", u["brand_id"])
    request.session["flash"] = "Saved. New submissions use these details."
    return RedirectResponse("/admin/facts", 303)


@app.post("/admin/facts/price")
def price_save(request: Request, name: str = Form(""), price: str = Form(""), old: str = Form("")):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    reg = load_reg(u)
    name = name.strip()[:80]
    val = parse_price(price) if price.strip() else None
    if not name or not val or val <= 0:
        return facts_view(request, u, reg, "Enter a configuration name and a price greater than zero, such as 13000000 or 1.3 Cr.",
                          {"name": name, "price": price})
    cfgs = reg.setdefault("configs", [])
    names = [c["name"].lower() for c in cfgs if c["name"] != old]
    if name.lower() in names:
        return facts_view(request, u, reg, f'"{name}" already has an approved price. Edit that row instead.',
                          {"name": name, "price": price})
    for c in cfgs:
        if old and c["name"] == old:
            c["name"], c["price"] = name, val
            break
    else:
        cfgs.append({"name": name, "price": val})
    save_reg(u, reg)
    log(None, u["name"], "PRICE_SAVED", f"{name} = {ui.money(val)}", u["brand_id"])
    request.session["flash"] = f"Saved {name} at {ui.money(val)}."
    return RedirectResponse("/admin/facts", 303)


@app.post("/admin/facts/price/remove")
def price_remove(request: Request, name: str = Form("")):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    reg = load_reg(u)
    if len(reg.get("configs", [])) <= 1:
        return facts_view(request, u, reg, "At least one approved price must remain, otherwise prices cannot be verified.")
    before = len(reg["configs"])
    reg["configs"] = [c for c in reg["configs"] if c["name"] != name]
    if len(reg["configs"]) != before:
        save_reg(u, reg)
        log(None, u["name"], "PRICE_REMOVED", name, u["brand_id"])
        request.session["flash"] = "Approved price removed."
    return RedirectResponse("/admin/facts", 303)


# -------------------------------------------------------- tamper check (hash)
def verify_page(request, u, result=""):
    body = (ui.page_header("Verify an Approved File",
                           "Check whether a file is exactly the one that was approved, before you post it.")
            + '<div class="cols"><div class="stack">'
            + ui.card("", f"""<form method="post" action="/verify" enctype="multipart/form-data" data-lock>
              <label for="vf">File to check</label><input id="vf" type="file" name="file" required>
              <p class="hint">Up to {MAX_MB} MB. The file is read only to compute its fingerprint and is not stored.</p>
              <p data-busy hidden class="muted" role="status">Checking...</p>
              <button class="btn" type="submit" data-busy-label="Checking...">{ui.icon("file-check", 16)}Verify File</button></form>""")
            + result + '</div><div>'
            + ui.card("How this works", """<p class="muted">Every approved file has a SHA-256 fingerprint. Even a tiny edit changes
              it completely. If the file you upload has the same fingerprint as an approved submission, it is identical to what was approved.</p>
              <p class="muted" style="margin:0">A file that does not match has simply not been approved in this form. That does not mean it is unsafe.</p>""")
            + '</div></div>')
    return page("Verify a File", body, u, "verify", [("Verify a File", None)], request)


@app.get("/verify", response_class=HTMLResponse)
def verify_form(request: Request):
    u = me(request)
    if not u:
        return RedirectResponse("/", 303)
    return verify_page(request, u)


@app.post("/verify", response_class=HTMLResponse)
def verify_check(request: Request, file: UploadFile = File(...)):
    u = me(request)
    if not u:
        return RedirectResponse("/", 303)
    h, size = hashlib.sha256(), 0
    while True:
        chunk = file.file.read(1024 * 1024)
        if not chunk:
            break
        size += len(chunk)
        if size > MAX_MB * 1024 * 1024:
            return verify_page(request, u, ui.alert(f"This file is larger than {MAX_MB} MB, so it could not be checked.",
                                                    "bad", "Could not verify."))
        h.update(chunk)
    digest = h.hexdigest()
    rows = q("SELECT id, approval_id, brand_id, user_id FROM submissions WHERE sha256=? AND approval_id IS NOT NULL AND source_url IS NULL",
             (digest,))
    fp = f'<dl class="kv" style="margin-top:12px"><dt>File fingerprint (SHA-256)</dt><dd class="mono">{digest}</dd></dl>'
    if rows:
        r = rows[0]
        link = ""
        if can_view(u, {"brand_id": r["brand_id"], "user_id": r["user_id"]}):
            link = f' <a href="/s/{r["id"]}">View submission #{r["id"]}</a>.'
        result = ui.card("", ui.alert(f'This file is identical to an approved submission. Approval ID: <b>{e(r["approval_id"])}</b>.{link}',
                                      "ok", "Match.") + fp)
    else:
        result = ui.card("", ui.alert("This file does not match any approved record. It may have been edited after approval, "
                                      "or it was never approved in this form.", "warn", "No match.") + fp)
    return verify_page(request, u, result)

    

# ============================================================ projects, facts, access (admin)
def own_project(u, pid):
    p = RE.get_project(pid)
    return p if p and p["brand_id"] == u["brand_id"] else None


def source_id_for(u, title):
    title = (title or "").strip()
    if not title:
        return None
    hit = q("SELECT id FROM sources WHERE brand_id=? AND title=?", (u["brand_id"], title[:200]))
    return hit[0]["id"] if hit else RE.add_source(u["brand_id"], title, "document", "official", u["name"])


def parse_fact_value(key, raw):
    raw = (raw or "").strip()
    if not raw:
        raise ValueError("Enter a value.")
    if key in ("price", "booking_amount"):
        v = parse_price(raw)
        if not v or v <= 0:
            raise ValueError("Enter a price greater than zero, for example 13000000 or 1.3 Cr.")
        return float(v), None
    if key in RE.NUMERIC_FIELDS or key.startswith("landmark:"):
        try:
            v = float(raw.replace(",", ""))
        except ValueError:
            raise ValueError("Enter a number.")
        if v <= 0:
            raise ValueError("The number must be greater than zero.")
        return v, None
    if key == "possession":
        if not re.fullmatch(r"20\d\d-(0[1-9]|1[0-2])", raw):
            raise ValueError("Write possession as year-month, for example 2028-12.")
        return None, raw
    if key == "ready_to_move":
        if raw.lower() not in ("yes", "no"):
            raise ValueError("Choose yes or no.")
        return None, raw.lower()
    return None, raw[:200]


def admin_back(request, pid, msg="", err=""):
    if msg:
        request.session["flash"] = msg
    if err:
        request.session["flash"] = err
    return RedirectResponse(f"/admin/projects/{pid}", 303)


@app.get("/admin/facts")
def facts_redirect(request: Request):
    return RedirectResponse("/admin/projects", 303)


@app.get("/admin/projects", response_class=HTMLResponse)
def projects_page(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    rows = []
    for p in RE.brand_projects(u["brand_id"], active_only=False):
        n_units = len(RE.unit_types(p["id"]))
        rera = (e(p["rera_no"]) + "<br>" + ui.badge({"authoritative": "Confirmed with official register", "format": "Format only",
                "registry_match": "Registry match"}.get(p["rera_level"], "Not recorded"),
                "ok" if p["rera_level"] == "authoritative" else "neutral")) if p["rera_no"] else '<span class="muted">Not recorded</span>'
        rows.append([f'<a href="/admin/projects/{p["id"]}"><b>{e(p["official_name"])}</b></a>',
                     e(", ".join(x for x in (p["locality"], p["city"]) if x) or "-"), rera, str(n_units),
                     f'<a class="btn secondary sm" href="/admin/projects/{p["id"]}">Open</a>'])
    add = (f'<form method="post" action="/admin/projects/add" data-lock><div class="field-grid">'
           f'<div><label for="pn">Official project name</label><input id="pn" name="official_name" required maxlength="120"></div>'
           f'<div><label for="ps">Short name</label><input id="ps" name="short_name" maxlength="60"></div>'
           f'<div><label for="pc">City</label><input id="pc" name="city" maxlength="60"></div>'
           f'<div><label for="pl">Locality</label><input id="pl" name="locality" maxlength="60"></div>'
           f'<div><label for="pr">RERA number</label><input id="pr" name="rera_no" maxlength="20" placeholder="P followed by 11 digits"></div>'
           f'<div><label for="pa">Other names (one per line)</label><textarea id="pa" name="aliases" rows="3" '
           f'placeholder="Hindi, Marathi and short spellings"></textarea></div></div>'
           f'<button class="btn" type="submit">{ui.icon("plus", 16)}Add project</button></form>')
    tools = ui.card("Set up faster", '<p class="muted">Fill our spreadsheet template, or upload a brochure or price list and let '
                    'the AI propose the products and facts. Nothing is saved until you check it and confirm.</p>'
                    '<a class="btn" href="/admin/import">Import spreadsheet or brochure</a> '
                    '<a class="btn secondary" href="/admin/fact-types">Fact types</a> '
                    '<form method="post" action="/admin/sample" style="display:inline" data-lock>'
                    '<button class="btn secondary" type="submit">Load sample data</button></form>')
    if load_reg(u).get("domain", "property") != "property":
        prow = [[r_[0], r_[3], r_[4]] for r_ in rows]
        add_g = (f'<form method="post" action="/admin/projects/add" data-lock><div class="field-grid">'
                 f'<div><label for="pn">Product name</label><input id="pn" name="official_name" required maxlength="120"></div>'
                 f'<div><label for="pa">Other names (one per line)</label><textarea id="pa" name="aliases" rows="3" '
                 f'placeholder="Short names, Hindi or Marathi spellings"></textarea></div></div>'
                 f'<button class="btn" type="submit">{ui.icon("plus", 16)}Add product</button></form>')
        body = (ui.page_header("Products and Facts", "Every product, its variants and the approved facts that content is checked against.")
                + '<div class="stack">' + tools
                + ui.card("Products", ui.table(["Product", "Variants", ""], prow, "Products") if prow else ui.empty("No products yet"))
                + ui.card("Add a product", add_g, "Other names help BrandGuard recognise the product in speech.") + "</div>")
        return page("Products and Facts", body, u, "facts", [("Products and Facts", None)], request)
    body = (ui.page_header("Projects and Facts", "Every project, its property types and the approved facts that content is checked against.")
            + '<div class="stack">' + tools
            + ui.card("Projects", ui.table(["Project", "Location", "RERA", "Property types", ""], rows, "Projects") if rows else ui.empty("No projects yet"))
            + ui.card("Add a project", add, "Aliases help BrandGuard recognise the project in Hindi, Marathi or short forms.") + "</div>")
    return page("Projects and Facts", body, u, "facts", [("Projects and Facts", None)], request)


@app.post("/admin/projects/add")
def project_add(request: Request, official_name: str = Form(""), short_name: str = Form(""), city: str = Form(""),
                locality: str = Form(""), rera_no: str = Form(""), aliases: str = Form("")):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    if not official_name.strip():
        request.session["flash"] = "Enter the project name."
        return RedirectResponse("/admin/projects", 303)
    pid = RE.add_project(u["brand_id"], official_name, short_name, city, locality, rera_no, lines(aliases), u["name"])
    log(None, u["name"], "PROJECT_ADDED", official_name.strip(), u["brand_id"])
    request.session["flash"] = "Project added. Now add its property types and facts."
    return RedirectResponse(f"/admin/projects/{pid}", 303)


def fact_rows_html(pid, unit_id=None):
    fs = RE._q("SELECT * FROM facts WHERE project_id=? AND key NOT LIKE 'g:%' AND "
               + ("unit_type_id=?" if unit_id else "unit_type_id IS NULL") + " ORDER BY key, id",
               (pid, unit_id) if unit_id else (pid,))
    out = []
    for f in fs:
        cur = RE._is_current(f)
        state = ui.badge("Current", "ok", "check") if cur else ui.badge("Not current", "neutral", "clock")
        src = q("SELECT title FROM sources WHERE id=?", (f["source_id"],)) if f["source_id"] else []
        end = ""
        if cur:
            end = (f'<form method="post" action="/admin/projects/{pid}/fact/end" style="display:inline"><input type="hidden" name="fact_id" value="{f["id"]}">'
                   f'<button class="btn secondary sm" type="submit">Withdraw</button></form>')
        label = (f["key"][2:].capitalize() if f["key"].startswith("g:") else
                 RE.FACT_LABEL.get(f["key"], f["key"].replace("offer:", "Offer: ").replace("landmark:", "Place: ")))
        out.append([e(label),
                    e(RE.fact_display(f)), f'{e(f["valid_from"] or "")} to {e(f["valid_to"] or "open")}',
                    e(src[0]["title"] if src else "-"), state, end])
    return out


@app.get("/admin/projects/{pid}", response_class=HTMLResponse)
def project_detail(request: Request, pid: int):
    u = admin_only(request)
    p = own_project(u, pid) if u else None
    if not p:
        return RedirectResponse("/admin/projects", 303)
    heads = ["Fact", "Value", "Valid", "Source", "State", ""]
    prop = load_reg(u).get("domain", "property") == "property"
    units_html = ""
    for ut in RE.unit_types(pid):
        cat = RE.CATEGORIES[ut["category"]][0]
        rows = fact_rows_html(pid, ut["id"])
        units_html += ui.card(f'{ut["label"]} ({cat})', ui.table(heads, rows, ut["label"]) if rows else ui.empty("No facts yet"),
                              actions=f'<a href="#fact-form">Add fact</a>')
    cat_opts = "".join(f'<option value="{k}">{e(v[0])}</option>' for k, v in RE.CATEGORIES.items())
    field_inputs = "".join(f'<div><label for="uf-{k}">{e(RE.FIELD_LABEL[k])}</label><input id="uf-{k}" name="{k}"></div>'
                           for k in ("bhk", "carpet_sqft", "builtup_sqft", "plot_sqft", "bathrooms", "balconies", "parking", "floor_range", "facing", "dimensions", "permitted_use"))
    unit_form = (f'<form method="post" action="/admin/projects/{pid}/unit" data-lock><div class="field-grid">'
                 f'<div><label for="uc">Property category</label><select id="uc" name="category">{cat_opts}</select></div>'
                 f'<div><label for="ul">Name</label><input id="ul" name="label" required placeholder="e.g. 2 BHK, Corner plot A"></div>'
                 f'<div><label for="uprice">Price (required)</label><input id="uprice" name="price" required placeholder="e.g. 1.3 Cr"></div>'
                 f'{field_inputs}'
                 f'<div><label for="usrc">Source document</label><input id="usrc" name="source" placeholder="e.g. Approved price sheet Oct 2026"></div></div>'
                 f'<p class="hint">Fill only what applies. Fields that do not apply to the category (for example BHK for a plot) are ignored.</p>'
                 f'<button class="btn" type="submit">{ui.icon("plus", 16)}Add property type</button></form>')
    unit_opts = '<option value="">Whole project</option>' + "".join(
        f'<option value="{ut["id"]}">{e(ut["label"])}</option>' for ut in RE.unit_types(pid))
    key_opts = "".join(f'<option value="{k}">{v}</option>' for k, v in (
        ("price", "Price"), ("booking_amount", "Booking amount"), ("carpet_sqft", "Carpet area (sq ft)"),
        ("builtup_sqft", "Built-up area (sq ft)"), ("plot_sqft", "Plot area (sq ft)"), ("bathrooms", "Bathrooms"),
        ("balconies", "Balconies"), ("parking", "Parking"), ("possession", "Possession (YYYY-MM)"),
        ("ready_to_move", "Ready to move (yes or no)")))
    offer_opts = "".join(f'<option value="offer:{k}">{v[0]}</option>' for k, v in RE.OFFER_TAGS.items())
    fact_form = (f'<form id="fact-form" method="post" action="/admin/projects/{pid}/fact" data-lock><div class="field-grid">'
                 f'<div><label for="fk">Fact</label><select id="fk" name="key">{key_opts}{offer_opts}</select></div>'
                 f'<div><label for="fu">Applies to</label><select id="fu" name="unit_type_id">{unit_opts}</select></div>'
                 f'<div><label for="fv">Value or terms</label><input id="fv" name="value" required></div>'
                 f'<div><label for="ffrom">Valid from</label><input id="ffrom" name="valid_from" type="date"></div>'
                 f'<div><label for="fto">Valid until (offers)</label><input id="fto" name="valid_to" type="date"></div>'
                 f'<div><label for="fs">Source document</label><input id="fs" name="source" required placeholder="Where this number comes from"></div>'
                 f'<div><label for="fr">Reason for change</label><input id="fr" name="reason" placeholder="e.g. Revised price sheet"></div></div>'
                 f'<p class="hint">Saving a new value closes the old one. Earlier approvals that used the old value are flagged for re-check, '
                 f'not changed.</p><button class="btn" type="submit">Save fact</button></form>')
    lm_form = (f'<form method="post" action="/admin/projects/{pid}/landmark" data-lock><div class="field-grid">'
               f'<div><label for="lmn">Place</label><input id="lmn" name="place" required placeholder="e.g. Wakad bridge"></div>'
               f'<div><label for="lmm">Travel time (minutes)</label><input id="lmm" name="minutes" required></div>'
               f'<div><label for="lmh">By</label><input id="lmh" name="how" placeholder="car, walking..."></div>'
               f'<div><label for="lms">Source document</label><input id="lms" name="source" required></div></div>'
               f'<button class="btn" type="submit">Save travel time</button></form>')
    proj_facts = fact_rows_html(pid, None)
    lm = RE._q("SELECT * FROM facts WHERE project_id=? AND key LIKE 'landmark:%' ORDER BY id", (pid,))
    lm_rows = [[e(f["key"].split(":", 1)[1].title()), e(f"{f['value_num']:g} min {f['unit'] or ''}"),
                f'{e(f["valid_from"] or "")} to {e(f["valid_to"] or "open")}', ui.badge("Current", "ok") if RE._is_current(f) else ui.badge("Not current", "neutral")]
               for f in lm]
    rera_info = (f'<dl class="kv"><dt>RERA number</dt><dd>{e(p["rera_no"] or "Not recorded")}</dd>'
                 f'<dt>Verification level</dt><dd>{e({"authoritative": "Confirmed with the official register", "format": "Format is valid. Not checked against any register.", "registry_match": "Matches the BrandGuard registry only"}.get(p["rera_level"], "Not recorded"))}</dd>'
                 + (f'<dt>Checked</dt><dd>{e(p["rera_checked_at"] or "")} by {e(p["rera_checked_by"] or "")} against {e(p["rera_source"] or "")}</dd>' if p["rera_level"] == "authoritative" else "")
                 + '</dl>')
    rera_form = (f'<form method="post" action="/admin/projects/{pid}/rera" data-lock><div class="field-grid">'
                 f'<div><label for="rn">RERA number</label><input id="rn" name="rera_no" value="{e(p["rera_no"] or "")}"></div>'
                 f'<div><label for="rs">Official register checked (optional)</label><input id="rs" name="source" placeholder="e.g. MahaRERA website, checked on screen"></div></div>'
                 f'<p class="hint">Fill the second box only if you have actually looked the number up on the official register. '
                 f'Otherwise it stays a format or registry match, never a confirmed registration.</p>'
                 f'<button class="btn" type="submit">Save RERA details</button></form>')
    names = RE.names_of(p)
    alias_form = (f'<form method="post" action="/admin/projects/{pid}/alias" data-lock><label for="al">Add another name</label>'
                  f'<input id="al" name="alias" required maxlength="80" placeholder="Hindi, Marathi or short spelling"><br><br>'
                  f'<button class="btn" type="submit">Add name</button></form>')
    gcard = product_facts_card(u, p)
    if prop:
        body = (ui.page_header(p["official_name"], e(", ".join(x for x in (p["locality"], p["city"]) if x)),
                               '<a class="btn secondary" href="/admin/projects">All projects</a>')
                + '<div class="stack">'
                + ui.card("Names BrandGuard listens for", "<p>" + e(" | ".join(names)) + "</p>" + alias_form)
                + ui.card("RERA", rera_info + rera_form)
                + ui.card("Project-wide facts", ui.table(heads, proj_facts, "Project facts") if proj_facts else ui.empty("None yet"))
                + units_html
                + ui.card("Add a property type", unit_form, "Apartment, penthouse, villa, bungalow, row house or plot.")
                + ui.card("Add or change a fact", fact_form, "Possession, ready to move, booking amount, offers and unit facts.")
                + gcard
                + ui.card("Travel times to landmarks", (ui.table(["Place", "Time", "Valid", "State"], lm_rows, "Landmarks") if lm_rows else "") + lm_form)
                + "</div>")
    else:
        variant_form = (f'<form method="post" action="/admin/projects/{pid}/unit" data-lock><input type="hidden" name="category" value="product">'
                        f'<div class="field-grid"><div><label for="vl">Variant name</label><input id="vl" name="label" required '
                        f'placeholder="e.g. VXi, 500 g pack, Gold plan"></div>'
                        f'<div><label for="vp">Price (optional)</label><input id="vp" name="price" placeholder="e.g. 9,470 or 8.5 lakh"></div>'
                        f'<div><label for="vs">Source document</label><input id="vs" name="source" placeholder="e.g. Price list Nov 2026"></div></div>'
                        f'<button class="btn" type="submit">{ui.icon("plus", 16)}Add variant</button></form>')
        body = (ui.page_header(p["official_name"], "Product, its variants and the approved facts content is checked against.",
                               '<a class="btn secondary" href="/admin/projects">All products</a>')
                + '<div class="stack">'
                + ui.card("Names BrandGuard listens for", "<p>" + e(" | ".join(names)) + "</p>" + alias_form)
                + gcard
                + (ui.card("Other product-wide facts", ui.table(heads, proj_facts, "Product facts")) if proj_facts else "")
                + units_html
                + ui.card("Add a variant", variant_form, "A model, plan, size or pack. Give each one its own price if it has one.")
                + "</div>")
    crumb = "Projects and Facts" if prop else "Products and Facts"
    return page(p["official_name"], body, u, "facts", [(crumb, "/admin/projects"), (p["official_name"], None)], request)


def product_facts_card(u, p):
    defs = GEN.list_defs(u["brand_id"])
    if not defs:
        return ui.card("Your own fact types", '<p class="muted">No fact types yet. Define what you sell on (discount, warranty, '
                       'fees, mileage, interest rate, delivery date...) once, then give each product its approved values.</p>'
                       '<a class="btn" href="/admin/fact-types">Define fact types</a>')
    def_opts = "".join(f'<option value="{d["id"]}">{e(d["label"])} ({e(GEN.TYPES[d["ftype"]])}'
                       f'{(", " + e(d["units"][0])) if d["units"] else ""})</option>' for d in defs)
    unit_opts = '<option value="">Whole product</option>' + "".join(
        f'<option value="{ut["id"]}">{e(ut["label"])}</option>' for ut in RE.unit_types(p["id"]))
    rows = []
    for d in defs:
        for f in RE._q("SELECT * FROM facts WHERE project_id=? AND key=? ORDER BY id", (p["id"], d["key"])):
            if not RE._is_current(f):
                continue
            ut = RE._q("SELECT label FROM unit_types WHERE id=?", (f["unit_type_id"],)) if f["unit_type_id"] else []
            rows.append([e(d["label"]), e(RE.fact_display(f)), e(ut[0]["label"] if ut else "Whole product"),
                         e(GEN.DIRECTIONS[d["direction"]]), f'{e(f["valid_from"] or "")} to {e(f["valid_to"] or "open")}',
                         f'<form method="post" action="/admin/projects/{p["id"]}/fact/end" style="display:inline">'
                         f'<input type="hidden" name="fact_id" value="{f["id"]}"><button class="btn secondary sm" type="submit">Withdraw</button></form>'])
    form = (f'<form method="post" action="/admin/projects/{p["id"]}/gfact" data-lock><div class="field-grid">'
            f'<div><label for="gd">Fact</label><select id="gd" name="def_id">{def_opts}</select></div>'
            f'<div><label for="gu">Applies to</label><select id="gu" name="unit_type_id">{unit_opts}</select></div>'
            f'<div><label for="gv">Approved value</label><input id="gv" name="value" required placeholder="30, 999, 22, 2027-03, yes"></div>'
            f'<div><label for="gt">Valid until (optional)</label><input id="gt" name="valid_to" type="date"></div>'
            f'<div><label for="gs">Source document</label><input id="gs" name="source" required placeholder="Where this value comes from"></div></div>'
            f'<button class="btn" type="submit">Save value</button> <a class="btn secondary" href="/admin/fact-types">Manage fact types</a></form>')
    return ui.card("Your own fact types", (ui.table(["Fact", "Approved", "Applies to", "Rule", "Valid", ""], rows, "Fact values")
                                          if rows else '<p class="muted">No approved values yet.</p>') + form,
                   "Claims about these are found in speech, captions, screen text and documents, and compared with these values.")


@app.post("/admin/projects/{pid}/gfact")
def project_gfact(request: Request, pid: int, def_id: int = Form(0), unit_type_id: str = Form(""), value: str = Form(""),
                  valid_to: str = Form(""), source: str = Form("")):
    u = admin_only(request)
    p = own_project(u, pid) if u else None
    if not p:
        return RedirectResponse("/admin/projects", 303)
    try:
        save_generic_fact(u, p, def_id, unit_type_id, value, valid_to, source)
    except ValueError as ex_:
        return admin_back(request, pid, err=str(ex_))
    return admin_back(request, pid, "Saved. Earlier approvals that relied on the old value have been flagged for re-check.")


def save_generic_fact(u, p, def_id, unit_type_id, value, valid_to, source, reason="Set by builder"):
    d = GEN.get_def_by_id(u["brand_id"], def_id)
    if not d:
        raise ValueError("Choose a fact type.")
    if not (source or "").strip():
        raise ValueError("Name the source document for this value.")
    ut_id = int(unit_type_id) if str(unit_type_id).isdigit() else None
    if ut_id and not RE._q("SELECT 1 FROM unit_types WHERE id=? AND project_id=?", (ut_id, p["id"])):
        raise ValueError("That variant does not belong to this product.")
    num, txt = GEN.parse_value(d["ftype"], value)
    if valid_to and valid_to < RE.today():
        raise ValueError("The end date is in the past.")
    unit = {"percent": "%", "money": "INR"}.get(d["ftype"], d["units"][0] if d["units"] else "")
    RE.set_fact(p["id"], d["key"], num, txt, ut_id, unit, source_id_for(u, source), None, valid_to or None, u["name"], reason)
    log(None, u["name"], "FACT_SET", f'{p["official_name"]}: {d["label"]} = {txt if txt is not None else f"{num:g}"}', u["brand_id"])


# ============================================================ import and sample data
SAMPLE_GENERAL = {
    "products": [{"name": "Zippy Hatchback", "aliases": ["Zippy", "ज़िप्पी", "झिप्पी"],
                  "variants": [{"name": "LXi", "price": "6.2 lakh"}, {"name": "VXi", "price": "7.1 lakh"}]},
                 {"name": "Easy Car Loan", "aliases": ["car loan", "कार लोन"], "variants": []}],
    "fact_types": [{"label": "Discount", "ftype": "percent", "unit": "", "aliases": "off, सूट, सवलत, डिस्काउंट", "direction": "max"},
                   {"label": "Mileage", "ftype": "number", "unit": "kmpl, km/l, km per litre", "aliases": "average, मायलेज", "direction": "max"},
                   {"label": "Warranty", "ftype": "number", "unit": "years, year, yrs, वर्ष, साल", "aliases": "वॉरंटी", "direction": "max"},
                   {"label": "Free insurance", "ftype": "yes_no", "unit": "", "aliases": "insurance free, मोफत विमा", "direction": "exact"},
                   {"label": "Interest rate", "ftype": "percent", "unit": "", "aliases": "rate of interest, ब्याज दर, व्याजदर", "direction": "min"},
                   {"label": "Processing fee", "ftype": "money", "unit": "", "aliases": "processing charges, प्रोसेसिंग फी", "direction": "exact"},
                   {"label": "Loan tenure", "ftype": "number", "unit": "years, year, वर्ष, साल", "aliases": "tenure", "direction": "max"}],
    "values": [{"product": "Zippy Hatchback", "variant": "", "fact": "Discount", "value": "5", "valid_to": ""},
               {"product": "Zippy Hatchback", "variant": "", "fact": "Mileage", "value": "22", "valid_to": ""},
               {"product": "Zippy Hatchback", "variant": "", "fact": "Warranty", "value": "3", "valid_to": ""},
               {"product": "Zippy Hatchback", "variant": "", "fact": "Free insurance", "value": "no", "valid_to": ""},
               {"product": "Easy Car Loan", "variant": "", "fact": "Interest rate", "value": "8.75", "valid_to": ""},
               {"product": "Easy Car Loan", "variant": "", "fact": "Processing fee", "value": "999", "valid_to": ""},
               {"product": "Easy Car Loan", "variant": "", "fact": "Loan tenure", "value": "7", "valid_to": ""}],
}
SAMPLE_PROPERTY = {
    "products": [{"name": "Sample Residency", "aliases": ["Sample Residency Baner", "सॅम्पल रेसिडेन्सी"],
                  "variants": [{"name": "1 BHK", "price": "65 lakh"}, {"name": "2 BHK", "price": "95 lakh"}]}],
    "fact_types": [{"label": "Discount", "ftype": "percent", "unit": "", "aliases": "off, सूट, सवलत", "direction": "max"}],
    "values": [{"product": "Sample Residency", "variant": "", "fact": "Discount", "value": "2", "valid_to": ""}],
}


def apply_import(u, rows, source_title):
    """Save the ticked rows of a checked proposal. Returns a summary string."""
    prop_domain = load_reg(u).get("domain", "property") == "property"
    src = source_id_for(u, source_title)
    made = {"fact types": 0, "products": 0, "variants": 0, "values": 0}
    for r in rows:
        if r["kind"] == "fact_type":
            d = r["data"]
            GEN.add_def(u["brand_id"], d["label"], d["ftype"], d.get("unit", ""), d.get("aliases", ""), d.get("direction", "exact"))
            made["fact types"] += 1
    projects = {p["official_name"].lower(): p for p in RE.brand_projects(u["brand_id"], active_only=False)}

    def product(name, aliases=()):
        p = projects.get(name.lower())
        if not p:
            pid = RE.add_project(u["brand_id"], name, "", "", "", "", list(aliases), u["name"])
            p = RE.get_project(pid)
            projects[name.lower()] = p
            made["products"] += 1
            log(None, u["name"], "PROJECT_ADDED", name, u["brand_id"])
        else:
            have = {a.lower() for a in RE.names_of(p)}
            for a in aliases:
                if a.lower() not in have:
                    RE._ex("INSERT INTO project_aliases (project_id, alias) VALUES (?,?)", (p["id"], a))
        return p
    for r in rows:
        if r["kind"] == "product":
            product(r["data"]["name"], r["data"].get("aliases", []))
    for r in rows:
        if r["kind"] != "variant":
            continue
        v = r["data"]
        p = product(v["product"])
        ut = next((x for x in RE.unit_types(p["id"]) if x["label"].lower() == v["name"].lower()), None)
        if not ut:
            bhk = RE._bhk_in(v["name"]) if prop_domain else None
            uid = RE.add_unit_type(p["id"], "apartment" if bhk else "product", v["name"], int(bhk) if bhk else None)
            made["variants"] += 1
        else:
            uid = ut["id"]
        if v.get("price"):
            RE.set_fact(p["id"], "price", float(parse_price(v["price"])), None, uid, "INR", src, None, None, u["name"], "Imported")
    for r in rows:
        if r["kind"] != "value":
            continue
        v = r["data"]
        p = product(v["product"])
        d = GEN.get_def(u["brand_id"], v["fact"])
        if not d:
            continue
        uid = None
        if v.get("variant"):
            ut = next((x for x in RE.unit_types(p["id"]) if x["label"].lower() == v["variant"].lower()), None)
            uid = ut["id"] if ut else None
        num, txt = GEN.parse_value(d["ftype"], v["value"])
        unit = {"percent": "%", "money": "INR"}.get(d["ftype"], d["units"][0] if d["units"] else "")
        RE.set_fact(p["id"], d["key"], num, txt, uid, unit, src, None, GEN.parse_date(v.get("valid_to", "")) or None,
                    u["name"], "Imported")
        made["values"] += 1
    for r in rows:
        if r["kind"] != "prop_fact":
            continue
        v = r["data"]
        p = product(v["product"])
        if v["key"] == "rera_no":
            RE._ex("UPDATE projects SET rera_no=?, rera_level=? WHERE id=?",
                   (v["value"], "format" if RE.rera_format_ok(v["value"]) else "none", p["id"]))
            made["values"] += 1
            continue
        uid = None
        if v.get("variant"):
            ut = next((x for x in RE.unit_types(p["id"]) if x["label"].lower() == v["variant"].lower()), None)
            uid = ut["id"] if ut else None
        if v["key"] in IMP.UNIT_LEVEL and not uid:
            continue                       # the unit type row was not ticked, so there is nothing to attach this to
        num, txt = parse_fact_value(v["key"], v["value"])
        unit = "INR" if v["key"] == "booking_amount" else ("sq ft" if v["key"].endswith("sqft") else "")
        RE.set_fact(p["id"], v["key"], num, txt, uid if v["key"] in IMP.UNIT_LEVEL else None, unit, src, None, None,
                    u["name"], "Imported")
        made["values"] += 1
    return ", ".join(f"{n} {k}" for k, n in made.items() if n) or "nothing new"


@app.get("/admin/import", response_class=HTMLResponse)
def import_page(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    domain = "property" if load_reg(u).get("domain", "property") == "property" else "general"
    thing = "project" if domain == "property" else "product"
    sheet = (f'<ol><li><a class="btn secondary sm" href="/admin/import/template">Download the template (.xlsx)</a></li>'
             f'<li>Fill one row per {thing} (or per variant). The columns are explained below.</li><li>Upload it here.</li></ol>'
             f'<form method="post" action="/admin/import/sheet" enctype="multipart/form-data" data-lock>'
             f'<input type="file" name="file" accept=".xlsx" required> <button class="btn" type="submit">Read spreadsheet</button></form>')
    guide_rows = [[f'<b>{e(c)}</b>', e(IMP.COLUMN_HELP.get(c, ("", ""))[0]), f'<span class="mono">{e(IMP.COLUMN_HELP.get(c, ("", ""))[1])}</span>']
                  for c in IMP.COLUMNS[domain]]
    guide = (ui.table(["Column", "What to write", "Example"], guide_rows, "Template columns")
             + f'<h3 class="sec-title">Rules</h3><ul class="muted" style="padding-left:20px">'
             f'<li>Leave a cell empty if it does not apply. Empty cells are skipped, never saved as zero.</li>'
             f'<li>For a {thing} with variants, write one row per variant and repeat the {thing} name.</li>'
             + (f'<li>Fill once-per-project columns (other names, booking amount, possession, RERA) on any one of its rows.</li>'
                if domain == "property" else f'<li>Fill "Other names" once per product, on any one of its rows.</li>') +
             f'<li>A value written on a variant row applies to that variant only. For a value that applies to the whole '
             f'{thing}, add a row with the variant column empty, or repeat the value on every row.</li>'
             f'<li>Dates can be 2027-03-15, 15 March 2027 or March 2027.</li></ul>'
             f'<h3 class="sec-title">Add your own columns</h3><p class="muted">Anything else your content talks about can be a new '
             f'column. Put the unit in brackets and BrandGuard works out the type from it and from the values:</p>'
             + ui.table(["Column you add", "BrandGuard treats it as", "Rule"], [
                 ["Mileage (kmpl)", "A number in kmpl", "Must not be more than this"],
                 ["Interest rate (%)", "A percentage", "Must not be less than this"],
                 ["Processing fee (₹)", "Money", "Must be exactly this"],
                 ["Free insurance", "Yes or no (when the values are yes / no)", "Must match"],
                 ["Colour options", "Text", "One of the written values must be used"]], "Own columns")
             + '<p class="hint">You can change the type and rule of any column later on the Fact types page.</p>')
    body_extra = ui.card("How to fill the spreadsheet", guide)
    ai_on = bool(os.environ.get("GROQ_API_KEY"))
    bro = (f'<p class="muted">PDF, Word, text or a photo of a price list. The AI proposes products and facts; every number '
           f'is checked against the document, and you confirm row by row.</p>'
           + ('' if ai_on else ui.alert("AI is not set up on this server, so brochure import is off. Use the spreadsheet.", "warn"))
           + f'<form method="post" action="/admin/import/brochure" enctype="multipart/form-data" data-lock>'
           f'<input type="file" name="file" accept=".pdf,.docx,.txt,.png,.jpg,.jpeg" required {"" if ai_on else "disabled"}> '
           f'<button class="btn" type="submit" {"" if ai_on else "disabled"}>Read brochure</button></form>')
    body = (ui.page_header("Import products and facts", "Nothing is saved until you check the rows and confirm.",
                           '<a class="btn secondary" href="/admin/projects">Products and facts</a>')
            + '<div class="stack"><div class="cols-2">' + ui.card("Spreadsheet (no AI)", sheet) + ui.card("Brochure or price list (AI)", bro)
            + "</div>" + body_extra + "</div>")
    return page("Import", body, u, "facts", [("Products and Facts", "/admin/projects"), ("Import", None)], request)


@app.get("/admin/import/template")
def import_template(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    domain = "property" if load_reg(u).get("domain", "property") == "property" else "general"
    name = "brandguard-projects-template.xlsx" if domain == "property" else "brandguard-products-template.xlsx"
    return Response(IMP.template_xlsx(domain), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _read_small(up, limit_mb=10):
    data = up.file.read(limit_mb * 1024 * 1024 + 1)
    if len(data) > limit_mb * 1024 * 1024:
        raise ValueError(f"The file is larger than {limit_mb} MB.")
    return data


@app.post("/admin/import/sheet")
def import_sheet(request: Request, file: UploadFile = File(...)):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    try:
        if not (file.filename or "").lower().endswith(".xlsx"):
            raise ValueError("Upload the filled template as an .xlsx file.")
        prop = IMP.parse_xlsx(_read_small(file))
    except ValueError as ex_:
        request.session["flash"] = str(ex_)
        return RedirectResponse("/admin/import", 303)
    rows = IMP.check(prop, None, parse_price, GEN.list_defs(u["brand_id"]), parse_fact_value)
    iid = IMP.save(u["brand_id"], "spreadsheet", file.filename or "Spreadsheet", {"rows": rows}, u["name"])
    return RedirectResponse(f"/admin/import/{iid}", 303)


@app.post("/admin/import/brochure")
def import_brochure(request: Request, file: UploadFile = File(...)):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    try:
        text = IMP.extract_text(file.filename, _read_small(file))
        if len(text.strip()) < 20:
            raise ValueError("Almost no text was found in this file.")
        prop = IMP.ai_propose(text)
    except ValueError as ex_:
        request.session["flash"] = str(ex_)
        return RedirectResponse("/admin/import", 303)
    except Exception as ex_:
        request.session["flash"] = (str(ex_) if "AI is not set up" in str(ex_)
                                    else "The AI could not read this document. Try again, or use the spreadsheet.")
        return RedirectResponse("/admin/import", 303)
    rows = IMP.check(prop, text, parse_price, GEN.list_defs(u["brand_id"]), parse_fact_value)
    iid = IMP.save(u["brand_id"], "brochure", file.filename or "Brochure", {"rows": rows}, u["name"])
    return RedirectResponse(f"/admin/import/{iid}", 303)


@app.get("/admin/import/{iid}", response_class=HTMLResponse)
def import_preview(request: Request, iid: int):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    meta, data = IMP.load(u["brand_id"], iid)
    if not meta:
        return RedirectResponse("/admin/import", 303)
    rows = data["rows"]
    kind_label = {"fact_type": "Fact type", "product": "Product", "variant": "Variant and price", "value": "Approved value",
                  "prop_fact": "Property fact"}
    trs = []
    for r in rows:
        box = (f'<input type="checkbox" name="row" value="{r["id"]}" {"checked" if r["ok"] else ""} '
               f'{"" if r["ok"] or "not found in the document" in r["problem"] else "disabled"} aria-label="Include">')
        trs.append([box, e(kind_label[r["kind"]]), e(r["text"]),
                    ui.badge("Ready", "ok", "check") if r["ok"] else f'{ui.badge("Check", "warn", "alert")} <span class="muted">{e(r["problem"])}</span>'])
    done = meta["applied_at"]
    by = "the AI read the document" if meta["kind"] == "brochure" else "read from your spreadsheet"
    form = (f'<form method="post" action="/admin/import/{iid}/apply" data-lock>'
            + ui.table(["", "What", "Row", "Check"], trs, "Proposed rows")
            + f'<div class="field-grid" style="margin-top:12px"><div><label for="st">Source document name</label>'
            f'<input id="st" name="source" required value="{e(meta["title"])}"></div></div>'
            f'<p class="hint">Rows marked Check are unticked. Values the AI proposed that are not in the document are unticked: tick them '
            f'only if you know they are right. Rows with errors cannot be imported; fix them in the file and upload again.</p>'
            f'<button class="btn" type="submit" {"disabled" if done else ""}>Confirm and save ticked rows</button></form>')
    body = (ui.page_header("Check before saving", f"{len(rows)} rows {by}. Untick anything that is wrong.",
                           '<a class="btn secondary" href="/admin/import">Back</a>')
            + '<div class="stack">' + (ui.alert(f"Already saved on {e(done)}.", "ok") if done else "")
            + ui.card("Proposed products and facts", form) + "</div>")
    return page("Check import", body, u, "facts", [("Products and Facts", "/admin/projects"), ("Import", "/admin/import"), ("Check", None)], request)


@app.post("/admin/import/{iid}/apply")
async def import_apply(request: Request, iid: int):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    meta, data = IMP.load(u["brand_id"], iid)
    if not meta or meta["applied_at"]:
        return RedirectResponse("/admin/import", 303)
    form = await request.form()
    pick = {int(x) for x in form.getlist("row") if str(x).isdigit()}
    rows = [r for r in data["rows"] if r["id"] in pick and (r["ok"] or "not found in the document" in r["problem"])]
    source = (form.get("source") or meta["title"]).strip()[:200]
    if not rows:
        request.session["flash"] = "Nothing was ticked, so nothing was saved."
        return RedirectResponse(f"/admin/import/{iid}", 303)
    try:
        summary = apply_import(u, rows, source)
    except ValueError as ex_:
        request.session["flash"] = f"Import stopped: {ex_}"
        return RedirectResponse(f"/admin/import/{iid}", 303)
    IMP.mark_applied(iid)
    log(None, u["name"], "IMPORT_APPLIED", f"{meta['kind']} {meta['title']}: {summary}", u["brand_id"])
    request.session["flash"] = f"Saved: {summary}."
    return RedirectResponse("/admin/projects", 303)


@app.post("/admin/sample")
def load_sample(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    sample = SAMPLE_PROPERTY if load_reg(u).get("domain", "property") == "property" else SAMPLE_GENERAL
    rows = IMP.check(json.loads(json.dumps(sample)), None, parse_price, GEN.list_defs(u["brand_id"]))
    summary = apply_import(u, [r for r in rows if r["ok"]], "BrandGuard sample data (not real)")
    log(None, u["name"], "SAMPLE_LOADED", summary, u["brand_id"])
    request.session["flash"] = f"Sample data loaded: {summary}. It is marked as sample data; remove it before real use."
    return RedirectResponse("/admin/projects", 303)


@app.get("/admin/fact-types", response_class=HTMLResponse)
def fact_types_page(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    rows = [[f'<b>{e(d["label"])}</b>', e(GEN.TYPES[d["ftype"]]), e(", ".join(d["units"]) or "-"),
             e(", ".join(d["aliases"]) or "-"), e(GEN.DIRECTIONS[d["direction"]]),
             f'<form method="post" action="/admin/fact-types/remove" style="display:inline"><input type="hidden" name="def_id" value="{d["id"]}">'
             f'<button class="btn secondary sm" type="submit">Remove</button></form>'] for d in GEN.list_defs(u["brand_id"])]
    t_opts = "".join(f'<option value="{k}">{e(v)}</option>' for k, v in GEN.TYPES.items())
    d_opts = "".join(f'<option value="{k}">{e(v)}</option>' for k, v in GEN.DIRECTIONS.items())
    form = (f'<form method="post" action="/admin/fact-types" data-lock><div class="field-grid">'
            f'<div><label for="dl">Name</label><input id="dl" name="label" required maxlength="60" placeholder="Discount, Mileage, Processing fee"></div>'
            f'<div><label for="dt">Type of value</label><select id="dt" name="ftype">{t_opts}</select></div>'
            f'<div><label for="du">Unit (numbers only)</label><input id="du" name="unit" placeholder="kmpl, km/l  or  years, yrs"></div>'
            f'<div><label for="da">Other words people use</label><input id="da" name="aliases" placeholder="off, सूट, सवलत"></div>'
            f'<div><label for="dd">Rule</label><select id="dd" name="direction">{d_opts}</select></div></div>'
            f'<p class="hint">Separate several spellings with commas. "Must not be more than this" suits discounts, mileage and '
            f'warranties. "Must not be less than this" suits interest rates and delivery dates.</p>'
            f'<button class="btn" type="submit">{ui.icon("plus", 16)}Save fact type</button></form>')
    body = (ui.page_header("Fact types", "What your content makes claims about. Define each once; give each product its values.",
                           '<a class="btn secondary" href="/admin/projects">Products and facts</a>')
            + '<div class="stack">' + ui.card("Your fact types", ui.table(["Name", "Type", "Unit", "Other words", "Rule", ""], rows, "Fact types")
                                              if rows else ui.empty("No fact types yet", "Add one below, or import a spreadsheet or brochure."))
            + ui.card("Add or change a fact type", form) + "</div>")
    return page("Fact types", body, u, "facts", [("Products and Facts", "/admin/projects"), ("Fact types", None)], request)


@app.post("/admin/fact-types")
def fact_types_save(request: Request, label: str = Form(""), ftype: str = Form(""), unit: str = Form(""),
                    aliases: str = Form(""), direction: str = Form("exact")):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    try:
        GEN.add_def(u["brand_id"], label, ftype, unit, aliases, direction)
        log(None, u["name"], "FACT_TYPE_SET", f"{label.strip()} ({ftype})", u["brand_id"])
        request.session["flash"] = f"Saved {label.strip()}."
    except ValueError as ex_:
        request.session["flash"] = str(ex_)
    return RedirectResponse("/admin/fact-types", 303)


@app.post("/admin/fact-types/remove")
def fact_types_remove(request: Request, def_id: int = Form(0)):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    d = GEN.get_def_by_id(u["brand_id"], def_id)
    if d:
        GEN.remove_def(u["brand_id"], def_id)
        log(None, u["name"], "FACT_TYPE_REMOVED", d["label"], u["brand_id"])
        request.session["flash"] = f"Removed {d['label']}. Its stored values are kept but no longer checked."
    return RedirectResponse("/admin/fact-types", 303)


def _save_fact(u, p, key, unit_type_id, value, source, reason, valid_from, valid_to):
    num, txt = parse_fact_value(key, value)
    if unit_type_id:
        ut = RE._q("SELECT * FROM unit_types WHERE id=? AND project_id=?", (unit_type_id, p["id"]))
        if not ut:
            raise ValueError("That property type does not belong to this project.")
        if key in ("possession", "ready_to_move", "booking_amount") or key.startswith("offer:"):
            unit_type_id = None
        elif key not in RE.CATEGORIES[ut[0]["category"]][1] and key != "price":
            raise ValueError("That fact does not apply to this property category.")
    elif key in RE.NUMERIC_FIELDS or key == "price":
        raise ValueError("Choose the property type this applies to.")
    if valid_from and valid_to and valid_to < valid_from:
        raise ValueError("The end date is before the start date.")
    sid_ = source_id_for(u, source)
    unit = "INR" if key in ("price", "booking_amount") else ("sq ft" if key.endswith("sqft") else "")
    fid, old = RE.set_fact(p["id"], key, num, txt, unit_type_id, unit, sid_, valid_from or None, valid_to or None,
                           u["name"], reason)
    log(None, u["name"], "FACT_SET", f'{p["official_name"]}: {RE.FACT_LABEL.get(key, key)} = {txt if txt else num}'
        + (f" (replaced {len(old)} earlier value)" if old else ""), u["brand_id"])
    return fid


@app.post("/admin/projects/{pid}/unit")
async def project_unit(request: Request, pid: int):
    u = admin_only(request)
    p = own_project(u, pid) if u else None
    if not p:
        return RedirectResponse("/admin/projects", 303)
    f = await request.form()
    cat, label = f.get("category", ""), (f.get("label", "") or "").strip()
    try:
        if cat not in RE.CATEGORIES or not label:
            raise ValueError("Choose a category and give the property type a name.")
        allowed = RE.CATEGORIES[cat][1]
        bhk = None
        if "bhk" in allowed and (f.get("bhk") or "").strip():
            try:
                bhk = int(float(f.get("bhk")))
            except ValueError:
                raise ValueError("BHK must be a number.")
        if cat == "product" and (f.get("price") or "").strip() and not (f.get("source") or "").strip():
            raise ValueError("Name the source document for the price.")
        uid = RE.add_unit_type(pid, cat, label, bhk)
        if cat != "product" or (f.get("price") or "").strip():
            _save_fact(u, p, "price", uid, f.get("price", ""), f.get("source", ""), "Added with property type", "", "")
        for k in allowed:
            if k in ("bhk",) or not (f.get(k) or "").strip():
                continue
            if k in RE.NUMERIC_FIELDS:
                _save_fact(u, p, k, uid, f.get(k), f.get("source", ""), "Added with property type", "", "")
            else:
                num, txt = None, f.get(k).strip()[:100]
                RE.set_fact(pid, k, num, txt, uid, "", source_id_for(u, f.get("source", "")), None, None, u["name"], "Added with property type")
        log(None, u["name"], "UNIT_ADDED", f'{p["official_name"]}: {label} ({cat})', u["brand_id"])
        return admin_back(request, pid, f"Added {label}.")
    except ValueError as ex_:
        return admin_back(request, pid, err=str(ex_))


@app.post("/admin/projects/{pid}/fact")
def project_fact(request: Request, pid: int, key: str = Form(""), unit_type_id: str = Form(""), value: str = Form(""),
                 source: str = Form(""), reason: str = Form(""), valid_from: str = Form(""), valid_to: str = Form("")):
    u = admin_only(request)
    p = own_project(u, pid) if u else None
    if not p:
        return RedirectResponse("/admin/projects", 303)
    if not key or not source.strip():
        return admin_back(request, pid, err="Choose a fact and name the source document.")
    if key.startswith("offer:") and not valid_to:
        return admin_back(request, pid, err="Offers need an end date, so that expired offers are never treated as current.")
    try:
        _save_fact(u, p, key, int(unit_type_id) if unit_type_id.isdigit() else None, value, source, reason, valid_from, valid_to)
    except ValueError as ex_:
        return admin_back(request, pid, err=str(ex_))
    return admin_back(request, pid, "Saved. Earlier approvals that relied on the old value have been flagged for re-check.")


@app.post("/admin/projects/{pid}/fact/end")
def project_fact_end(request: Request, pid: int, fact_id: int = Form(0)):
    u = admin_only(request)
    p = own_project(u, pid) if u else None
    if not p:
        return RedirectResponse("/admin/projects", 303)
    f = RE._q("SELECT * FROM facts WHERE id=? AND project_id=?", (fact_id, pid))
    if f:
        RE.end_fact(fact_id, u["name"])
        log(None, u["name"], "FACT_ENDED", f'{p["official_name"]}: {RE.FACT_LABEL.get(f[0]["key"], f[0]["key"])}', u["brand_id"])
    return admin_back(request, pid, "Fact withdrawn. It is no longer treated as current.")


@app.post("/admin/projects/{pid}/landmark")
def project_landmark(request: Request, pid: int, place: str = Form(""), minutes: str = Form(""), how: str = Form(""),
                     source: str = Form("")):
    u = admin_only(request)
    p = own_project(u, pid) if u else None
    if not p:
        return RedirectResponse("/admin/projects", 303)
    try:
        if not place.strip() or not source.strip():
            raise ValueError("Enter the place and the source document.")
        num, _ = parse_fact_value("landmark:x", minutes)
        RE.set_fact(pid, "landmark:" + place.strip().lower()[:60], num, None, None, (how.strip() or "min")[:30],
                    source_id_for(u, source), None, None, u["name"], "Travel time")
        log(None, u["name"], "FACT_SET", f'{p["official_name"]}: {place.strip()} {num:g} min', u["brand_id"])
    except ValueError as ex_:
        return admin_back(request, pid, err=str(ex_))
    return admin_back(request, pid, "Saved.")


@app.post("/admin/projects/{pid}/alias")
def project_alias(request: Request, pid: int, alias: str = Form("")):
    u = admin_only(request)
    p = own_project(u, pid) if u else None
    if not p:
        return RedirectResponse("/admin/projects", 303)
    if alias.strip():
        ex("INSERT INTO project_aliases (project_id, alias) VALUES (?,?)", (pid, alias.strip()[:80]))
        log(None, u["name"], "PROJECT_EDITED", f'{p["official_name"]}: added name "{alias.strip()}"', u["brand_id"])
    return admin_back(request, pid, "Name added.")


@app.post("/admin/projects/{pid}/rera")
def project_rera(request: Request, pid: int, rera_no: str = Form(""), source: str = Form("")):
    u = admin_only(request)
    p = own_project(u, pid) if u else None
    if not p:
        return RedirectResponse("/admin/projects", 303)
    rera_no = rera_no.strip().upper()
    if rera_no and not RE.rera_format_ok(rera_no):
        return admin_back(request, pid, err="A RERA number looks like P followed by 11 digits.")
    level = "none" if not rera_no else "format"
    ex("UPDATE projects SET rera_no=?, rera_level=?, rera_source=NULL, rera_checked_at=NULL, rera_checked_by=NULL WHERE id=?",
       (rera_no, level, pid))
    if rera_no and source.strip():
        ex("UPDATE projects SET rera_level='authoritative', rera_source=?, rera_checked_at=?, rera_checked_by=? WHERE id=?",
           (source.strip()[:120], RE.today(), u["name"], pid))
        log(None, u["name"], "RERA_RECORDED", f'{p["official_name"]}: {rera_no} confirmed against {source.strip()}', u["brand_id"])
    else:
        log(None, u["name"], "PROJECT_EDITED", f'{p["official_name"]}: RERA number {rera_no or "removed"}', u["brand_id"])
    return admin_back(request, pid, "RERA details saved.")


@app.get("/admin/access", response_class=HTMLResponse)
def access_page(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    users = q("""SELECT us.id, us.name FROM users us JOIN memberships m ON m.user_id=us.id
                 WHERE m.brand_id=? ORDER BY us.name""", (u["brand_id"],))
    projs = RE.brand_projects(u["brand_id"])
    rows = []
    for us in users:
        gs = [g for g in q("SELECT * FROM grants WHERE user_id=? AND brand_id=? ORDER BY id", (us["id"], u["brand_id"])) if RE._grant_active(g)]
        items = ""
        for g in gs:
            pr = RE.get_project(g["project_id"])
            exp = f' <span class="muted">until {e(g["expires"])}</span>' if g["expires"] else ""
            items += (f'<form method="post" action="/admin/access/revoke" style="display:inline-block;margin:2px 6px 2px 0">'
                      f'<input type="hidden" name="grant_id" value="{g["id"]}">{ui.badge(pr["official_name"] if pr else "?", "info")}{exp} '
                      f'<button class="btn secondary sm" type="submit" aria-label="Remove access to {e(pr["official_name"] if pr else "")}">Remove</button></form>')
        rows.append([e(us["name"]), items or '<span class="muted">No project access</span>'])
    boxes = "".join(f'<label class="check"><input type="checkbox" name="project_id" value="{p["id"]}"> {e(p["official_name"])}'
                    f'{" (" + e(p["city"]) + ")" if p["city"] else ""}</label>' for p in projs)
    uopts = "".join(f'<option value="{us["id"]}">{e(us["name"])}</option>' for us in users)
    form = (f'<form method="post" action="/admin/access/grant" data-lock><label for="gu">Partner or freelancer</label>'
            f'<select id="gu" name="user_id">{uopts}</select><p></p><b>Projects</b><div>{boxes}</div><p></p>'
            f'<label for="gx">Access ends on (optional)</label><input id="gx" type="date" name="expires">'
            f'<p class="hint">Tick as many projects as needed. A city or company link does not grant every project, only the ticked ones.</p>'
            f'<button class="btn" type="submit">Grant access</button></form>')
    body = (ui.page_header("Project Access", "Who may submit content for which project. Nobody needs a per-video assignment.")
            + '<div class="stack">' + ui.card("Current access", ui.table(["Person", "Projects"], rows, "Access") if rows else ui.empty("No partners linked"))
            + ui.card("Grant access", form) + "</div>")
    return page("Project Access", body, u, "access", [("Project Access", None)], request)


@app.post("/admin/access/grant")
async def access_grant(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    f = await request.form()
    uid = f.get("user_id", "")
    if not str(uid).isdigit() or not q("SELECT 1 FROM memberships WHERE user_id=? AND brand_id=?", (int(uid), u["brand_id"])):
        request.session["flash"] = "That person is not linked to your company."
        return RedirectResponse("/admin/access", 303)
    n = 0
    for pid in f.getlist("project_id"):
        if str(pid).isdigit() and own_project(u, int(pid)):
            RE.grant(int(uid), u["brand_id"], int(pid), u["name"], (f.get("expires") or None))
            log(None, u["name"], "ACCESS_GRANTED", f"user {uid} -> {RE.get_project(int(pid))['official_name']}", u["brand_id"])
            n += 1
    request.session["flash"] = f"Access granted for {n} project(s)." if n else "Tick at least one project."
    return RedirectResponse("/admin/access", 303)


@app.post("/admin/access/revoke")
def access_revoke(request: Request, grant_id: int = Form(0)):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    if RE.revoke(grant_id, u["brand_id"], u["name"]):
        log(None, u["name"], "ACCESS_REVOKED", f"grant {grant_id}", u["brand_id"])
        request.session["flash"] = "Access removed. It takes effect immediately."
    return RedirectResponse("/admin/access", 303)


# ============================================================ accounts, sign-in and invites
def client_ip(request):
    fwd = request.headers.get("x-forwarded-for", "")
    return (fwd.split(",")[0].strip() if fwd else (request.client.host if request.client else ""))


def public_base(request):
    base = os.environ.get("PUBLIC_URL", "").rstrip("/")
    if base:
        return base
    proto = request.headers.get("x-forwarded-proto", request.url.scheme)
    return f"{proto}://{request.headers.get('host', request.url.netloc)}"


def start_session(request, user_id):
    request.session.clear()          # a fresh session on every sign-in
    request.session["uid"] = user_id


def auth_form(title, intro, fields, action, button, error="", extra="", industry=None):
    inputs = "".join(
        f'<div><label for="f-{n}">{e(lbl)}</label><input id="f-{n}" name="{n}" type="{t}" value="{e(v)}" '
        f'{"required " if req else ""}autocomplete="{ac}"></div>' for n, lbl, t, v, req, ac in fields)
    alert = ui.alert(e(error), "bad", "Could not continue.") if error else ""
    if industry is not None:
        opts = "".join(f'<option value="{k}"{" selected" if k == industry else ""}>{e(lbl)}</option>' for k, lbl in packs.OPTIONS)
        inputs += (f'<div><label for="f-industry">Your industry</label><select id="f-industry" name="industry">{opts}</select>'
                   f'<p class="hint">You get a starter list of claims to avoid for this industry. You can edit it any time.</p></div>')
    return (f'<div style="max-width:460px;margin:0 auto"><h1>{e(title)}</h1><p class="page-desc">{intro}</p>{alert}'
            f'<form method="post" action="{action}" data-lock class="stack">{inputs}'
            f'<button class="btn" type="submit">{e(button)}</button></form>{extra}</div>')


@app.get("/signup", response_class=HTMLResponse)
def signup_form(request: Request):
    if me(request):
        return RedirectResponse("/", 303)
    return page("Create company", auth_form(
        "Set up your company", "You become the builder for this company. You can invite everyone else afterwards.",
        [("company", "Company name", "text", "", True, "organization"), ("name", "Your name", "text", "", True, "name"),
         ("email", "Work email", "email", "", True, "email"),
         ("password", f"Password (at least {ACC.MIN_PASSWORD} characters)", "password", "", True, "new-password")],
        "/signup", "Create company",
        extra='<p class="muted" style="margin-top:14px">Already have an account? <a href="/login">Sign in</a></p>', industry=""))


@app.post("/signup")
def signup(request: Request, company: str = Form(""), name: str = Form(""), email: str = Form(""),
           password: str = Form(""), industry: str = Form("")):
    industry = industry if industry in packs.PACKS else ""
    try:
        uid, bid = ACC.create_company(company, name, email, password, industry)
    except ValueError as ex_:
        return page("Create company", auth_form(
            "Set up your company", "You become the builder for this company.",
            [("company", "Company name", "text", company, True, "organization"), ("name", "Your name", "text", name, True, "name"),
             ("email", "Work email", "email", email, True, "email"),
             ("password", f"Password (at least {ACC.MIN_PASSWORD} characters)", "password", "", True, "new-password")],
            "/signup", "Create company", str(ex_), industry=industry))
    start_session(request, uid)
    log(None, name.strip(), "COMPANY_CREATED", f"Company created: {company.strip()}", bid)
    request.session["flash"] = "Company created. Follow the steps on this page to finish setting up."
    return RedirectResponse("/", 303)


@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request):
    if me(request):
        return RedirectResponse("/", 303)
    return page("Sign in", auth_form(
        "Sign in", "Use the email and password you registered with.",
        [("email", "Email", "email", "", True, "email"), ("password", "Password", "password", "", True, "current-password")],
        "/login", "Sign in",
        extra='<p class="muted" style="margin-top:14px">New here? <a href="/signup">Create a company account</a>. '
              'Invited by a company? Open the link in your invite email.</p>'))


@app.post("/login")
def login(request: Request, email: str = Form(""), password: str = Form("")):
    try:
        user = ACC.authenticate(email, password, client_ip(request))
        err = "" if user else "Wrong email or password."
    except ValueError as ex_:
        user, err = None, str(ex_)
    if not user:
        return page("Sign in", auth_form(
            "Sign in", "Use the email and password you registered with.",
            [("email", "Email", "email", email, True, "email"), ("password", "Password", "password", "", True, "current-password")],
            "/login", "Sign in", err))
    start_session(request, user["id"])
    return RedirectResponse("/", 303)


@app.get("/invite/{token}", response_class=HTMLResponse)
def invite_form(request: Request, token: str):
    inv = ACC.get_invite(token)
    if not inv:
        return page("Invite", ui.alert("This invite link is not valid any more. It may have been used or have expired. "
                                       "Ask the company to send a new one.", "bad", "Invite not valid."))
    return page("Join " + inv["company"], invite_html(inv, token))


def invite_html(inv, token, error=""):
    label = ACC.KINDS[inv["kind"]][0]
    existing = bool(q("SELECT 1 FROM users WHERE lower(email)=?", (inv["email"],)))
    fields = ([] if existing else [("name", "Your name", "text", "", True, "name")]) + [
        ("password", "Your password" if existing else f"Choose a password (at least {ACC.MIN_PASSWORD} characters)",
         "password", "", True, "current-password" if existing else "new-password")]
    intro = (f"{e(inv['company'])} has invited <b>{e(inv['email'])}</b> as <b>{e(label)}</b>. "
             + ("This email already has an account. Enter its password to join." if existing else "Create your account to join."))
    return auth_form("Join " + inv["company"], intro, fields, f"/invite/{e(token)}", "Join company", error)


@app.post("/invite/{token}")
def invite_accept(request: Request, token: str, name: str = Form(""), password: str = Form("")):
    inv = ACC.get_invite(token)
    if not inv:
        return page("Invite", ui.alert("This invite link is not valid any more.", "bad", "Invite not valid."))
    if ACC.too_many_attempts(inv["email"], client_ip(request)):
        return page("Join " + inv["company"], invite_html(inv, token, "Too many attempts. Wait a few minutes and try again."))
    try:
        uid, bid, kind = ACC.accept_invite(token, name, password)
    except ValueError as ex_:
        ACC._note_fail(inv["email"], client_ip(request))
        return page("Join " + inv["company"], invite_html(inv, token, str(ex_)))
    pids = json.loads(inv["project_ids"] or "[]")
    who = q("SELECT name FROM users WHERE id=?", (uid,))[0]["name"]
    if ACC.KINDS[kind][1] == "submitter" and "*" in pids:
        ex("UPDATE memberships SET all_projects=1 WHERE user_id=? AND brand_id=?", (uid, bid))
        pids = [p["id"] for p in RE.brand_projects(bid)]
    if ACC.KINDS[kind][1] == "submitter":
        for pid in pids:
            p = RE.get_project(int(pid))
            if p and p["brand_id"] == bid:
                RE.grant(uid, bid, int(pid), "invite by " + (inv["invited_by"] or "company"))
    log(None, who, "INVITE_ACCEPTED", f"{who} joined as {ACC.KINDS[kind][0]}", bid)
    start_session(request, uid)
    request.session["flash"] = f"Welcome to {inv['company']}."
    return RedirectResponse("/", 303)


@app.get("/admin/team", response_class=HTMLResponse)
def team_page(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    bid = u["brand_id"]
    link = request.session.pop("invite_link", "")
    notice = ""
    if link:
        notice = ui.alert(f'Invite link (shown once, single use, valid {ACC.INVITE_DAYS} days): '
                          f'<br><input readonly value="{e(link)}" style="width:100%" onclick="this.select()">',
                          "ok", "Invite created.")
    mrows = []
    for m in ACC.members(bid):
        label = "Builder" if m["kind"] == "builder" else ACC.KINDS.get(m["kind"], (m["kind"],))[0]
        rm = ""
        if m["kind"] != "builder":
            rm = (f'<form method="post" action="/admin/team/remove" style="display:inline"><input type="hidden" name="user_id" value="{m["id"]}">'
                  f'<button class="btn secondary sm" type="submit">Remove</button></form>')
        mrows.append([e(m["name"]), e(m["email"] or ""), e(label), e(str(m["joined"] or "")[:10]), rm])
    irows = []
    for i in ACC.pending_invites(bid):
        irows.append([e(i["email"]), e(ACC.KINDS[i["kind"]][0]), e(i["expires_at"][:10]),
                      f'<form method="post" action="/admin/team/revoke" style="display:inline"><input type="hidden" name="invite_id" value="{i["id"]}">'
                      f'<button class="btn secondary sm" type="submit">Cancel</button></form>'])
    projects = RE.brand_projects(bid)
    word = "projects" if load_reg(u).get("domain", "property") == "property" else "products"
    boxes = (f'<label style="display:block;font-weight:600"><input type="checkbox" name="all_projects" value="1" checked '
             f'id="allp"> All {word}, including ones you add later</label>'
             + (f'<p class="hint" style="margin-top:6px">Or untick that and choose only some:</p>' if projects else
                f'<p class="hint" style="margin-top:6px">You have no {word} yet. With this ticked, the person gets each one as you add it.</p>')
             + "".join(f'<label style="display:block;margin-left:22px"><input type="checkbox" name="project_id" value="{p["id"]}"> '
                       f'{e(p["official_name"])}</label>' for p in projects))
    kinds = "".join(f'<option value="{k}">{e(v[0])}</option>' for k, v in ACC.KINDS.items())
    mail_note = ("The invite is also emailed." if mailer.configured()
                 else "Email sending is not set up on this server, so copy the link and send it yourself.")
    form = (f'<form method="post" action="/admin/team/invite" data-lock><label for="ie">Email address</label>'
            f'<input id="ie" name="email" type="email" required><p></p><label for="ik">Role</label>'
            f'<select id="ik" name="kind">{kinds}</select><p></p><b>What they may promote</b><div>{boxes}</div>'
            f'<p class="hint">Partners, freelancers and agencies only see and submit for the {word} chosen here. Managers get full '
            f'builder access. {mail_note}</p>'
            f'<button class="btn" type="submit">Create invite</button></form>')
    body = (ui.page_header("Team and Invites", f"Invite people to your company and control which {word} they can use.")
            + '<div class="stack">' + notice
            + ui.card("People", ui.table(["Name", "Email", "Role", "Joined", ""], mrows, "People"))
            + ui.card("Pending invites", ui.table(["Email", "Role", "Expires", ""], irows, "Pending invites")
                      if irows else ui.empty("No pending invites"))
            + ui.card("Invite someone", form) + bulk_invite_card(word) + email_card(u) + "</div>")
    return page("Team and Invites", body, u, "team", [("Team and Invites", None)], request)


INVITE_COLUMNS = ["Email", "Name", "Role", "Access"]
ROLE_WORDS = {"channel partner": "partner", "partner": "partner", "cp": "partner", "freelancer": "freelancer",
              "agency": "agency", "marketing agency": "agency", "manager": "manager", "builder": "manager", "admin": "manager"}


def email_card(u):
    p = mailer.provider()
    state = {"brevo": "Email is sent through Brevo.", "smtp": "Email is sent through your SMTP server."}.get(
        p, "Email is not set up, so invite links are shown on screen and you send them yourself.")
    how = ("" if p else
           '<p class="muted">To send invites automatically: create a free Brevo account (300 emails a day), verify the address '
           'you send from, create an API key, then add <span class="mono">BREVO_API_KEY</span> and <span class="mono">MAIL_FROM</span> '
           'in the server settings (Render: Environment). Render\'s free plan blocks normal email ports, so Gmail SMTP will not '
           'work there; Brevo does.</p>')
    btn = ('' if not p else f'<form method="post" action="/admin/email-test" data-lock><button class="btn secondary" type="submit">'
           f'Send a test email to {e(u["email"] or "me")}</button></form>')
    return ui.card("Email", f"<p>{e(state)}</p>{how}{btn}")


@app.post("/admin/email-test")
def email_test(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    ok = bool(u["email"]) and mailer.send(u["email"], "BrandGuard test email",
                                           "This is a test from BrandGuard. If you can read it, invites will be emailed.")
    request.session["flash"] = (f"Test email sent to {u['email']}. Check the inbox and the spam folder." if ok
                                else f"The test email was not sent. {mailer.last_error or 'Your account has no email address.'}")
    return RedirectResponse("/admin/team", 303)


def bulk_invite_card(word):
    return ui.card("Invite many people from a spreadsheet",
                   f'<p class="muted">One row per person. Columns: <b>Email</b>, <b>Name</b> (optional, used in the email), '
                   f'<b>Role</b> (Channel partner, Freelancer, Agency or Manager) and <b>Access</b> (write <b>All</b>, or the '
                   f'{word} names separated by commas; empty means All). You see every row before invites are created.</p>'
                   f'<p><a class="btn secondary sm" href="/admin/team/template">Download the template (.xlsx)</a></p>'
                   f'<form method="post" action="/admin/team/bulk" enctype="multipart/form-data" data-lock>'
                   f'<input type="file" name="file" accept=".xlsx,.csv" required> '
                   f'<button class="btn" type="submit">Read the list</button></form>')


@app.get("/admin/team/template")
def team_template(request: Request):
    if not admin_only(request):
        return RedirectResponse("/", 303)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(INVITE_COLUMNS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")
    for col, w in zip("ABCD", (34, 24, 20, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="brandguard-invite-list.xlsx"'})


def read_people(filename, data):
    """[{email, name, role, access}] from an .xlsx or .csv list."""
    rows = []
    if (filename or "").lower().endswith(".csv"):
        import csv
        text = data.decode("utf-8-sig", errors="replace")
        for r in csv.DictReader(io.StringIO(text)):
            rows.append({(k or "").strip().lower(): (v or "").strip() for k, v in r.items()})
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ValueError("This file could not be read. Upload the filled template (.xlsx) or a .csv file.")
        got = list(wb.worksheets[0].iter_rows(values_only=True))
        if not got:
            return []
        head = [str(h or "").strip().lower() for h in got[0]]
        for r in got[1:501]:
            rows.append({head[i]: str(v).strip() if v is not None else "" for i, v in enumerate(r) if i < len(head)})
    out = [r for r in rows if any(r.values())]
    if out and "email" not in out[0]:
        raise ValueError("The first row must have an Email column.")
    return out[:500]


def check_people(u, rows):
    projects = RE.brand_projects(u["brand_id"])
    names = {}
    for p in projects:
        for n in RE.names_of(p):
            names[n.strip().lower()] = p
    seen, out = set(), []
    for r in rows:
        email = ACC.clean_email(r.get("email", ""))
        role_raw = (r.get("role") or "channel partner").strip().lower()
        kind = ROLE_WORDS.get(role_raw) or (role_raw if role_raw in ACC.KINDS else "")
        access = (r.get("access") or "all").strip()
        problem, pids, all_p = "", [], False
        if not email:
            problem = f'"{r.get("email", "")}" is not an email address.'
        elif email in seen:
            problem = "This email is listed twice."
        elif not kind:
            problem = f'Role "{r.get("role")}" is not one of Channel partner, Freelancer, Agency, Manager.'
        elif q("SELECT 1 FROM users WHERE lower(email)=? AND role='admin' AND brand_id=?", (email, u["brand_id"])):
            problem = "Already a builder in your company."
        elif ACC.KINDS[kind][1] == "submitter":
            if access.lower() in ("all", "*", "everything"):
                all_p = True
            else:
                for part in [x.strip() for x in access.split(",") if x.strip()]:
                    p = names.get(part.lower())
                    if not p:
                        problem = f'"{part}" is not one of your projects or products.'
                        break
                    pids.append(p["id"])
        seen.add(email)
        out.append({"email": email or r.get("email", ""), "name": (r.get("name") or "")[:80], "kind": kind, "pids": pids,
                    "all": all_p, "ok": not problem, "problem": problem,
                    "access": "All" if all_p else ", ".join(RE.get_project(i)["official_name"] for i in pids) or "-"})
    return out


@app.post("/admin/team/bulk")
def team_bulk(request: Request, file: UploadFile = File(...)):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    try:
        people = check_people(u, read_people(file.filename, _read_small(file, 2)))
        if not people:
            raise ValueError("No people were found in the file.")
    except ValueError as ex_:
        request.session["flash"] = str(ex_)
        return RedirectResponse("/admin/team", 303)
    iid = IMP.save(u["brand_id"], "invites", file.filename or "Invite list", {"people": people}, u["name"])
    return RedirectResponse(f"/admin/team/bulk/{iid}", 303)


@app.get("/admin/team/bulk/{iid}", response_class=HTMLResponse)
def team_bulk_preview(request: Request, iid: int):
    u = admin_only(request)
    meta, data = IMP.load(u["brand_id"], iid) if u else (None, None)
    if not meta or meta["kind"] != "invites":
        return RedirectResponse("/admin/team", 303)
    people = data["people"]
    rows = [[e(p["email"]), e(p["name"] or "-"), e(ACC.KINDS[p["kind"]][0] if p["kind"] else "-"), e(p["access"]),
             ui.badge("Ready", "ok", "check") if p["ok"] else f'{ui.badge("Skipped", "warn", "alert")} <span class="muted">{e(p["problem"])}</span>']
            for p in people]
    n_ok = sum(1 for p in people if p["ok"])
    sends = "Each person gets an email with their own link." if mailer.configured() else \
        "Email is not set up, so you will get a list of links to send yourself."
    form = ('' if meta["applied_at"] else
            f'<form method="post" action="/admin/team/bulk/{iid}/send" data-lock><button class="btn" type="submit" '
            f'{"disabled" if not n_ok else ""}>Create {n_ok} invite(s)</button></form>')
    body = (ui.page_header("Check the invite list", f"{n_ok} of {len(people)} rows are ready. Rows with a problem are skipped. {sends}",
                           '<a class="btn secondary" href="/admin/team">Back</a>')
            + '<div class="stack">' + (ui.alert(f"Already sent on {e(meta['applied_at'])}.", "ok") if meta["applied_at"] else "")
            + ui.card("People", ui.table(["Email", "Name", "Role", "Access", "Check"], rows, "Invite list") + form) + "</div>")
    return page("Invite list", body, u, "team", [("Team and Invites", "/admin/team"), ("Invite list", None)], request)


@app.post("/admin/team/bulk/{iid}/send", response_class=HTMLResponse)
def team_bulk_send(request: Request, iid: int):
    u = admin_only(request)
    meta, data = IMP.load(u["brand_id"], iid) if u else (None, None)
    if not meta or meta["kind"] != "invites" or meta["applied_at"]:
        return RedirectResponse("/admin/team", 303)
    IMP.mark_applied(iid)
    company = q("SELECT name FROM brands WHERE id=?", (u["brand_id"],))[0]["name"]
    results = []
    for p in data["people"]:
        if not p["ok"]:
            continue
        try:
            tok = ACC.create_invite(u["brand_id"], p["email"], p["kind"], p["pids"], u["name"], all_projects=p["all"])
        except ValueError as ex_:
            results.append((p["email"], "", False, str(ex_)))
            continue
        link = f"{public_base(request)}/invite/{tok}"
        hello = f"Hello {p['name']},\n\n" if p["name"] else ""
        sent = mailer.send(p["email"], f"{u['name']} invited you to {company} on BrandGuard",
                           f"{hello}{u['name']} has invited you to join {company} on BrandGuard as {ACC.KINDS[p['kind']][0]}.\n\n"
                           f"Open this link to create your account (valid {ACC.INVITE_DAYS} days, single use):\n{link}\n")
        results.append((p["email"], link, sent, ""))
    n_sent = sum(1 for r in results if r[2])
    n_made = sum(1 for r in results if r[1])
    log(None, u["name"], "INVITE_CREATED", f"Bulk: {n_made} invite(s) created, {n_sent} emailed", u["brand_id"])
    rows = [[e(em), ui.badge("Emailed", "ok", "check") if sent else (ui.badge("Not created", "bad", "x-circle") if not link
             else ui.badge("Send yourself", "warn", "alert")),
             (f'<input readonly value="{e(link)}" style="width:100%;margin:0" onclick="this.select()">' if link and not sent else e(err or ""))]
            for em, link, sent, err in results]
    note = ("Links are shown only on this page, once. Copy the ones marked 'Send yourself' now."
            if n_made > n_sent else "Everyone was emailed.")
    body = (ui.page_header("Invites created", f"{n_made} invite(s) created, {n_sent} emailed.",
                           '<a class="btn" href="/admin/team">Back to Team</a>')
            + '<div class="stack">' + ui.alert(note, "info" if n_made > n_sent else "ok")
            + (ui.alert(f"Email problem: {e(mailer.last_error)}", "warn") if n_made > n_sent and mailer.configured() and mailer.last_error else "")
            + ui.card("Results", ui.table(["Email", "Status", "Link"], rows, "Invite results")) + "</div>")
    return page("Invites created", body, u, "team", [("Team and Invites", "/admin/team"), ("Invites created", None)], request)


@app.post("/admin/team/invite")
async def team_invite(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    f = await request.form()
    pids = [int(p) for p in f.getlist("project_id") if str(p).isdigit() and own_project(u, int(p))]
    all_p = bool(f.get("all_projects"))
    try:
        tok = ACC.create_invite(u["brand_id"], f.get("email", ""), f.get("kind", ""), pids, u["name"], all_projects=all_p)
    except ValueError as ex_:
        request.session["flash"] = str(ex_)
        return RedirectResponse("/admin/team", 303)
    link = f"{public_base(request)}/invite/{tok}"
    request.session["invite_link"] = link
    company = q("SELECT name FROM brands WHERE id=?", (u["brand_id"],))[0]["name"]
    sent = mailer.send(ACC.clean_email(f.get("email", "")), f"{u['name']} invited you to {company} on BrandGuard",
                       f"{u['name']} has invited you to join {company} on BrandGuard.\n\nOpen this link to create your "
                       f"account (valid {ACC.INVITE_DAYS} days, single use):\n{link}\n")
    log(None, u["name"], "INVITE_CREATED", f"{ACC.clean_email(f.get('email', ''))} as {ACC.KINDS[f.get('kind')][0]}"
        + (" (emailed)" if sent else ""), u["brand_id"])
    who = ACC.clean_email(f.get("email", ""))
    request.session["flash"] = (f"Invite emailed to {who}. The link is also shown below." if sent else
                                f"Invite created, but the email was not sent ({mailer.last_error}). Copy the link below and send it."
                                if mailer.configured() else "Invite created. Copy the link below and send it to them.")
    return RedirectResponse("/admin/team", 303)


@app.post("/admin/team/revoke")
def team_revoke(request: Request, invite_id: int = Form(0)):
    u = admin_only(request)
    if u and ACC.revoke_invite(u["brand_id"], invite_id):
        log(None, u["name"], "INVITE_CANCELLED", f"Invite {invite_id}", u["brand_id"])
        request.session["flash"] = "Invite cancelled."
    return RedirectResponse("/admin/team", 303)


@app.post("/admin/team/remove")
def team_remove(request: Request, user_id: int = Form(0)):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    if ACC.remove_member(u["brand_id"], user_id):
        for g in q("SELECT id FROM grants WHERE user_id=? AND brand_id=? AND revoked_at IS NULL", (user_id, u["brand_id"])):
            RE.revoke(g["id"], u["brand_id"], u["name"])
        log(None, u["name"], "MEMBER_REMOVED", f"User {user_id} removed. Their access ended immediately.", u["brand_id"])
        request.session["flash"] = "Removed. Their project access ended immediately."
    return RedirectResponse("/admin/team", 303)


# ============================================================ bulk upload queue
BULK_MAX_FILES = int(os.environ.get("BULK_MAX_FILES", "10"))
BULK_MAX_PENDING = int(os.environ.get("BULK_MAX_PENDING", "30"))


def process_job(job):
    """Worker callback: review one stored file. Returns (submission_id, error message)."""
    rows = q("SELECT * FROM users WHERE id=?", (job["user_id"],))
    if not rows:
        return None, "The account that uploaded this file no longer exists."
    p = Path(job["path"])
    if not p.exists():
        return None, "The stored file is missing. It may have been cleared when the server restarted. Upload it again."
    sid, msg = review_submission(rows[0], job["brand_id"], p, job["sha"], job["filename"], job["language"],
                                 hint_pid=job["hint_pid"])
    return sid, msg


def notify_batch(batch):
    u = q("SELECT * FROM users WHERE id=?", (batch["user_id"],))
    if not u:
        return
    u = u[0]
    sm = JOBS.summary(batch["id"])
    log(None, "BrandGuard engine", "BATCH_DONE",
        f"Batch {batch['id']}: {sm['approved']} approved, {sm['rejected']} rejected, {sm['check']} need a person, {sm['failed']} failed.",
        batch["brand_id"])
    if u["pw_hash"] and u["email"]:
        base = os.environ.get("PUBLIC_URL", "").rstrip("/")
        mailer.send(u["email"], f"Your BrandGuard batch #{batch['id']} is ready",
                    f"{sm['total']} videos were reviewed.\n\nApproved: {sm['approved']}\nRejected: {sm['rejected']}\n"
                    f"Need a person to check: {sm['check']}\nCould not be reviewed: {sm['failed']}\n\n"
                    f"Open the results: {base}/bulk/{batch['id']}\n")


def run_queue_once():
    """One worker step. Used by the background thread and by tests."""
    return JOBS.run_once(process_job, notify_batch)


@app.on_event("startup")
def start_worker():
    JOBS.reset_stuck()
    sweep_uploads()
    if os.environ.get("BG_WORKER", "1") != "0":
        JOBS.start(process_job, notify_batch)


JOB_BADGE = {"QUEUED": ("Waiting", "neutral", "clock"), "PROCESSING": ("Reviewing now", "warn", "clock"),
             "DONE": ("Reviewed", "ok", "check-circle"), "FAILED": ("Could not review", "bad", "x-circle")}


@app.get("/bulk", response_class=HTMLResponse)
def bulk_form(request: Request):
    u = me(request)
    if not u or u["role"] != "submitter":
        return RedirectResponse("/", 303)
    brands = user_brands(u)
    cid = current_company(request, u)
    projs = RE.authorized_projects(u["id"], cid) if cid else []
    popts = '<option value="">Detect from the video</option>' + "".join(
        f'<option value="{p["id"]}">{e(p["official_name"])}</option>' for p in projs)
    bopts = "".join(f'<option value="{b["id"]}"{" selected" if b["id"] == cid else ""}>{e(b["name"])}</option>' for b in brands)
    exts = ",".join(sorted(ALLOWED))
    form = (f'<form method="post" action="/bulk" enctype="multipart/form-data" data-lock>'
            f'<label for="bc">Company</label><select id="bc" name="brand_id">{bopts}</select><p></p>'
            f'<label for="bl">Spoken language</label><select id="bl" name="language"><option value="auto">Auto-detect</option>'
            f'<option value="mr">Marathi</option><option value="hi">Hindi</option><option value="en">English</option></select><p></p>'
            f'<label for="bp">Project (optional)</label><select id="bp" name="project_id">{popts}</select><p></p>'
            f'<label for="bf">Videos (up to {BULK_MAX_FILES} at a time)</label>'
            f'<input id="bf" type="file" name="files" multiple required accept="{exts}">'
            f'<p class="hint">Up to {MAX_MB} MB each. They are saved first and reviewed one at a time in the background, '
            f'so you can close this page. You get an email when the batch is done if email is set up for this server, '
            f'and the results stay on this page either way.</p>'
            f'<button class="btn" type="submit">{ui.icon("upload", 16)}Upload and review</button></form>')
    recent = []
    for b in JOBS.recent_batches(u["id"]):
        sm = JOBS.summary(b["id"])
        recent.append([f'<a href="/bulk/{b["id"]}">Batch #{b["id"]}</a>', e(b["created_at"]), str(sm["total"]),
                       "Done" if sm["waiting"] == 0 else f'{sm["waiting"]} waiting'])
    body = (ui.page_header("Bulk Upload", "Upload many videos at once. They are queued and reviewed one by one.")
            + '<div class="stack">' + ui.card("New batch", form)
            + ui.card("Your recent batches", ui.table(["Batch", "Started", "Videos", "Status"], recent, "Batches")
                      if recent else ui.empty("No batches yet")) + "</div>")
    return page("Bulk Upload", body, u, "bulk", [("Bulk Upload", None)], request)


@app.post("/bulk")
def bulk_submit(request: Request, brand_id: int = Form(...), language: str = Form("auto"), project_id: str = Form(""),
                files: list[UploadFile] = File(...)):
    u = me(request)
    if not u or u["role"] != "submitter":
        return RedirectResponse("/", 303)
    if not allowed_brand(u, brand_id):
        return err_page(u, "You are not linked to that company.")
    files = [f for f in files if f and f.filename]
    if not files:
        return err_page(u, "Choose at least one file.")
    if len(files) > BULK_MAX_FILES:
        return err_page(u, f"Upload at most {BULK_MAX_FILES} files at a time.")
    if JOBS.pending_count(u["id"]) + len(files) > BULK_MAX_PENDING:
        return err_page(u, "You already have many videos waiting. Wait for them to finish, then upload more.")
    hint = int(project_id) if project_id.isdigit() else None
    if hint and not RE.can_use_project(u["id"], hint):
        return err_page(u, "You do not have access to that project.")
    items, skipped = [], []
    for f in files:
        try:
            path, sha = save_upload(f)
            items.append({"filename": f.filename, "path": path, "sha": sha})
        except Exception as ex_:
            skipped.append(f"{f.filename}: {ex_ if isinstance(ex_, (ValueError, RuntimeError)) else 'upload failed'}")
    if not items:
        return err_page(u, "None of the files could be accepted. " + " ".join(skipped))
    bid = JOBS.create_batch(brand_id, u["id"], items, language if language in ("auto", "mr", "hi", "en") else "auto", hint)
    log(None, u["name"], "BATCH_STARTED", f"Batch {bid}: {len(items)} file(s) queued", brand_id)
    request.session["flash"] = (f"{len(items)} video(s) queued." + (f" Skipped: {' '.join(skipped)}" if skipped else ""))[:400]
    return RedirectResponse(f"/bulk/{bid}", 303)


@app.get("/bulk/{batch_id}", response_class=HTMLResponse)
def bulk_view(request: Request, batch_id: int):
    u = me(request)
    b = JOBS.get_batch(batch_id)
    if not u or not b or not (b["user_id"] == u["id"] or (u["role"] == "admin" and u["brand_id"] == b["brand_id"])):
        return RedirectResponse("/", 303)
    sm = JOBS.summary(batch_id)
    rows = []
    for j in JOBS.batch_jobs(batch_id):
        txt, tone, ic = JOB_BADGE[j["status"]]
        extra = ""
        if j["status"] == "QUEUED":
            extra = f'<span class="muted"> {JOBS.ahead_of(j)} ahead</span>'
        if j["status"] == "DONE" and j["sub_id"]:
            result = f'{ui.chip(j["sub_status"])} <a href="/s/{j["sub_id"]}">Open report</a>'
        elif j["status"] == "FAILED":
            result = e(j["error"] or "")
        else:
            result = ""
        rows.append([e(j["filename"] or ""), ui.badge(txt, tone, ic) + extra, result])
    done = sm["waiting"] == 0
    refresh = "" if done else '<meta http-equiv="refresh" content="6">'
    stats = (f'<p><b>{sm["total"]}</b> videos. <b>{sm["approved"]}</b> approved, <b>{sm["rejected"]}</b> rejected, '
             f'<b>{sm["check"]}</b> need a person, <b>{sm["failed"]}</b> failed'
             + ("." if done else f', <b>{sm["waiting"]}</b> still waiting. This page refreshes by itself.') + '</p>')
    body = (refresh + ui.page_header(f"Batch #{batch_id}", "Videos are reviewed one at a time in the background.")
            + '<div class="stack">' + ui.card("Summary", stats)
            + ui.card("Videos", ui.table(["File", "Progress", "Result"], rows, "Batch videos"))
            + f'<p><a class="btn secondary" href="/bulk">Back to bulk upload</a></p></div>')
    return page(f"Batch #{batch_id}", body, u, "bulk", [("Bulk Upload", "/bulk"), (f"Batch #{batch_id}", None)], request)


# ============================================================ AI review decisions
def _ai_target(request, sid, i):
    u, s = admin_only(request), None
    s = get_sub(sid) if u else None
    if not (u and s and s["brand_id"] == u["brand_id"]):
        return None, None, None, None
    rep = rep_of(s)
    fs = rep.get("findings", [])
    if not (0 <= i < len(fs)) or not fs[i].get("ai") or fs[i]["type"] != "BANNED_PARAPHRASE":
        return u, s, rep, None
    return u, s, rep, fs[i]


def _save_after_ai_decision(s, rep):
    finalize(rep)
    if rep.get("needs_confirmation") and not rep["verdict"].startswith("BLOCKED"):
        rep["verdict"] = "CONFIRM - project needs confirming"
    ex("UPDATE submissions SET report=?, verdict=? WHERE id=?", (json.dumps(rep, ensure_ascii=False), rep["verdict"], s["id"]))


@app.post("/ai/{sid}/{i}/dismiss")
def ai_dismiss(request: Request, sid: int, i: int):
    u, s, rep, f = _ai_target(request, sid, i)
    if not f:
        return RedirectResponse(f"/s/{sid}" if u else "/", 303)
    f["dismissed"] = True
    AIR.dismiss(s["brand_id"], f["evidence"])
    _save_after_ai_decision(s, rep)          # status is left for the builder to decide: a dismissal never auto-approves
    log(sid, u["name"], "AI_DISMISSED", f'"{f["evidence"]}" judged harmless', s["brand_id"])
    request.session["flash"] = "Dismissed. This company will not be shown that wording again."
    return RedirectResponse(f"/s/{sid}", 303)


@app.post("/ai/{sid}/{i}/ban")
def ai_ban(request: Request, sid: int, i: int):
    u, s, rep, f = _ai_target(request, sid, i)
    if not f:
        return RedirectResponse(f"/s/{sid}" if u else "/", 303)
    reg = load_reg(u)
    phrase = norm(str(f["evidence"]))[:120]
    if phrase and phrase.lower() not in {x.lower() for x in reg.get("banned_phrases", [])}:
        reg.setdefault("banned_phrases", []).append(phrase)
        save_reg(u, reg)
    f["confirmed"] = True
    _save_after_ai_decision(s, rep)
    log(sid, u["name"], "AI_BANNED", f'"{phrase}" added to banned phrases (resembles "{f.get("matches", "")}")', s["brand_id"])
    request.session["flash"] = "Added to your banned phrases. Future videos using these words are flagged automatically."
    return RedirectResponse(f"/s/{sid}", 303)


# ============================================================ suggested wording
@app.post("/suggest/{sid}/{i}")
def suggest_wording(request: Request, sid: int, i: int):
    u, s = me(request), get_sub(sid)
    if not (u and s and can_view(u, s)):
        return RedirectResponse("/", 303)
    rep = rep_of(s)
    fs = rep.get("findings", [])
    if not (0 <= i < len(fs)) or fs[i]["type"] not in SUGGESTABLE:
        return RedirectResponse(f"/s/{sid}", 303)
    sugg = rep.setdefault("suggestions", {})
    if str(i) in sugg:
        return RedirectResponse(f"/s/{sid}", 303)
    if len(sugg) >= 10:
        request.session["flash"] = "Enough suggestions for this video. Fix these first."
        return RedirectResponse(f"/s/{sid}", 303)
    f = fs[i]
    reg = json.loads(q("SELECT registry FROM brands WHERE id=?", (s["brand_id"],))[0]["registry"])
    sentence = LRN._sentence_for(rep, f)
    if not sentence:
        sugg[str(i)] = {"sentence": "", "rewrite": None, "error": "The original sentence is not available.", "by": "none"}
    elif f["type"] == "TERMINOLOGY":
        prefer = ""
        for rule in reg.get("terminology", []):
            a_, _, b_ = norm(rule).partition("=>")
            if a_.strip().lower() == str(f["evidence"]).lower():
                prefer = b_.strip()
        text = RW.replace_term(sentence, str(f["evidence"]), prefer)
        sugg[str(i)] = {"sentence": sentence, "rewrite": text, "by": "rule",
                        "error": None if text else "No preferred wording is set for this term."}
    else:
        text, err = RW.ai_rewrite(sentence, str(f["evidence"]), f["explanation"], reg)
        sugg[str(i)] = {"sentence": sentence, "rewrite": text, "by": "ai", "error": err}
    ex("UPDATE submissions SET report=? WHERE id=?", (json.dumps(rep, ensure_ascii=False), sid))
    log(sid, u["name"], "REWRITE_SUGGESTED", f"Finding {i}: {f['type']}", s["brand_id"])
    return RedirectResponse(f"/s/{sid}", 303)


# ============================================================ replay: before and after lessons
def replay_status(s, use_lessons):
    old = rep_of(s)
    inp = old.get("inputs", {})
    reg = json.loads(q("SELECT registry FROM brands WHERE id=?", (s["brand_id"],))[0]["registry"])
    user = q("SELECT * FROM users WHERE id=?", (s["user_id"],))[0]
    screen = None
    if inp.get("screen") is not None and old.get("screen_meta"):
        screen = {"lines": [{"start": l["start"], "end": l["end"], "text": l["text"], "raw": l.get("raw", l["text"]),
                             "conf": l["conf"]} for l in inp["screen"]], **old["screen_meta"]}
    rep = analyze(user, s["brand_id"], inp.get("segments", []), inp.get("caption", ""), reg,
                  confirmed=old.get("project_ids") or None, is_link=bool(s["source_url"]), screen=screen,
                  use_lessons=use_lessons)
    return status_of(rep)


def replay(brand_id, limit=60):
    """Re-run stored disputed videos with lessons off and on. Nothing is saved."""
    rows = q("""SELECT * FROM submissions WHERE brand_id=? AND dispute_note IS NOT NULL
                AND status IN ('APPROVED','REJECTED_FINAL') ORDER BY id DESC LIMIT ?""", (brand_id, limit))
    out = []
    for s in rows:
        old = rep_of(s)
        learnable = any(f["type"] in LRN.LEARNABLE for f in old.get("findings", []))
        try:
            before, after = replay_status(s, False), replay_status(s, True)
        except Exception:
            before = after = "ERROR"
        out.append({"id": s["id"], "file": s["filename"], "overturned": s["status"] == "APPROVED",
                    "learnable": learnable, "before": before, "after": after})
    rej = ("REJECTED", "BLOCKED")
    ov = [r for r in out if r["overturned"]]
    up = [r for r in out if not r["overturned"]]
    return {"rows": out,
            "overturned": len(ov), "overturned_wording": sum(1 for r in ov if r["learnable"]),
            "rejected_before": sum(1 for r in ov if r["before"] in rej),
            "cleared": sum(1 for r in ov if r["before"] in rej and r["after"] not in rej),
            "upheld": len(up), "wrongly_cleared": sum(1 for r in up if r["before"] in rej and r["after"] not in rej)}


@app.get("/admin/learned/replay", response_class=HTMLResponse)
def learned_replay(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    r = replay(u["brand_id"])
    rows = [[f'<a href="/s/{x["id"]}">#{x["id"]}</a> {e(x["file"] or "")}', "Overturned" if x["overturned"] else "Upheld",
             "Yes" if x["learnable"] else "No (price or fact)", ui.chip(x["before"]) if x["before"] in ui.STATUS else e(x["before"]),
             ui.chip(x["after"]) if x["after"] in ui.STATUS else e(x["after"])] for x in r["rows"]]
    stats = (f'<p>Of <b>{r["overturned"]}</b> disputes you overturned, <b>{r["rejected_before"]}</b> are still rejected without lessons. '
             f'With lessons, <b>{r["cleared"]}</b> of those now clear.</p>'
             f'<p>Of <b>{r["upheld"]}</b> disputes you upheld (the rejection was right), <b>{r["wrongly_cleared"]}</b> would wrongly clear. '
             f'This number should always be 0.</p>'
             f'<p class="muted">Only wording mistakes can be learned. Price and fact disputes are never cleared by lessons. '
             f'This replay re-runs the stored transcripts with the current rules and saves nothing.</p>')
    table = (ui.table(["Video", "Your decision", "Wording issue?", "Without lessons", "With lessons"], rows, "Replay")
             if rows else ui.empty("No resolved disputes yet", "Resolve a dispute first, then replay it here."))
    body = (ui.page_header("Replay disputed videos", "Measures what the lessons change, using your own past decisions.")
            + '<div class="stack">' + ui.card("Result", stats) + ui.card("Videos", table)
            + '<p><a class="btn secondary" href="/admin/learned">Back to Lessons</a></p></div>')
    return page("Replay", body, u, "learned", [("Lessons", "/admin/learned"), ("Replay", None)], request)


# ============================================================ lessons
@app.get("/admin/learned", response_class=HTMLResponse)
def learned_page(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    st = LRN.stats(u["brand_id"])
    rows = []
    for c in LRN.list_lessons(u["brand_id"]):
        state = ui.badge("Active", "ok", "check-circle") if c["active"] else ui.badge("Off", "neutral", "clock")
        cue = "any context" if c["scope"] == "phrase" else (c["cue"] or "-")
        note = f'<div class="muted">{e(c["pending_reason"])}</div>' if c["pending_reason"] else ""
        tog = (f'<form method="post" action="/admin/learned/toggle" style="display:inline">'
               f'<input type="hidden" name="lesson_id" value="{c["id"]}">'
               f'<input type="hidden" name="on" value="{0 if c["active"] else 1}">'
               f'<button class="btn secondary sm" type="submit">{"Switch off" if c["active"] else "Switch on"}</button></form>')
        scope = (f'<form method="post" action="/admin/learned/scope" style="display:inline">'
                 f'<input type="hidden" name="lesson_id" value="{c["id"]}">'
                 f'<input type="hidden" name="scope" value="{"context" if c["scope"] == "phrase" else "phrase"}">'
                 f'<button class="btn secondary sm" type="submit">{"Only in this context" if c["scope"] == "phrase" else "Allow anywhere"}</button></form>')
        rows.append([f'<b>{e(c["summary"] or "")}</b>{note}', e(c["trigger"]), e(cue), e(c["example"] or ""),
                     f'<a href="/s/{c["source_sub"]}">#{c["source_sub"]}</a>', str(c["uses"]), state, tog + " " + scope])
    table = (ui.table(["Lesson", "Phrase", "Harmless when", "Example", "From", "Used", "State", ""], rows, "Lessons")
             if rows else ui.empty("No lessons yet", "When you overturn a dispute about wording, a lesson appears here."))
    body = (ui.page_header("Lessons", "What BrandGuard learned from disputes you resolved.")
            + '<div class="stack">'
            + ui.alert(f"Lessons only cover wording and understanding mistakes. {LRN.NEVER_LEARNED.capitalize()} are never learned. "
                       "A lesson clears a flag only when the same phrase appears with the same harmless cue words. "
                       "Anything new still goes to a person or is flagged.", "info", "How this works.")
            + ui.card("Impact", f'<p><b>{st["lessons"]}</b> lessons, <b>{st["active"]}</b> active. '
                      f'<b>{st["cleared"]}</b> flags cleared across <b>{st["videos"]}</b> videos.</p>'
                      f'<p class="muted">AI lesson summaries: {"on" if os.environ.get("GROQ_API_KEY") else "off (lessons use a plain summary and may need your approval)"}.</p>'
                      '<p><a class="btn secondary sm" href="/admin/learned/replay">Replay disputed videos</a></p>')
            + ui.card("Lessons", table) + "</div>")
    return page("Lessons", body, u, "learned", [("Lessons", None)], request)


@app.post("/admin/learned/toggle")
def learned_toggle(request: Request, lesson_id: int = Form(0), on: int = Form(0)):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    if LRN.set_active(u["brand_id"], lesson_id, bool(on)):
        log(None, u["name"], "LESSON_CHANGED", f"Lesson {lesson_id} switched {'on' if on else 'off'}", u["brand_id"])
        request.session["flash"] = "Lesson switched on." if on else "Lesson switched off. That wording is checked normally again."
    return RedirectResponse("/admin/learned", 303)


@app.post("/admin/learned/scope")
def learned_scope(request: Request, lesson_id: int = Form(0), scope: str = Form("context")):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    if LRN.set_scope(u["brand_id"], lesson_id, scope):
        log(None, u["name"], "LESSON_CHANGED", f"Lesson {lesson_id} scope set to {scope}", u["brand_id"])
    return RedirectResponse("/admin/learned", 303)

    


# ============================================================ import brand rules and projects from a PDF (builder)
GUIDE_HEADINGS = {
    "banned_phrases": ("banned phrases", "banned phrase", "banned words", "banned claims", "prohibited phrases",
                       "prohibited claims", "restricted phrases", "do not say", "never say", "phrases to avoid"),
    "confidential": ("confidential", "confidential terms", "confidential information", "internal terms",
                     "never disclose", "do not disclose"),
    "allowed_contacts": ("approved contacts", "approved contact numbers", "approved contact", "authorised contacts",
                         "authorized contacts", "contact numbers", "official contacts"),
    "glossary": ("glossary", "speech glossary", "project names", "key terms", "brand terms"),
}
GUIDE_LABEL = {"banned_phrases": "Banned phrases", "confidential": "Confidential terms",
               "allowed_contacts": "Approved contact numbers", "glossary": "Speech glossary"}
_BULLET = re.compile(r"^\s*(?:[•●▪‣\*\-–—]|\(?\d{1,2}[.)])\s+")

# project lines: "Project: Name", then "City: Pune", then one line per property type
PROJECT_KEYS = {
    "city": ("city",),
    "locality": ("locality", "area", "location", "micro market"),
    "rera": ("rera", "rera number", "rera no", "rera registration"),
    "possession": ("possession", "possession date"),
    "ready": ("ready to move", "ready to move in"),
    "booking": ("booking amount", "token amount"),
    "short": ("short name",),
    "aliases": ("other names", "aliases", "also called", "hindi name", "marathi name"),
}
_PROJ_START = re.compile(r"^(?:project|project name)\s*[:\-–]\s*(.+)$", re.I)
_KV = re.compile(r"^([A-Za-z ]{3,30}?)\s*[:\-–]\s*(.+)$")
_UNIT = re.compile(r"^(.{2,50}?)\s*[:\-–]\s*(.*\bprice\b.*)$", re.I)


def pdf_text(data, filename):
    """Plain text of a PDF or TXT. Scanned PDFs have no text layer and are refused clearly."""
    if filename.lower().endswith(".txt"):
        return data.decode("utf-8", "ignore")
    from pypdf import PdfReader
    import io as _io
    r = PdfReader(_io.BytesIO(data))
    if r.is_encrypted:
        raise ValueError("This PDF is password protected. Remove the password and upload it again.")
    if len(r.pages) > 40:
        raise ValueError("This PDF has more than 40 pages. Upload only the brand guidelines section.")
    text = "\n".join((p.extract_text() or "") for p in r.pages)
    if len(text.strip()) < 20:
        raise ValueError("No readable text was found. This looks like a scanned image, which is not supported yet. "
                         "Upload a PDF exported from a document, or a .txt file.")
    return text


def _heading_key(line):
    """Returns (rule key, inline text after the colon) when the line is a known rules heading."""
    base = re.sub(r"\s+", " ", _BULLET.sub("", line)).strip()
    t = base.strip(":").strip().lower()
    for k, names in GUIDE_HEADINGS.items():
        if t in names:
            return k, ""
    m = re.match(r"^([A-Za-z ]{3,30})\s*:\s*(.+)$", base)
    if m:
        head = m.group(1).strip().lower()
        for k, names in GUIDE_HEADINGS.items():
            if head in names:
                return k, m.group(2)
    return None, ""


def _clean_item(key, s):
    s = re.sub(r"[\x00-\x1f]", " ", s or "")
    s = re.sub(r"\s+", " ", s).strip().strip(".;,").strip()
    if len(s) < 2 or len(s) > 80 or re.match(r"(?i)^page \d+", s):
        return None
    if key == "allowed_contacts":
        compact = re.sub(r"(?<=\d)[\s-](?=\d)", "", s)
        ph = re.search(r"(?<!\d)(?:\+?91)?([6-9]\d{9})(?!\d)", compact)
        if ph:
            return ph.group(1)
        em = re.search(r"[^@\s]+@[^@\s]+\.[A-Za-z]{2,}", s)
        return em.group(0).lower() if em else None
    return s


def _category_of(label):
    t = label.lower()
    if "penthouse" in t:
        return "penthouse"
    if re.search(r"row[\s-]?house", t):
        return "row_house"
    if "villa" in t:
        return "villa"
    if "bungalow" in t:
        return "bungalow"
    if "plot" in t:
        return "plot"
    if re.search(r"apartment|flat|\bbhk\b|\brk\b", t):
        return "apartment"
    return None


def _project_key(line):
    m = _KV.match(line)
    if not m:
        return None, ""
    head = re.sub(r"\s+", " ", m.group(1)).strip().lower()
    for k, names in PROJECT_KEYS.items():
        if head in names:
            return k, m.group(2).strip()
    return None, ""


def _unit_line(line):
    """'2 BHK Apartment: Price 1.3 Cr; Carpet 720 sq ft; Bathrooms 2' -> (unit dict, skipped pieces), or None."""
    m = _UNIT.match(line)
    if not m:
        return None
    label = re.sub(r"\s+", " ", m.group(1)).strip()
    cat = _category_of(label)
    if not cat:
        return None
    bm = re.search(r"(\d)\s*bhk", label, re.I)
    unit = {"label": label[:80], "category": cat, "bhk": int(bm.group(1)) if bm else None}
    skipped = []
    for piece in re.split(r"[;|]", m.group(2)):
        p = piece.strip().strip(".")
        if not p:
            continue
        low = p.lower()
        mm = re.match(r"^price\s*[:\-]?\s*(.+)$", low)
        if mm:
            unit["price"] = p[mm.start(1):][:40]
            continue
        got = False
        for key, pat in (("carpet_sqft", r"^carpet(?:\s+area)?\s*[:\-]?\s*([\d,\.]+)"),
                         ("builtup_sqft", r"^built[\s-]?up(?:\s+area)?\s*[:\-]?\s*([\d,\.]+)"),
                         ("plot_sqft", r"^plot(?:\s+area)?\s*[:\-]?\s*([\d,\.]+)"),
                         ("bathrooms", r"^(?:bathrooms?|baths?)\s*[:\-]?\s*(\d+)"),
                         ("bathrooms", r"^(\d+)\s*(?:bathrooms?|baths?)\b"),
                         ("balconies", r"^(?:balconies|balcony)\s*[:\-]?\s*(\d+)"),
                         ("balconies", r"^(\d+)\s*(?:balconies|balcony)\b"),
                         ("parking", r"^parking(?:\s+spaces?)?\s*[:\-]?\s*(\d+)"),
                         ("parking", r"^(\d+)\s*parking\b")):
            mm = re.match(pat, low)
            if mm:
                unit[key] = mm.group(1).replace(",", "")
                got = True
                break
        if got:
            continue
        for key, pat in (("facing", r"^facing\s*[:\-]?\s*(.+)$"), ("floor_range", r"^floors?(?:\s+range)?\s*[:\-]?\s*(.+)$"),
                         ("dimensions", r"^dimensions?\s*[:\-]?\s*(.+)$"),
                         ("permitted_use", r"^(?:permitted\s+use|use)\s*[:\-]?\s*(.+)$")):
            mm = re.match(pat, low)
            if mm:
                unit[key] = p[mm.start(1):][:60]
                got = True
                break
        if not got:
            skipped.append(p[:60])
    if "price" not in unit:
        return None
    return unit, skipped


def _poss_norm(raw):
    raw = (raw or "").strip()
    m = re.fullmatch(r"(20\d\d)-(0[1-9]|1[0-2])", raw)
    if m:
        return raw
    if re.search(r"[A-Za-z]{3}[a-z]*\s+20\d\d", raw):
        return RE._parse_poss(raw)
    return None


def _is_project_line(line):
    return bool(_PROJ_START.match(line) or _project_key(line)[0] or _unit_line(line))


def parse_guide(text):
    """Reads a brand guide made of plain headings (rules) and 'Project:' blocks (projects, prices, areas).
    No guessing: only recognised lines are taken, unrecognised lines inside a project are listed, and the
    builder confirms everything before it is saved."""
    out = {k: [] for k in GUIDE_HEADINGS}
    projects = []
    curp = None

    def add(key, raw):
        v = _clean_item(key, raw)
        if v and v.lower() not in [x.lower() for x in out[key]] and len(out[key]) < 100:
            out[key].append(v)

    cur = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        pm = _PROJ_START.match(line)
        if pm:
            curp = {"name": re.sub(r"\s+", " ", pm.group(1)).strip()[:120], "short": "", "city": "", "locality": "",
                    "rera": "", "possession": "", "ready": "", "booking": "", "aliases": [], "units": [],
                    "skipped": [], "notes": []}
            if len(projects) < 20:
                projects.append(curp)
            cur = None
            continue
        key, inline = _heading_key(line)
        if key:
            cur, curp = key, None
            for part in re.split(r"[;,]", inline):
                add(cur, part)
            continue
        if curp is not None:
            pk, val = _project_key(line)
            if pk:
                if pk == "aliases":
                    curp["aliases"] += [a.strip()[:80] for a in re.split(r"[;,]", val) if a.strip()]
                elif pk == "possession":
                    norm = _poss_norm(val)
                    if norm:
                        curp["possession"] = norm
                    else:
                        curp["notes"].append(f'Possession "{val[:40]}" needs a month and year, for example December 2028.')
                elif pk == "ready":
                    v = val.strip().lower()
                    if v in ("yes", "no"):
                        curp["ready"] = v
                    else:
                        curp["notes"].append(f'Ready to move "{val[:20]}" must be yes or no.')
                else:
                    curp[pk] = val.strip()[:120]
                continue
            ul = _unit_line(line)
            if ul:
                unit, sk = ul
                curp["units"].append(unit)
                curp["skipped"] += [f'{unit["label"]}: {s}' for s in sk]
                continue
        unknown_heading = (line.endswith(":") and len(line) <= 60) or (line.isupper() and len(line) > 3)
        if unknown_heading:
            cur, curp = None, None
            continue
        if curp is not None:
            if len(curp["skipped"]) < 8:
                curp["skipped"].append(line[:90])
            continue
        if _is_project_line(line):
            continue
        if cur and (_BULLET.match(line) or len(line) <= 80):
            for part in _BULLET.sub("", line).split(";"):
                add(cur, part)
    out["projects"] = projects
    return out


def _s(v, n=120):
    return re.sub(r"\s+", " ", re.sub(r"[\x00-\x1f]", " ", v)).strip()[:n] if isinstance(v, str) else ""


def _norm_projects(raw):
    """Rebuilds the project list from posted JSON, keeping only known text fields."""
    out = []
    for p in (raw if isinstance(raw, list) else [])[:20]:
        if not isinstance(p, dict) or not _s(p.get("name")):
            continue
        pr = {k: _s(p.get(k)) for k in ("name", "short", "city", "locality", "rera", "possession", "ready", "booking")}
        pr["aliases"] = [_s(a, 80) for a in (p.get("aliases") or [])[:20] if _s(a, 80)] if isinstance(p.get("aliases"), list) else []
        pr["units"] = []
        for un in (p.get("units") or [])[:30] if isinstance(p.get("units"), list) else []:
            if not isinstance(un, dict) or un.get("category") not in RE.CATEGORIES or not _s(un.get("label"), 80):
                continue
            nu = {"label": _s(un.get("label"), 80), "category": un["category"],
                  "bhk": un.get("bhk") if isinstance(un.get("bhk"), int) and 1 <= un["bhk"] <= 9 else None}
            for k in ("price", "carpet_sqft", "builtup_sqft", "plot_sqft", "bathrooms", "balconies", "parking",
                      "facing", "floor_range", "dimensions", "permitted_use"):
                if _s(un.get(k), 60):
                    nu[k] = _s(un.get(k), 60)
            pr["units"].append(nu)
        out.append(pr)
    return out


def _price_text(raw):
    v = parse_price(raw or "")
    return fmt_money(v) if v else "could not read"


@app.get("/admin/guide", response_class=HTMLResponse)
def guide_form(request: Request, err: str = ""):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    items_html = "".join("<li><b>" + e(GUIDE_LABEL[k]) + '</b> <span class="muted">(heading, for example: '
                         + e(GUIDE_HEADINGS[k][0].title()) + ")</span></li>" for k in GUIDE_HEADINGS)
    example = ("Project: Demo Lakeview\nCity: Pune\nLocality: Baner\nRERA: P44444444444\nPossession: June 2029\n"
               "Ready to move: No\nBooking amount: 6 lakh\nOther names: Lakeview; Lake View\n"
               "2 BHK Apartment: Price 1.05 Cr; Carpet 640 sq ft; Built-up 790 sq ft; Bathrooms 2; Balconies 2; Parking 1")
    help_ = ('<p class="muted">BrandGuard reads two kinds of content, then shows you everything it found before anything is saved.</p>'
             '<p style="margin-bottom:4px"><b>1. Rules</b> <span class="muted">(a heading, then one item per line)</span></p>'
             "<ul style='margin:0 0 12px;padding-left:20px'>" + items_html + "</ul>"
             '<p style="margin-bottom:4px"><b>2. Projects, prices and facts</b> <span class="muted">(one property type per line)</span></p>'
             f'<pre class="mono" style="white-space:pre-wrap;background:#f7f9fc;border:1px solid var(--line);border-radius:7px;padding:12px;margin:0 0 8px">{e(example)}</pre>'
             '<p class="muted" style="margin:0">Keep each property type on one line, not in a table. Property types can be Apartment, '
             'Penthouse, Villa, Bungalow, Row house or Plot. Scanned images are not supported. A new project is not visible to '
             'any partner until you grant access on the Project Access page.</p>')
    form = ('<form method="post" action="/admin/guide/preview" enctype="multipart/form-data" data-lock>'
            '<label for="gf">Brand guideline document (PDF or TXT)</label>'
            '<input id="gf" type="file" name="file" accept=".pdf,.txt" required>'
            '<p class="hint">Up to 10 MB. Nothing is saved until you confirm the list on the next screen.</p>'
            f'<button class="btn" type="submit">{ui.icon("upload", 16)}Read document</button></form>')
    body = (ui.page_header("Import from a Document",
                           f"Upload your brand guide and BrandGuard turns it into rules, projects and approved prices for {e(load_reg(u).get('brand', ''))}.",
                           '<a class="btn secondary" href="/admin/rules">Back to Brand Rules</a>')
            + '<div class="stack">' + (ui.alert(e(err), "bad") if err else "")
            + ui.card("Upload", form) + ui.card("What the document should look like", help_) + "</div>")
    return page("Import", body, u, "rules", [("Brand Rules", "/admin/rules"), ("Import", None)], request)


@app.post("/admin/guide/preview", response_class=HTMLResponse)
def guide_preview(request: Request, file: UploadFile = File(...)):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    name = (file.filename or "document")[:120]
    if not name.lower().endswith((".pdf", ".txt")):
        return guide_form(request, err="Upload a PDF or a TXT file.")
    data = file.file.read(10 * 1024 * 1024 + 1)
    if len(data) > 10 * 1024 * 1024:
        return guide_form(request, err="This file is larger than 10 MB.")
    try:
        found = parse_guide(pdf_text(data, name))
    except Exception as ex_:
        msg = str(ex_) if isinstance(ex_, ValueError) else "This file could not be read. Try exporting it as a PDF again."
        return guide_form(request, err=msg)
    reg = load_reg(u)
    projects = [p for p in found["projects"] if p["units"] or p["city"] or p["rera"] or p["possession"] or p["booking"]]
    total = sum(len(found[k]) for k in GUIDE_HEADINGS) + len(projects)
    if not total:
        return guide_form(request, err="Nothing was found. Check that the document has headings such as "
                                       "\"Banned phrases\", or blocks that start with \"Project: Name\", with the details on separate lines.")
    cards = []
    for key in GUIDE_HEADINGS:
        items = found[key]
        if not items:
            continue
        have = [x.lower() for x in reg.get(key, [])]
        boxes = ""
        for it in items:
            dup = it.lower() in have
            state = " disabled" if dup else " checked"
            note = ' <span class="muted">(already in your rules)</span>' if dup else ""
            boxes += (f'<label class="check"><input type="checkbox" name="{key}" value="{e(it)}"{state}> '
                      f'{e(it)}{note}</label>')
        cards.append(ui.card(f"{GUIDE_LABEL[key]} ({len(items)} found)", boxes))
    existing = {p["official_name"].lower(): p for p in RE.brand_projects(u["brand_id"], active_only=False)}
    for i, p in enumerate(projects):
        ex_p = existing.get(p["name"].lower())
        state = ui.badge("Already exists: values will be updated", "info", "info") if ex_p else ui.badge("New project", "ok", "plus")
        det = [("City", p["city"]), ("Locality", p["locality"]), ("RERA number", p["rera"]),
               ("Possession", p["possession"]), ("Ready to move", p["ready"]),
               ("Booking amount", _price_text(p["booking"]) if p["booking"] else ""),
               ("Other names", ", ".join(p["aliases"]))]
        dl = "".join(f"<dt>{e(k)}</dt><dd>{e(v)}</dd>" for k, v in det if v)
        rows = []
        for un in p["units"]:
            rows.append([e(un["label"]), e(RE.CATEGORIES[un["category"]][0]),
                         f'{e(un["price"])}<br><span class="muted">reads as {e(_price_text(un["price"]))}</span>',
                         e(un.get("carpet_sqft") or un.get("plot_sqft") or "-"), e(un.get("builtup_sqft") or "-"),
                         e(un.get("bathrooms") or "-"), e(un.get("balconies") or "-"), e(un.get("parking") or "-")])
        tbl = (ui.table(["Property type", "Category", "Price", "Carpet / plot sq ft", "Built-up sq ft", "Bathrooms", "Balconies", "Parking"],
                        rows, "Property types in " + p["name"]) if rows else ui.empty("No property types found for this project"))
        warn = ""
        if p["notes"] or p["skipped"]:
            lines_ = "".join(f"<li>{e(x)}</li>" for x in (p["notes"] + p["skipped"])[:12])
            warn = ui.alert("<b>Not understood, so skipped:</b><ul style='margin:6px 0 0;padding-left:18px'>" + lines_ + "</ul>", "warn")
        pick = (f'<label class="check"><input type="checkbox" name="proj" value="{i}" checked> '
                f'<b>Add this project and its prices</b></label>')
        cards.append(ui.card(p["name"], f'<p>{state}</p>{pick}<dl class="kv" style="margin:10px 0 14px">{dl}</dl>{tbl}{warn}',
                             "Check that each price reads as the amount you expect (lakh versus crore)."))
    payload = e(json.dumps(projects, ensure_ascii=False))
    form = (f'<form method="post" action="/admin/guide/apply" data-lock><input type="hidden" name="title" value="{e(name)}">'
            f'<input type="hidden" name="payload" value="{payload}">'
            + "".join(cards)
            + f'<div class="form-actions" style="margin-top:16px"><button class="btn" type="submit">{ui.icon("check", 16)}Save selected</button>'
              '<a class="btn secondary" href="/admin/guide">Cancel</a></div></form>')
    body = (ui.page_header("Check what was found", f"From {e(name)}. Untick anything that should not be saved.")
            + '<div class="stack">' + ui.alert("Nothing is saved yet. Rules and prices apply to new submissions only. "
                                                "Changing an approved price creates a new version and flags earlier approvals for re-check.", "info")
            + form + "</div>")
    return page("Import", body, u, "rules", [("Brand Rules", "/admin/rules"), ("Import", None)], request)


def _import_fact(u, p, key, uid, raw, title, stats, warns, label):
    """Saves one fact. Unchanged values are skipped, so re-importing never creates empty versions."""
    try:
        num, txt = parse_fact_value(key, raw)
        state, cur, _d = RE.resolve_fact(p["id"], key, uid)
        if state == "ok" and cur is not None:
            same = (abs((cur["value_num"] or 0) - num) < 1e-6) if num is not None else (cur["value_text"] == txt)
            if same:
                stats["unchanged"] += 1
                return
        _save_fact(u, p, key, uid, raw, title, "Imported from brand guide", "", "")
        stats["facts"] += 1
    except ValueError as ex_:
        warns.append(f"{label}: {RE.FACT_LABEL.get(key, key)} not saved. {ex_}")


def _import_project(u, pr, title, stats, warns):
    existing = next((p for p in RE.brand_projects(u["brand_id"], active_only=False)
                     if p["official_name"].lower() == pr["name"].lower()), None)
    if existing:
        pid = existing["id"]
        ex("UPDATE projects SET city=?, locality=?, short_name=? WHERE id=?",
           (pr["city"] or existing["city"], pr["locality"] or existing["locality"], pr["short"] or existing["short_name"], pid))
    else:
        pid = RE.add_project(u["brand_id"], pr["name"], pr["short"], pr["city"], pr["locality"], "", [], u["name"])
        log(None, u["name"], "PROJECT_ADDED", pr["name"] + " (from brand guide)", u["brand_id"])
        stats["projects"] += 1
    p = RE.get_project(pid)
    changed_before = stats["facts"] + stats["units"]
    if pr["rera"]:
        r = pr["rera"].strip().upper()
        if not RE.rera_format_ok(r):
            warns.append(f'{pr["name"]}: RERA number "{pr["rera"]}" was skipped. It should look like P followed by 11 digits.')
        elif r != (p["rera_no"] or ""):
            ex("UPDATE projects SET rera_no=?, rera_level='format', rera_source=NULL, rera_checked_at=NULL, rera_checked_by=NULL WHERE id=?", (r, pid))
    have = [a.lower() for a in RE.aliases_of(pid)]
    for a in pr["aliases"]:
        if a.lower() not in have:
            ex("INSERT INTO project_aliases (project_id, alias) VALUES (?,?)", (pid, a))
            have.append(a.lower())
    p = RE.get_project(pid)
    for key, field in (("possession", pr["possession"]), ("ready_to_move", pr["ready"]), ("booking_amount", pr["booking"])):
        if field:
            _import_fact(u, p, key, None, field, title, stats, warns, pr["name"])
    for un in pr["units"]:
        match = None
        for ut in RE.unit_types(pid, active_only=False):
            if (un["bhk"] is not None and ut["category"] == un["category"] and ut["bhk"] == un["bhk"]) \
                    or ut["label"].lower() == un["label"].lower():
                match = ut
                break
        try:
            if not parse_fact_value("price", un["price"])[0]:
                raise ValueError("Enter a price greater than zero.")
        except ValueError as ex_:
            warns.append(f'{pr["name"]} {un["label"]}: skipped. The price could not be read. {ex_}')
            continue
        if match:
            uid = match["id"]
        else:
            bhk = un["bhk"] if "bhk" in RE.CATEGORIES[un["category"]][1] else None
            uid = RE.add_unit_type(pid, un["category"], un["label"], bhk)
            stats["units"] += 1
        label = f'{pr["name"]} {un["label"]}'
        _import_fact(u, p, "price", uid, un["price"], title, stats, warns, label)
        for key in ("carpet_sqft", "builtup_sqft", "plot_sqft", "bathrooms", "balconies", "parking",
                    "facing", "floor_range", "dimensions", "permitted_use"):
            if un.get(key):
                _import_fact(u, p, key, uid, un[key], title, stats, warns, label)
    if existing and stats["facts"] + stats["units"] > changed_before:
        stats["updated"] += 1
    log(None, u["name"], "PROJECT_EDITED", f'{pr["name"]}: details imported from brand guide', u["brand_id"])


@app.post("/admin/guide/apply")
async def guide_apply(request: Request):
    u = admin_only(request)
    if not u:
        return RedirectResponse("/", 303)
    f = await request.form()
    reg = load_reg(u)
    added = {}
    for key in GUIDE_HEADINGS:
        cur = reg.setdefault(key, [])
        have = [x.lower() for x in cur]
        for raw in f.getlist(key)[:100]:
            v = _clean_item(key, raw)
            if v and v.lower() not in have:
                cur.append(v)
                have.append(v.lower())
                added[key] = added.get(key, 0) + 1
    title = (f.get("title") or "document")[:120]
    stats = {"projects": 0, "updated": 0, "units": 0, "facts": 0, "unchanged": 0}
    warns = []
    chosen = [int(i) for i in f.getlist("proj") if str(i).isdigit()]
    try:
        projects = _norm_projects(json.loads((f.get("payload") or "[]")[:300000]))
    except ValueError:
        projects = []
    source_title = "Brand guide: " + title
    for i in chosen:
        if 0 <= i < len(projects):
            _import_project(u, projects[i], source_title, stats, warns)
    if not added and not (stats["projects"] or stats["updated"] or stats["facts"] or stats["units"]):
        request.session["flash"] = ("Nothing new to save. " + (warns[0] if warns else "Everything selected already matched."))
        return RedirectResponse("/admin/rules", 303)
    if added:
        save_reg(u, reg)
    RE.add_source(u["brand_id"], source_title, "brand_guide", "official", u["name"])
    parts = [f"{n} {GUIDE_LABEL[k].lower()}" for k, n in added.items()]
    if stats["projects"]:
        parts.append(f'{stats["projects"]} new project(s)')
    if stats["updated"]:
        parts.append(f'{stats["updated"]} existing project(s) updated')
    if stats["units"]:
        parts.append(f'{stats["units"]} property type(s)')
    if stats["facts"]:
        parts.append(f'{stats["facts"]} approved fact(s) saved')
    log(None, u["name"], "RULES_IMPORTED", f"{title}: " + ", ".join(parts), u["brand_id"])
    msg = "Imported: " + ", ".join(parts) + "."
    if stats["unchanged"]:
        msg += f' {stats["unchanged"]} value(s) were already up to date.'
    if warns:
        msg += f' {len(warns)} item(s) skipped: ' + warns[0]
    request.session["flash"] = msg[:400]
    return RedirectResponse("/admin/projects" if (stats["projects"] or stats["updated"]) else "/admin/rules", 303)


# ============================================================ demo files for judges (public, no sign-in)
AUDIO_DIR = Path(__file__).resolve().parent / "demo_audio"
AUDIO_EXT = (".mp4", ".webm", ".m4a", ".mp3", ".wav", ".ogg", ".aac")
VIDEO_EXT = (".mp4", ".webm")
TEST_CLIPS = [
    {"key": "approval", "n": 1, "title": "Approval test",
     "say": "Demo Heights in Wakad, Pune. A 2 BHK home at 1 crore 30 lakh. Carpet area 720 square feet, with 2 bathrooms. "
            "Possession in December 2028.",
     "result": "Auto-Approved. The price, carpet area, bathrooms and possession all match the approved facts, so an "
               "approval ID is issued with no human needed.",
     "steps": ["Download the file below.", "Click Enter as Freelancer on the home page.",
               "Open Submit a Video, choose Demo Developer, and upload the file.",
               "The result page opens. Look at the Claim-by-claim checks table: each spoken claim is matched to an approved fact."]},
    {"key": "rejection", "n": 2, "title": "Rejection test",
     "say": "Demo Heights in Wakad. 2 BHK for 1 crore 30 lakh. Guaranteed returns, the price will double in two years. "
            "Last chance!",
     "result": "Rejected. The price is correct, but \"guaranteed returns\" and \"price will double\" are banned phrases "
               "(\"last chance\" is also flagged if the sample brand guide has been imported). Each finding shows the "
               "exact second it appeared, whether it was spoken or shown on screen, why it matters and what to change.",
     "steps": ["Download the file below.", "Enter as Freelancer and open Submit a Video.",
               "Choose Demo Developer and upload the file.",
               "Read the findings on the result page. The creator is told exactly what to fix."]},
    {"key": "dispute", "n": 3, "title": "Dispute test",
     "say": "Demo Heights, Wakad. 3 BHK flats available at 1 crore 50 lakh. Possession in December 2028.",
     "result": "Rejected, because 3 BHK is approved at 1 crore 90 lakh, not 1 crore 50 lakh. You then dispute it and "
               "play the builder to resolve it. This shows that a person always has the final word.",
     "steps": ["Download the file below, then Enter as Freelancer, open Submit a Video, choose Demo Developer and upload it.",
               "The submission is Rejected for a price mismatch. Click Raise a dispute.",
               "Write why the finding is wrong (for example: \"I quoted the offer price for a limited floor\"). "
               "Optionally type what you actually said. Click Submit dispute. The status becomes Disputed.",
               "In the sidebar, switch the demo role to Builder. Open Review Queue, then open the disputed submission.",
               "The builder sees the evidence and the creator's explanation. Click Overturn decision to approve it "
               "(an approval ID is issued), or Uphold rejection to keep it rejected (a reason is required).",
               "Switch back to Freelancer to see the outcome on the submission, and open Audit Trail as Builder to see "
               "every step recorded with who did it and when."]},
    {"key": "ocr", "file": "test", "n": 4, "title": "On-screen text (OCR) test video",
     "say": "",
     "result": "BrandGuard reads the text on screen frame by frame. On the result page, the check \"On-screen text\" shows "
               "Done, and the card \"Text read from the screen\" lists what appeared in the video, such as the price "
               "(₹1,30,00,000), carpet area, bathrooms and possession date, each with its time. The on-screen price is "
               "compared with the approved price and with the price spoken in the video.",
     "steps": ["Download the video below.", "Click Enter as Freelancer, open Submit a Video, choose Demo Developer and upload it.",
               "Wait for the review. On the free server this takes a while (see the note at the top of this page).",
               "On the result page, open the \"Text read from the screen\" card and the \"On-screen text\" and "
               "\"Voice versus screen\" checks."]},
]


def _clip_file(key):
    clip = next((c for c in TEST_CLIPS if c["key"] == key), None)
    base = clip.get("file", key) if clip else key
    for ext in AUDIO_EXT:
        p = AUDIO_DIR / f"{base}{ext}"
        if p.is_file():
            return p
    return None


@app.get("/test-audio", response_class=HTMLResponse)
def test_audio_page(request: Request):
    u = me(request)
    cards = []
    for c in TEST_CLIPS:
        f = _clip_file(c["key"])
        if f:
            src = f'/test-audio/file/{c["key"]}?inline=1'
            if f.suffix.lower() in VIDEO_EXT:
                player = (f'<video controls preload="none" playsinline style="width:100%;max-height:360px;margin:8px 0 12px;'
                          f'background:#000;border-radius:8px" src="{src}">Your browser cannot play video.</video>')
            else:
                player = (f'<audio controls preload="none" style="width:100%;margin:8px 0 12px" src="{src}">'
                          f'Your browser cannot play audio.</audio>')
            media = player + (f'<a class="btn" href="/test-audio/file/{c["key"]}" download>{ui.icon("upload", 16)}'
                              f'Download demo file {c["n"]}</a>')
        else:
            media = ui.alert("This demo file has not been added yet.", "warn")
        steps = "".join(f"<li>{e(s)}</li>" for s in c["steps"])
        said = (f'<p class="muted" style="margin-bottom:4px">What is said in the clip</p>'
                f'<p style="margin-top:0"><i>"{e(c["say"])}"</i></p>') if c["say"] else ""
        inner = (said + media
                 + f'<h3 class="sec-title">What to do</h3><ol style="margin:0 0 12px;padding-left:20px">{steps}</ol>'
                 f'<h3 class="sec-title">What you should see</h3><p style="margin:0">{e(c["result"])}</p>')
        cards.append(ui.card(f'Demo file {c["n"]}: {c["title"]}', inner, cls=""))
    guide = AUDIO_DIR / "Demo_Developer_Brand_Guide.pdf"
    extra = ""
    if guide.is_file():
        extra = ui.card("Optional: sample brand guide",
                        '<p class="muted">As Builder, open Brand Rules and click Import from PDF, then upload this file. '
                        'BrandGuard lists the rules it found, you confirm them, and they apply to new submissions straight away.</p>'
                        '<a class="btn secondary" href="/test-audio/file/guide" download>Download sample brand guide (PDF)</a>')
    start = ('<a class="btn" href="/enter/freelancer">Enter as Freelancer</a> '
             '<a class="btn secondary" href="/enter/admin">Enter as Builder</a>') if not u else ""
    body = (ui.page_header("Demo Files for Judges",
                           "Ready-made demo recordings and a demo video to test the main workflow: one approved, one rejected, "
                           "one disputed, and one that shows on-screen text being read. No sign-up is needed. Use the demo "
                           "roles in the sidebar or the buttons below.", start)
            + '<div class="stack">'
            + ui.alert("These files are made for demonstration only. Download a file here, then upload it on the Submit a "
                       "Video page to see how BrandGuard reviews it. The company, projects and prices are fictional.",
                       "info", "For demo use.")
            + ui.alert("This demo runs on a free Render server, which has very little CPU and memory. A review needs speech "
                       "recognition plus reading every frame of the video, so it can take a few minutes, and the first "
                       "request after a quiet period is slower because the server has to wake up. Please keep the page open "
                       "and wait. The same review takes seconds on a normal server.",
                       "warn", "Please be patient.")
            + "".join(cards) + extra + "</div>")
    if u:
        return page("Demo Files", body, u, "testaudio", [("Demo Files", None)], request)
    return page("Demo Files", body)


@app.get("/test-audio/file/{key}")
def test_audio_file(key: str, inline: int = 0):
    if key == "guide":
        p = AUDIO_DIR / "Demo_Developer_Brand_Guide.pdf"
        name = "Demo_Developer_Brand_Guide.pdf"
    else:
        clip = next((c for c in TEST_CLIPS if c["key"] == key), None)
        p = _clip_file(key) if clip else None
        name = f'BrandGuard_Demo_{clip["n"]}_{re.sub(r"[^A-Za-z0-9]+", "_", clip["title"]).strip("_")}{p.suffix}' if p else ""
    if not p or not p.is_file():
        return HTMLResponse("File not found.", status_code=404)
    return FileResponse(str(p), filename=name, content_disposition_type="inline" if inline else "attachment")
