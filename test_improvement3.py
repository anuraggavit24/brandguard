"""Checks for Improvement 3: any company, generic fact types, spreadsheet and brochure import, sample data.
Run: python test_improvement3.py"""
import io, json, os, re, subprocess, sys, tempfile
os.environ["DATA_DIR"] = tempfile.mkdtemp()
os.environ.pop("GROQ_API_KEY", None)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import main
import generic as GEN
import importer as IMP
import realestate as RE
import ai_review as AIR
from fastapi.testclient import TestClient

wav = os.path.join(os.environ["DATA_DIR"], "t.mp4")
subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=3", "-f", "lavfi", "-i",
                "color=c=navy:s=320x240:d=3", "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", wav],
               capture_output=True)
SCRIPT = {"t": ""}
main.transcribe = lambda w, l, g: [{"start": float(i), "end": float(i) + 1, "text": t, "conf": 0.95}
                                   for i, t in enumerate(SCRIPT["t"].split("||"))] if SCRIPT["t"] else []
main.OCR.read_screen = lambda *a, **k: None
fails = []


def check(name, cond, extra=""):
    print(("PASS " if cond else "FAIL ") + name + (f"  [{extra}]" if extra and not cond else ""))
    if not cond:
        fails.append(name)


# ---------------- value parsing and matching helpers
check("percent value", GEN.parse_value("percent", "30%") == (30.0, None))
check("money value in words", GEN.parse_value("money", "9,470")[0] == 9470 and GEN.parse_value("money", "1.2 lakh")[0] == 120000)
check("date value", GEN.parse_value("date", "15 March 2027") == (None, "2027-03-15") and GEN.parse_value("date", "2027-03") == (None, "2027-03"))
check("yes/no value", GEN.parse_value("yes_no", "Yes") == (None, "yes"))
for bad in [("percent", "150"), ("date", "soon"), ("yes_no", "maybe"), ("money", "abc")]:
    try:
        GEN.parse_value(*bad)
        check(f"bad value refused: {bad}", False)
    except ValueError:
        check(f"bad value refused: {bad}", True)
check("dates found in speech", sorted(d[0] for d in GEN.find_dates("delivery by March 2027 or 15/04/2027")) == ["2027-03", "2027-04-15"])
check("number with unit", GEN.number_with_unit("gives 22 kmpl on highway", ["kmpl", "km/l"]) == [(22.0, "22 kmpl")])
d_ins = {"label": "Free insurance", "aliases": ["मोफत विमा"]}
check("polarity: positive", GEN.polarity(GEN.fold("get free insurance with this car"), d_ins) == "yes")
check("polarity: negated", GEN.polarity(GEN.fold("there is no free insurance"), d_ins) == "no")
check("compare directions", GEN.compare(20, 22, "max", 0.01) and not GEN.compare(25, 22, "max", 0.01)
      and GEN.compare(9, 8.75, "min", 0.01) and not GEN.compare(7.5, 8.75, "min", 0.01))

# ---------------- end to end: a car and loan company, no property words anywhere
with TestClient(main.app) as a:
    a.post("/signup", data={"company": "Zippy Motors", "name": "Meera", "email": "meera@zippy.test",
                            "password": "a-long-password-1", "industry": "finance"})
    brand = main.q("SELECT * FROM brands WHERE name='Zippy Motors'")[0]
    BID = brand["id"]
    check("non-property company is a general company", json.loads(brand["registry"])["domain"] == "general")
    page = a.get("/admin/projects").text
    check("products page uses product words, not RERA", "Products and Facts" in page and "RERA" not in page)
    a.post("/admin/sample")
    projs = {p["official_name"]: p for p in RE.brand_projects(BID)}
    check("sample data creates the car and the loan", set(projs) == {"Zippy Hatchback", "Easy Car Loan"}, str(set(projs)))
    check("sample data creates the fact types", {d["label"] for d in GEN.list_defs(BID)} >= {"Discount", "Mileage", "Interest rate", "Processing fee"})
    car, loan = projs["Zippy Hatchback"], projs["Easy Car Loan"]
    check("variants are products, not apartments", {u["category"] for u in RE.unit_types(car["id"])} == {"product"})
    detail = a.get(f"/admin/projects/{car['id']}").text
    check("product page shows its own fact types and no RERA card", "Your own fact types" in detail and "Mileage" in detail and "RERA number" not in detail)
    a.post(f"/admin/projects/{car['id']}/unit", data={"category": "product", "label": "ZXi", "price": "", "source": ""})
    check("a variant can be added without a price", any(u["label"] == "ZXi" for u in RE.unit_types(car["id"])))
    d_disc = GEN.get_def(BID, "Discount")
    a.post(f"/admin/projects/{car['id']}/gfact", data={"def_id": d_disc["id"], "value": "abc", "source": "x"})
    check("a bad value is refused with a message", "percentage" in a.get(f"/admin/projects/{car['id']}").text.lower())

    # invite a dealer to both products
    r = a.post("/admin/team/invite", data={"email": "dealer@x.test", "kind": "partner", "project_id": [str(car["id"]), str(loan["id"])]})
    tok = re.search(r"/invite/([\w-]+)", r.text).group(1)
    p = TestClient(main.app)
    p.post(f"/invite/{tok}", data={"name": "Dealer Dev", "password": "dealer-strong-pass"})

    def submit(text, pid=""):
        SCRIPT["t"] = text
        r = p.post("/upload", data={"brand_id": str(BID), "language": "en", "project_id": str(pid)},
                   files={"file": ("t.mp4", open(wav, "rb"), "video/mp4")}, follow_redirects=False)
        sid = int(r.headers["location"].rsplit("/", 1)[-1])
        row = main.q("SELECT * FROM submissions WHERE id=?", (sid,))[0]
        rep = main.rep_of(row)
        return row["status"], rep, {(c["kind"], c["verdict"]) for c in rep["claims"]}

    st, rep, cl = submit("The new Zippy Hatchback VXi at 7.1 lakh with 5% discount and 22 kmpl mileage")
    check("correct price, discount and mileage are all verified", st == "AUTO_APPROVED" and ("custom", "VERIFIED") in cl
          and not any(v != "VERIFIED" for _, v in cl), f"{st} {cl}")
    st, rep, cl = submit("Zippy Hatchback with flat 10% discount")
    check("a higher discount than approved is a violation", st == "REJECTED" and ("custom", "VIOLATION") in cl, f"{st} {cl}")
    st, rep, cl = submit("Zippy Hatchback gives 18 kmpl")
    check("a lower mileage than approved is fine (rule: must not be more)", ("custom", "VERIFIED") in cl and st != "REJECTED", f"{st} {cl}")
    st, rep, cl = submit("Zippy Hatchback gives 28 kmpl")
    check("a higher mileage is a violation", ("custom", "VIOLATION") in cl, f"{cl}")
    st, rep, cl = submit("Zippy Hatchback comes with free insurance")
    check("a yes/no fact claimed when not approved is a violation", ("custom", "VIOLATION") in cl, f"{cl}")
    st, rep, cl = submit("Easy Car Loan at interest rate of just 7.5%")
    check("an interest rate below the approved rate is a violation", st == "REJECTED" and ("custom", "VIOLATION") in cl, f"{st} {cl}")
    st, rep, cl = submit("Easy Car Loan interest rate 9% and processing fee 999 rupees")
    check("rate above the minimum and the exact fee are verified, the fee is not read as a car price",
          ("custom", "VERIFIED") in cl and not any(k == "price" for k, _ in cl) and st == "AUTO_APPROVED", f"{st} {cl}")
    st, rep, cl = submit("Easy Car Loan processing fee only 499 rupees")
    check("a wrong fee is a violation", ("custom", "VIOLATION") in cl, f"{cl}")
    st, rep, cl = submit("Zippy Hatchback VXi 7.1 lakh and get 12% cashback")
    check("an unknown percentage goes to a person", ("custom", "UNVERIFIABLE") in cl and st == "SUPERVISION", f"{st} {cl}")
    st, rep, cl = submit("झिप्पी हॅचबॅक वर पाच टक्के सूट")
    check("a short Marathi name alone asks the dealer to confirm the product", st == "NEEDS_CONFIRMATION", st)
    st, rep, cl = submit("झिप्पी हॅचबॅक वर पाच टक्के सूट", car["id"])
    check("Marathi discount claim is verified once the product is known", ("custom", "VERIFIED") in cl, f"{cl} {rep['identity']['state']}")
    st, rep, cl = submit("Zippy Hatchback with 2 BHK style space and 1200 sq ft boot")
    check("property checks do not run for a car company", not any(k in ("area", "bathrooms") for k, _ in cl), f"{cl}")
    rpage = p.get("/reviews").text if p.get("/reviews").status_code == 200 else ""

    # small product, small price, Marathi words (the original question)
    a.post("/admin/projects/add", data={"official_name": "Glow Mixer", "aliases": "mixer"})
    gm = [x for x in RE.brand_projects(BID) if x["official_name"] == "Glow Mixer"][0]
    a.post(f"/admin/projects/{gm['id']}/unit", data={"category": "product", "label": "Standard", "price": "9470", "source": "Price list"})
    a.post(f"/admin/projects/{gm['id']}/gfact", data={"def_id": d_disc["id"], "value": "30", "source": "Festive offer"})
    tok2 = None
    pid_gm = gm["id"]
    ru = main.q("SELECT id FROM users WHERE email='dealer@x.test'")[0]["id"]
    RE.grant(ru, BID, pid_gm, "test")
    st, rep, cl = submit("Glow Mixer 9 hazar charshe satar rupate la with 30% discount")
    check("small product: 9,470 rupees in Marathi words and 30% are verified", st == "AUTO_APPROVED"
          and ("price", "VERIFIED") in cl and ("custom", "VERIFIED") in cl, f"{st} {cl}")
    st, rep, cl = submit("Glow Mixer 8 hazar rupaye with 30% discount")
    check("small product: wrong small price is caught", ("price", "VIOLATION") in cl or ("price", "REVIEW") in cl, f"{cl}")

    # ---------------- spreadsheet import (one sheet, plain columns, instructions on the website)
    data = a.get("/admin/import/template").content
    check("template downloads as xlsx", data[:2] == b"PK")
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    heads = [x.value for x in ws[1]]
    check("template is one sheet of plain column names, no instructions inside", len(wb.worksheets) == 1 and heads[0] == "Product"
          and "Discount (%)" in heads and ws.max_row == 1, str(heads))
    check("import page explains every column", all(h in a.get("/admin/import").text for h in heads)
          and "Add your own columns" in a.get("/admin/import").text)
    ws.cell(row=1, column=len(heads) + 1, value="Mileage (kmpl)")
    ws.append(["Quick Scooter", "Quicky", "Base", "85000", "10", "2026-12-31", "2", "", "", "45"])
    ws.append(["Quick Scooter", "", "Pro", "95000", "abc", "", "", "", "", "40"])
    buf = io.BytesIO(); wb.save(buf)
    r = a.post("/admin/import/sheet", files={"file": ("mine.xlsx", buf.getvalue(), "application/octet-stream")}, follow_redirects=False)
    iid = int(r.headers["location"].rsplit("/", 1)[-1])
    prev = a.get(f"/admin/import/{iid}").text
    _, dat = IMP.load(BID, iid)
    texts = {x["text"]: x for x in dat["rows"]}
    check("preview lists rows and the problems", "Quick Scooter" in prev and any(not x["ok"] for x in dat["rows"]))
    check("a new column with a unit becomes a number fact type", any(x["kind"] == "fact_type" and x["data"]["label"] == "Mileage"
          and x["data"]["ftype"] == "number" and x["data"]["unit"] == "kmpl" and x["data"]["direction"] == "max" for x in dat["rows"]))
    check("a bad value is marked, not saved", not texts["Quick Scooter / Pro: Discount = abc"]["ok"])
    check("nothing is saved before confirming", not any(x["official_name"] == "Quick Scooter" for x in RE.brand_projects(BID)))
    good = [str(x["id"]) for x in dat["rows"] if x["ok"]]
    bad = [str(x["id"]) for x in dat["rows"] if not x["ok"]]
    a.post(f"/admin/import/{iid}/apply", data={"row": good + bad, "source": "Price list Nov 2026"})
    qs = [x for x in RE.brand_projects(BID) if x["official_name"] == "Quick Scooter"][0]
    base = [x for x in RE.unit_types(qs["id"]) if x["label"] == "Base"][0]
    check("confirmed rows are saved", RE.resolve_fact(qs["id"], "price", base["id"])[1]["value_num"] == 85000)
    disc = RE._q("SELECT * FROM facts WHERE project_id=? AND key='g:discount'", (qs["id"],))
    check("the discount is saved with its end date, the bad one is not", len(disc) == 1 and disc[0]["valid_to"] == "2026-12-31")
    r2 = a.post(f"/admin/import/{iid}/apply", data={"row": good}, follow_redirects=False)
    check("an import cannot be saved twice", r2.headers["location"] == "/admin/import")
    r3 = a.post("/admin/import/sheet", files={"file": ("x.xlsx", b"not excel", "application/octet-stream")})
    check("a broken file gets a clear message", "could not be read" in r3.text)
    old = io.BytesIO()
    from openpyxl import Workbook
    ow = Workbook(); ow.active.title = "Products"; ow.active.append(["Product", "Other names", "Variant", "Price"])
    ow.active.append(["Old Format Fan", "", "", "2,499"]); ft = ow.create_sheet("Fact types"); ft.append(["Name", "Type", "Unit", "Other words", "Rule"])
    fa = ow.create_sheet("Facts"); fa.append(["Product", "Variant", "Fact", "Value", "Valid until"]); ow.save(old)
    check("the older three-sheet template is still read", IMP.parse_xlsx(old.getvalue())["products"][0]["name"] == "Old Format Fan")

    # ---------------- brochure import (AI mocked)
    doc = ("Sparkle Water Purifier. Model S1 price Rs 12,999. Model S2 price Rs 15,499. Festive discount 10%. "
           "Warranty 2 years. Free installation.")
    IMP.call_model = lambda s, u_: json.dumps({
        "products": [{"name": "Sparkle Water Purifier", "aliases": ["Sparkle"], "variants": [
            {"name": "S1", "price": "12,999"}, {"name": "S2", "price": "15,499"}, {"name": "S3", "price": "18,999"}]}],
        "fact_types": [{"label": "Festive discount", "ftype": "percent", "unit": "", "aliases": [], "direction": "max"},
                       {"label": "Purifier warranty", "ftype": "number", "unit": "years", "aliases": [], "direction": "max"}],
        "values": [{"product": "Sparkle Water Purifier", "variant": "", "fact": "Festive discount", "value": "10%"},
                   {"product": "Sparkle Water Purifier", "variant": "", "fact": "Purifier warranty", "value": "5"}]})
    r = a.post("/admin/import/brochure", files={"file": ("brochure.txt", doc.encode(), "text/plain")}, follow_redirects=False)
    iid2 = int(r.headers["location"].rsplit("/", 1)[-1])
    _, dat2 = IMP.load(BID, iid2)
    rows2 = {x["text"]: x for x in dat2["rows"]}
    check("AI values found in the document are ready", rows2["Sparkle Water Purifier / S1: 12,999"]["ok"])
    check("an invented price is marked not found in the document", not rows2["Sparkle Water Purifier / S3: 18,999"]["ok"]
          and "not found" in rows2["Sparkle Water Purifier / S3: 18,999"]["problem"])
    check("an invented value is marked not found", not rows2["Sparkle Water Purifier: Purifier warranty = 5"]["ok"])
    a.post(f"/admin/import/{iid2}/apply", data={"row": [str(x["id"]) for x in dat2["rows"] if x["ok"]], "source": "Sparkle brochure"})
    sp = [x for x in RE.brand_projects(BID) if x["official_name"] == "Sparkle Water Purifier"][0]
    check("unticked invented rows were not saved", {u["label"] for u in RE.unit_types(sp["id"])} == {"S1", "S2"})
    IMP.call_model = lambda s, u_: (_ for _ in ()).throw(RuntimeError("AI is not set up on this server. Use the spreadsheet template instead."))
    r = a.post("/admin/import/brochure", files={"file": ("b.txt", doc.encode(), "text/plain")})
    check("without AI the brochure path explains itself", "Use the spreadsheet" in r.text)
    r = a.post("/admin/import/brochure", files={"file": ("b.exe", b"MZ", "application/octet-stream")})
    check("unsupported brochure type is refused", "Upload a PDF" in r.text)

    # ---------------- property companies keep working, and their sample is property-shaped
    h = TestClient(main.app)
    h.post("/signup", data={"company": "Homes Co", "name": "H", "email": "h@homes.test", "password": "a-long-password-1",
                            "industry": "real_estate"})
    hb = main.q("SELECT * FROM brands WHERE name='Homes Co'")[0]
    check("real estate signup is a property company", json.loads(hb["registry"])["domain"] == "property")
    h.post("/admin/sample")
    hp = RE.brand_projects(hb["id"])
    check("property sample creates BHK apartments", hp and {u["category"] for u in RE.unit_types(hp[0]["id"])} == {"apartment"})
    check("property pages still show RERA", "RERA" in h.get(f"/admin/projects/{hp[0]['id']}").text)

print("\n%d failed" % len(fails) if fails else "\nALL PASSED")
sys.exit(1 if fails else 0)
