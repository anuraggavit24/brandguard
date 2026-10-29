"""Checks for bulk invites from a spreadsheet and the email settings. Run: python test_invites_bulk.py"""
import io, json, os, re, sys, tempfile
os.environ["DATA_DIR"] = tempfile.mkdtemp()
for k in ("GROQ_API_KEY", "BREVO_API_KEY", "MAIL_FROM", "SMTP_HOST"):
    os.environ.pop(k, None)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import main
import mailer
import realestate as RE
from fastapi.testclient import TestClient
from openpyxl import load_workbook

fails = []


def check(name, cond, extra=""):
    print(("PASS " if cond else "FAIL ") + name + (f"  [{extra}]" if extra and not cond else ""))
    if not cond:
        fails.append(name)


SENT = []
with TestClient(main.app) as a:
    a.post("/signup", data={"company": "Zippy", "name": "M", "email": "m@z.test", "password": "a-long-password-1", "industry": "ecommerce"})
    a.post("/admin/sample")
    bid = main.q("SELECT id FROM brands WHERE name='Zippy'")[0]["id"]
    page = a.get("/admin/team").text
    check("team page offers the spreadsheet invite and explains email", "Invite many people from a spreadsheet" in page
          and "BREVO_API_KEY" in page)
    data = a.get("/admin/team/template").content
    wb = load_workbook(io.BytesIO(data)); ws = wb.active
    check("invite template has plain columns", [c.value for c in ws[1]] == ["Email", "Name", "Role", "Access"])
    ws.append(["ravi@x.test", "Ravi", "Channel partner", "All"])
    ws.append(["sara@x.test", "Sara", "Freelancer", "Zippy Hatchback"])
    ws.append(["bad-email", "", "Agency", ""])
    ws.append(["ravi@x.test", "", "Agency", ""])
    ws.append(["joe@x.test", "", "Plumber", ""])
    ws.append(["kim@x.test", "", "Agency", "Unknown Thing"])
    ws.append(["mgr@x.test", "", "Manager", ""])
    buf = io.BytesIO(); wb.save(buf)
    r = a.post("/admin/team/bulk", files={"file": ("people.xlsx", buf.getvalue(), "x")}, follow_redirects=False)
    loc = r.headers["location"]
    prev = a.get(loc).text
    check("preview shows 3 ready and explains each skipped row", "3 of 7 rows are ready" in prev and "listed twice" in prev
          and "not one of Channel partner" in prev and "Unknown Thing" in prev and "not an email" in prev)
    check("nothing is created before confirming", not main.q("SELECT 1 FROM invites WHERE email='ravi@x.test'"))
    res = a.post(loc + "/send").text
    check("without email set up, links are shown once to send yourself", "Send yourself" in res and res.count("/invite/") >= 3)
    check("three invites are created", len(main.q("SELECT 1 FROM invites WHERE brand_id=?", (bid,))) == 3)
    sara = main.q("SELECT project_ids FROM invites WHERE email='sara@x.test'")[0]["project_ids"]
    car = [p["id"] for p in RE.brand_projects(bid) if p["official_name"] == "Zippy Hatchback"][0]
    check("named access is kept per person", json.loads(sara) == [car])
    check("'All' access is kept", json.loads(main.q("SELECT project_ids FROM invites WHERE email='ravi@x.test'")[0]["project_ids"]) == ["*"])
    check("the same list cannot be sent twice", a.post(loc + "/send", follow_redirects=False).status_code == 303)

    # with email configured (Brevo mocked)
    os.environ["BREVO_API_KEY"], os.environ["MAIL_FROM"] = "k", "noreply@z.test"
    mailer._brevo = lambda to, subject, text: (SENT.append((to, subject, text)) or (True, ""))
    csv = "Email,Name,Role,Access\nanu@x.test,Anu,Agency,All\n".encode()
    r = a.post("/admin/team/bulk", files={"file": ("people.csv", csv, "text/csv")}, follow_redirects=False)
    res = a.post(r.headers["location"] + "/send").text
    check("with email set up, each person is emailed their own link", "Emailed" in res and SENT and SENT[-1][0] == "anu@x.test"
          and "/invite/" in SENT[-1][2] and "Hello Anu" in SENT[-1][2])
    r = a.post("/admin/team/invite", data={"email": "solo@x.test", "kind": "partner", "all_projects": "1"})
    check("a single invite is emailed and says so", "Invite emailed to solo@x.test" in r.text and SENT[-1][0] == "solo@x.test")
    r = a.post("/admin/email-test")
    check("test email goes to the builder", SENT[-1][0] == "m@z.test")
    mailer._brevo = lambda to, subject, text: (False, "Brevo said HTTP 401: Key not found")
    r = a.post("/admin/team/invite", data={"email": "fail@x.test", "kind": "partner", "all_projects": "1"})
    check("a failed email gives the reason and the link is still shown", "Key not found" in r.text and "/invite/" in r.text)

    # the tour points at the menu item, not the logo
    import ui
    check("tour looks for menu links inside the menu, not the logo", "$('#side nav a[href=" in ui.JS)

print("\n%d failed" % len(fails) if fails else "\nALL PASSED")
sys.exit(1 if fails else 0)
